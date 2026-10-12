"""gen_pgas_pptx.py: PGAS と BAR1 の記事の静止画 7 枚。"""
import os
import sys

sys.path.insert(0, "/Users/akazawt/works/data-science/tools/pptxkit")
from pptx.dml.color import RGBColor  # noqa: E402
from pptx.enum.dml import MSO_LINE_DASH_STYLE  # noqa: E402
from pptx.enum.shapes import MSO_SHAPE  # noqa: E402

from pptxkit import Deck, shapes as sh  # noqa: E402

TPL = "/Users/akazawt/works/data-science/tools/pptxkit/templates/aws-dark-2024.pptx"
OUT = sys.argv[1] if len(sys.argv) > 1 else "pgas-figures.pptx"

GRN = RGBColor(36, 160, 70)
GRN_HI = RGBColor(59, 209, 111)
GRN_DK = RGBColor(18, 64, 34)
CHIP = RGBColor(22, 26, 30)
EDGE = RGBColor(70, 74, 80)
GREY = RGBColor(52, 56, 64)
LINK = RGBColor(40, 150, 240)
LINK_DK = RGBColor(16, 44, 72)
YEL = RGBColor(251, 211, 50)
YEL_DK = RGBColor(80, 66, 14)
MUTED = RGBColor(150, 150, 158)
WHITE = RGBColor(245, 245, 248)
DASH = MSO_LINE_DASH_STYLE.DASH

d = Deck(TPL, check_layout=False)
t = d.theme


def label(s, x, y, w, h, body, size=16, color=WHITE, bold=False, align="center"):
    return sh.textbox(s, t, x, y, w, h, body, size=size, color=color, bold=bold, align=align,
                      anchor="middle")


def box(s, x, y, w, h, fill, stroke=None, body="", size=15, color=WHITE, radius=0.1,
        dash=None, sw=1.25):
    sp = sh.rect(s, t, x, y, w, h, fill=fill, stroke=stroke, sw=sw, dash=dash,
                 shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=radius)
    if body:
        sh._write(sp.text_frame, t, body, size, color, True, "center", "middle")
        tf = sp.text_frame
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = sh.U(2)
    return sp


def title(s, head, sub):
    label(s, 43, 26, 1190, 46, head, size=30, bold=True, align="left")
    label(s, 43, 72, 1190, 26, sub, size=16, color=MUTED, align="left")


def arrow(s, x1, y1, x2, y2, color=WHITE, width=2, dashed=False):
    sh.arrow(s, t, x1, y1, x2, y2, color=color, width=width, dashed=dashed)




PNK = RGBColor(255, 120, 160)
PNK_DK = RGBColor(80, 30, 46)


def gpu(s, x, y, w, h, name, sub=None):
    box(s, x, y, w, h, GRN_DK, stroke=GRN, body=name if not sub else name + "\n" + sub, size=14)


# ------------------------------------------------------------------ 1. map
s = d.slide().s
title(s, "From the PGAS view down to the BAR1 window",
      "Blue: PGAS view.  Pink: goes through BAR1.  Yellow: NIC.  Grey: NVLink (no BAR1).  Green: GPU memory")
box(s, 43, 120, 1194, 70, LINK_DK, stroke=LINK, body="PGAS view: every GPU's symmetric heap in one address book (NVSHMEM)", size=17)
box(s, 43, 230, 380, 120, CHIP, stroke=EDGE, body="same node, NVLink\nmap peer heap, plain store\n(no BAR1)", size=15)
box(s, 450, 230, 380, 120, CHIP, stroke=PNK, body="same node, PCIe P2P\nmap peer heap, plain store\nthrough the peer's BAR1", size=15)
box(s, 857, 230, 380, 120, CHIP, stroke=YEL, body="other node\nNIC carries the data\nNIC DMA into BAR1", size=15)
for cx in (233, 640, 1047):
    arrow(s, cx, 190, cx, 230)
box(s, 450, 410, 787, 70, PNK_DK, stroke=PNK, body="BAR1: GPU memory seen as PCIe addresses", size=17)
arrow(s, 640, 350, 640, 410, color=PNK, width=3)
arrow(s, 1047, 350, 1047, 410, color=PNK, width=3)
box(s, 43, 520, 1194, 70, GRN_DK, stroke=GRN, body="GPU memory (VRAM)", size=17)
arrow(s, 233, 350, 233, 520, color=MUTED, width=3)
arrow(s, 840, 480, 840, 520, color=PNK, width=3)

# ------------------------------------------------------------------ 2. PGAS
s = d.slide().s
title(s, "PGAS: one address book over every GPU's memory",
      "Green: GPU memory.  Yellow: the symmetric heap.  Pink: a and b, allocated by every PE together")
for i in range(4):
    x = 43 + i * 300
    box(s, x, 130, 280, 420, GRN_DK, stroke=GRN, radius=0.04)
    label(s, x, 140, 280, 30, f"PE {i} (GPU {i})", size=17, bold=True)
    box(s, x + 20, 190, 240, 330, YEL_DK, stroke=YEL, radius=0.04)
    label(s, x + 20, 196, 240, 26, "symmetric heap", size=14)
    box(s, x + 40, 260, 200, 60, PNK_DK, stroke=PNK, body="a (1 MiB)", size=15)
    box(s, x + 40, 340, 200, 60, PNK_DK, stroke=PNK, body="b (1 MiB)", size=15)
label(s, 43, 570, 1194, 30, "a on PE 2 = name + PE number", size=15, color=MUTED, align="left")

# ------------------------------------------------------------------ 3. offset
s = d.slide().s
title(s, "\"Symmetric\" means the same offset from the heap base",
      "Measured on g7.48xlarge, NVSHMEM 3.8.0.  Yellow: heap base.  Pink: a")
label(s, 43, 112, 584, 30, "default (CUDA VMM)", size=18, bold=True)
label(s, 653, 112, 584, 30, "CUDA VMM disabled", size=18, bold=True)
for k, (base_l, a_l) in enumerate([("0x10050000000", "0x10050003600"), ("0x10050000000", "0x10050003600")]):
    y = 170 + k * 170
    label(s, 63, y, 200, 30, f"PE {k}", size=16, bold=True, align="left")
    box(s, 63, y + 36, 540, 100, CHIP, stroke=EDGE, radius=0.04)
    box(s, 83, y + 56, 200, 60, YEL_DK, stroke=YEL, body="heap base", size=14)
    box(s, 383, y + 56, 200, 60, PNK_DK, stroke=PNK, body="a = " + a_l, size=14)
for k, a_l in enumerate(["0x7f8bc0003600", "0x7ff060003600"]):
    y = 170 + k * 170
    label(s, 673, y, 200, 30, f"PE {k}", size=16, bold=True, align="left")
    box(s, 673, y + 36, 540, 100, CHIP, stroke=EDGE, radius=0.04)
    box(s, 693, y + 56, 200, 60, YEL_DK, stroke=YEL, body="heap base", size=14)
    box(s, 993, y + 56, 200, 60, PNK_DK, stroke=PNK, body="a = " + a_l, size=14)
label(s, 43, 520, 584, 30, "same value on every PE (by chance)", size=15, color=MUTED)
label(s, 653, 520, 584, 30, "different values, same tail ...3600", size=15, color=MUTED)

# ------------------------------------------------------------------ 4. P2P matrix
s = d.slide().s
title(s, "Can one GPU write another GPU's memory directly over PCIe?",
      "Measured with nvidia-smi topo -p2p rw.  Green: OK.  Grey: not supported")
for col, (name, n, ok) in enumerate([("g6e.12xlarge: 4 × L40S", 4, False), ("g7.48xlarge: 8 × RTX PRO 4500", 8, True)]):
    x0 = 43 + col * 610
    label(s, x0, 112, 590, 30, name, size=18, bold=True)
    cell = 60 if n == 4 else 46
    for i in range(n):
        for j in range(n):
            x = x0 + 40 + j * cell
            y = 170 + i * cell
            if i == j:
                box(s, x, y, cell - 6, cell - 6, CHIP, stroke=EDGE, body="-", size=12)
            else:
                box(s, x, y, cell - 6, cell - 6, GRN_DK if ok else GREY, stroke=GRN if ok else EDGE, body="OK" if ok else "NS", size=12)
label(s, 43, 560, 590, 30, "NVSHMEM with 4 PEs: \"Peer GPU is not accessible\"", size=15, color=MUTED)
label(s, 653, 560, 590, 30, "NVSHMEM with 4 PEs: maps every peer heap", size=15, color=MUTED)

# ------------------------------------------------------------------ 5. BAR window
s = d.slide().s
title(s, "BAR1: the window that shows GPU memory on PCIe",
      "Measured on g6e.12xlarge (L40S).  Pink: BAR1.  Green: GPU memory.  Grey: other BARs.  Yellow: other PCIe devices")
box(s, 43, 130, 1194, 90, CHIP, stroke=EDGE, radius=0.04)
label(s, 63, 135, 500, 30, "PCIe address space", size=15, bold=True, align="left")
box(s, 83, 170, 200, 40, GREY, stroke=EDGE, body="BAR0 16 MiB: registers", size=13)
box(s, 303, 170, 640, 40, PNK_DK, stroke=PNK, body="BAR1 64 GiB: window onto GPU memory", size=13)
box(s, 963, 170, 200, 40, GREY, stroke=EDGE, body="BAR3 32 MiB", size=13)
box(s, 303, 330, 470, 110, GRN_DK, stroke=GRN, body="GPU memory 45 GiB", size=17)
arrow(s, 540, 210, 540, 330, color=PNK, width=3, dashed=True)
box(s, 873, 330, 300, 110, YEL_DK, stroke=YEL, body="CPU, NIC, other GPUs\nread and write here", size=15)
arrow(s, 1023, 330, 900, 210, color=WHITE, width=2)

# ------------------------------------------------------------------ 6. NIC path
s = d.slide().s
title(s, "How the NIC reaches GPU memory",
      "Green: GPU memory.  Pink: BAR1.  Yellow: NIC.  Grey: setup before transfers")
box(s, 43, 150, 340, 90, GREY, stroke=EDGE, body="1. register the heap\nwith the NIC", size=15)
box(s, 443, 150, 340, 90, GREY, stroke=EDGE, body="2. driver gives\nBAR1 addresses", size=15)
box(s, 843, 150, 340, 90, GREY, stroke=EDGE, body="3. NIC remembers\nthose addresses", size=15)
arrow(s, 383, 195, 443, 195, color=MUTED)
arrow(s, 783, 195, 843, 195, color=MUTED)
box(s, 843, 340, 340, 90, YEL_DK, stroke=YEL, body="NIC", size=17)
box(s, 443, 340, 340, 90, PNK_DK, stroke=PNK, body="BAR1 address", size=17)
box(s, 43, 340, 340, 90, GRN_DK, stroke=GRN, body="GPU memory", size=17)
arrow(s, 843, 385, 783, 385, color=YEL, width=3)
arrow(s, 443, 385, 383, 385, color=PNK, width=3)
label(s, 43, 460, 1194, 30, "per transfer: NIC DMA to the BAR1 address", size=15, color=MUTED, align="left")

# ------------------------------------------------------------------ 7. static vs dynamic
s = d.slide().s
title(s, "Small window vs large window",
      "Pink: mapped part of BAR1.  Green: GPU memory")
label(s, 43, 112, 584, 30, "BAR1 smaller than memory (Kepler: min 256 MB)", size=18, bold=True)
label(s, 653, 112, 584, 30, "BAR1 holds all memory + reserve", size=18, bold=True)
box(s, 63, 170, 200, 60, CHIP, stroke=EDGE, body="", radius=0.04)
box(s, 63, 170, 90, 60, PNK_DK, stroke=PNK, body="mapped", size=13)
box(s, 63, 320, 544, 90, GRN_DK, stroke=GRN, body="GPU memory", size=16)
arrow(s, 108, 230, 108, 320, color=PNK, width=3, dashed=True)
label(s, 63, 430, 544, 60, "map only what is registered\nBAR1 Used grows, can run out", size=15, color=MUTED)
box(s, 673, 170, 544, 60, PNK_DK, stroke=PNK, body="mapped once at boot: all of memory", size=13)
box(s, 673, 320, 544, 90, GRN_DK, stroke=GRN, body="GPU memory", size=16)
for cx in (733, 945, 1157):
    arrow(s, cx, 230, cx, 320, color=PNK, width=3, dashed=True)
label(s, 673, 430, 544, 60, "static BAR1 (when enabled): registration reuses it\nobserved on g7.48xlarge: BAR1 Used 0-3 MiB", size=15, color=MUTED)
d.save(OUT)
# ZN-24: テンプレートのロゴ、著作権表示、ヘッダー、フッターを必ず消してから使う
sys.path.insert(0, "/Users/akazawt/works/data-science/claudecode/.claude/skills/review-html/packs/zenn-doc/templates")
import strip_branding  # noqa: E402
strip_branding.main(OUT, OUT)
print(OUT)
