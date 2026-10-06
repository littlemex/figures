"""gen_compile_pptx.py -- PyTorch の forward が NEFF になるまでの図を、ネイティブ shape の pptx 4 枚で組む。

    /usr/bin/python3 gen_compile_pptx.py <out.pptx>

1. 全体: forward -> FX graph -> HLO -> NEFF と、変換を受け持つ道具。CPU で済む範囲と NeuronCore が要る範囲
2. neuronx-cc の中: ログに出る 5 つの段の名前と、hlo2penguin と MLIR のパス
3. CPU コンパイルが残すファイル: キャッシュのディレクトリと、どの段の成果物か
4. 共通の IR を挟む理由: N x M の変換と N + M の変換
根拠: vLLM Neuron v0.21.0.1.0.0 (Neuron SDK 2.31.0, neuronx-cc 2.26) を CPU で動かしたときのファイルとログ。
図の文字は英語。色の役割は ZN-25 (緑 = 自分が書く、青 = 道具が受け持つ、黄 = 注意)。
"""
import sys

sys.path.insert(0, "/Users/akazawt/works/data-science/tools/pptxkit")
from pptx.dml.color import RGBColor  # noqa: E402
from pptx.enum.shapes import MSO_SHAPE  # noqa: E402

from pptxkit import Deck, shapes as sh  # noqa: E402

TPL = "/Users/akazawt/works/data-science/tools/pptxkit/templates/aws-dark-2024.pptx"
OUT = sys.argv[1] if len(sys.argv) > 1 else "compile-figures.pptx"

GRN = RGBColor(36, 160, 70)
GRN_HI = RGBColor(59, 209, 111)
GRN_DK = RGBColor(18, 64, 34)
CHIP = RGBColor(22, 26, 30)
EDGE = RGBColor(70, 74, 80)
GREY = RGBColor(52, 56, 64)
LINK = RGBColor(40, 150, 240)
LINK_DK = RGBColor(16, 44, 72)
YEL = RGBColor(251, 211, 50)
YEL_DK = RGBColor(80, 66, 14)
MUTED = RGBColor(150, 150, 158)
WHITE = RGBColor(245, 245, 248)

d = Deck(TPL, check_layout=False)
t = d.theme


def label(s, x, y, w, h, body, size=16, color=WHITE, bold=False, align="center"):
    return sh.textbox(s, t, x, y, w, h, body, size=size, color=color, bold=bold, align=align,
                      anchor="middle")


def box(s, x, y, w, h, fill, stroke=None, body="", size=15, color=WHITE, radius=0.1,
        dash=None, sw=1.25):
    sp = sh.rect(s, t, x, y, w, h, fill=fill, stroke=stroke, sw=sw, dash=dash,
                 shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=radius)
    if body:
        sh._write(sp.text_frame, t, body, size, color, True, "center", "middle")
        tf = sp.text_frame
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = sh.U(2)
    return sp


def title(s, head, sub):
    label(s, 43, 26, 1190, 46, head, size=30, bold=True, align="left")
    label(s, 43, 72, 1190, 26, sub, size=16, color=MUTED, align="left")


def bracket(s, x0, x1, y, text, color):
    sh.arrow(s, t, x0, y, x1, y, color=color, width=2)
    sh.arrow(s, t, x1, y, x0, y, color=color, width=2)
    label(s, x0, y + 8, x1 - x0, 30, text, size=15, color=color)


# ------------------------------------------------------------------ 1. the whole pipeline
s = d.slide()
title(s.s, "From forward to a Trainium binary",
      "Green: you write it.  Blue: the toolchain.  Yellow: needs a NeuronCore")
stages = [("forward", "Python", GRN_DK, GRN), ("FX graph", "fxgraph.txt", LINK_DK, LINK),
          ("HLO", "graph.hlo", LINK_DK, LINK), ("NEFF", "graph_<hash>.neff", LINK_DK, LINK)]
tools = ["TorchDynamo", "torch-xla", "neuronx-cc"]
W, GAP, X0, Y = 170, 80, 43, 190
for i, (name, file, fill, edge) in enumerate(stages):
    x = X0 + i * (W + GAP)
    box(s.s, x, Y, W, 90, fill, stroke=edge, body=name, size=22)
    label(s.s, x, Y + 96, W, 26, file, size=14, color=MUTED)
    if i < 3:
        sh.arrow(s.s, t, x + W + 4, Y + 45, x + W + GAP - 4, Y + 45, color=WHITE, width=2.5)
        label(s.s, x + W - 20, Y - 40, GAP + 40, 30, tools[i], size=14, color=LINK)
xr = X0 + 4 * (W + GAP) - GAP + 40
box(s.s, xr, Y, 1237 - xr, 90, YEL_DK, stroke=YEL, body="Neuron Runtime\non a NeuronCore", size=17)
sh.arrow(s.s, t, xr - 36, Y + 45, xr - 4, Y + 45, color=YEL, width=2.5)
bracket(s.s, X0, X0 + 4 * W + 3 * GAP, 360, "runs on a CPU", WHITE)
bracket(s.s, xr, 1237, 360, "needs Trainium", YEL)
box(s.s, X0 + 3 * (W + GAP), 450, W, 70, GREY, stroke=EDGE, body="compile cache", size=16)
sh.arrow(s.s, t, X0 + 3 * (W + GAP) + W / 2, Y + 124, X0 + 3 * (W + GAP) + W / 2, 446, color=MUTED, width=2)
label(s.s, X0 + 3 * (W + GAP) - 60, 524, W + 120, 26, "cache hit: reuse", size=14, color=MUTED)

# ------------------------------------------------------------------ 2. inside neuronx-cc
s = d.slide()
title(s.s, "Inside neuronx-cc, as its log shows",
      "Blue: stages named in log-neuron-cc.txt.  Grey: what the log shows it runs")
names = ["HLOToTensorizer", "Frontend", "StaticIOTranspose", "WalrusDriver", "NeffWrapper"]
W, GAP, X0, Y = 180, 22, 132, 240
box(s.s, 22, Y, 90, 80, CHIP, stroke=EDGE, body="graph\n.hlo", size=15)
for i, nm in enumerate(names):
    x = X0 + i * (W + GAP)
    box(s.s, x, Y, W, 80, LINK_DK, stroke=LINK, body=nm, size=14)
    if i < 4:
        sh.arrow(s.s, t, x + W + 2, Y + 40, x + W + GAP - 2, Y + 40, color=WHITE, width=2)
sh.arrow(s.s, t, 112, Y + 40, X0 - 2, Y + 40, color=WHITE, width=2)
box(s.s, 1147, Y, 90, 80, CHIP, stroke=EDGE, body="graph\n.neff", size=15)
sh.arrow(s.s, t, X0 + 5 * (W + GAP) - GAP + 2, Y + 40, 1145, Y + 40, color=WHITE, width=2)
box(s.s, X0, Y + 110, W, 110, GREY, stroke=EDGE, body="hlo2penguin\nHLO -> penguin.py\nMLIR passes", size=14)
sh.arrow(s.s, t, X0 + W / 2, Y + 82, X0 + W / 2, Y + 108, color=MUTED, width=2)
x3 = X0 + 3 * (W + GAP)
x4 = X0 + 4 * (W + GAP)

# ------------------------------------------------------------------ 3. files left by a CPU compile
s = d.slide()
title(s.s, "What a CPU compile leaves on disk",
      "Files under VLLM_CACHE_ROOT/.../compile_cache/<hash>/")
rows = [("FX graph", "fxgraph.txt", "passes/00_original.txt ... 05_*.txt", LINK_DK, LINK),
        ("HLO", "graph.hlo", "hlo_passes/step1_torch_xla_trace.hlo ... step4_*.hlo", LINK_DK, LINK),
        ("neuronx-cc", "command.txt", "log-neuron-cc.txt", LINK_DK, LINK),
        ("NEFF", "graph_<hash>.neff", "", LINK_DK, LINK)]
y = 130
for name, main, extra, fill, edge in rows:
    box(s.s, 43, y, 230, 80, fill, stroke=edge, body=name, size=19)
    sh.arrow(s.s, t, 277, y + 40, 330, y + 40, color=MUTED, width=2)
    box(s.s, 334, y + 10, 330, 60, CHIP, stroke=EDGE, body=main, size=17)
    if extra:
        label(s.s, 690, y + 10, 547, 60, extra, size=15, color=MUTED, align="left")
    y += 104
box(s.s, 43, 560, 1194, 48, YEL_DK, stroke=YEL, radius=0.3,
    body="Observed next: NEFF load error", size=16)

# ------------------------------------------------------------------ 4. why a shared IR
s = d.slide()
title(s.s, "Why a shared IR in the middle", "Counting model.  Yellow: the shared layer in the middle")
fw = ["Framework 1", "Framework 2", "Framework 3"]
hw = ["Hardware 1", "Hardware 2", "Hardware 3"]
for side, x0 in ((0, 43), (1, 665)):
    box(s.s, x0, 112, 572, 470, CHIP, stroke=EDGE, radius=0.04)
    for i in range(3):
        box(s.s, x0 + 30, 160 + i * 130, 150, 64, GRN_DK, stroke=GRN, body=fw[i], size=16)
        box(s.s, x0 + 392, 160 + i * 130, 150, 64, LINK_DK, stroke=LINK, body=hw[i], size=16)
for i in range(3):
    for j in range(3):
        sh.arrow(s.s, t, 223, 192 + i * 130, 435, 192 + j * 130, color=MUTED, width=1.5)
box(s.s, 665 + 211, 302, 150, 80, YEL_DK, stroke=YEL, body="shared IR", size=18)
for i in range(3):
    sh.arrow(s.s, t, 665 + 180, 192 + i * 130, 665 + 211, 342, color=MUTED, width=1.5)
    sh.arrow(s.s, t, 665 + 361, 342, 665 + 392, 192 + i * 130, color=MUTED, width=1.5)
label(s.s, 43, 590, 572, 34, "3 × 3 = 9 translators", size=18, bold=True)
label(s.s, 665, 590, 572, 34, "3 + 3 = 6 translators", size=18, bold=True)

d.save(OUT)
# ZN-24: テンプレートのロゴ、著作権表示、ヘッダー、フッターを必ず消してから使う
sys.path.insert(0, "/Users/akazawt/works/data-science/claudecode/.claude/skills/review-html/packs/zenn-doc/templates")
import strip_branding  # noqa: E402
strip_branding.main(OUT, OUT)
print(OUT)
