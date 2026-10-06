"""gen_lustre_pptx.py -- FSx for Lustre クライアントの DKMS 化の記事の図を、ネイティブ shape の pptx で組む。

    /usr/bin/python3 gen_lustre_pptx.py <out.pptx>

1. モジュールがカーネルのリリースに縛られる理由: vermagic とリリースごとのパッケージ
2. 5 段の入れ子: Ubuntu のリリース -> suite -> カーネル系列 -> カーネルのリリース -> モジュールのパッケージ
3. モジュールをどこでビルドするか: AWS (公開パッケージ) / AMI を焼くとき / ホストでカーネルが入るとき (DKMS)
4. ビルドが走る瞬間: apt -> dpkg -> postinst.d と header_postinst.d のフック -> dkms autoinstall
5. 失敗が apt を巻き込む: run-parts 11 -> dpkg error -> half-configured -> apt 100
6. 対象外のカーネルの扱い 3 通り: 設定なし / BUILD_EXCLUSIVE_KERNEL / no-autoinstall-errors
7. 公開状況: suite ごとのカーネル系列が観測できた最古の更新日 (2026-09-12 取得)
8. スクリプトの処理の流れ
9. --refresh-policy とカーネルの導入の順序
10. AMI の 2 段構え
根拠: 記事の本文と実行例、dkms(8) (Ubuntu 24.04)、lustre_installer.sh (littlemex/distributed-ai 52fff76)。
図の文字は英語。色の役割は ZN-25 (緑 = 自分が書く・置く、青 = OS とツールが受け持つ、黄 = 注意)。
"""
import sys

sys.path.insert(0, "/Users/akazawt/works/data-science/tools/pptxkit")
from pptx.dml.color import RGBColor  # noqa: E402
from pptx.enum.dml import MSO_LINE_DASH_STYLE  # noqa: E402
from pptx.enum.shapes import MSO_SHAPE  # noqa: E402

from pptxkit import Deck, shapes as sh  # noqa: E402

TPL = "/Users/akazawt/works/data-science/tools/pptxkit/templates/aws-dark-2024.pptx"
OUT = sys.argv[1] if len(sys.argv) > 1 else "lustre-figures.pptx"

GRN = RGBColor(36, 160, 70)
GRN_HI = RGBColor(59, 209, 111)
GRN_DK = RGBColor(18, 64, 34)
CHIP = RGBColor(22, 26, 30)
EDGE = RGBColor(70, 74, 80)
GREY = RGBColor(52, 56, 64)
LINK = RGBColor(40, 150, 240)
LINK_DK = RGBColor(16, 44, 72)
YEL = RGBColor(251, 211, 50)
YEL_DK = RGBColor(80, 66, 14)
RED = RGBColor(230, 90, 90)
MUTED = RGBColor(150, 150, 158)
WHITE = RGBColor(245, 245, 248)

d = Deck(TPL, check_layout=False)
t = d.theme


def label(s, x, y, w, h, body, size=16, color=WHITE, bold=False, align="center"):
    return sh.textbox(s, t, x, y, w, h, body, size=size, color=color, bold=bold, align=align,
                      anchor="middle")


def box(s, x, y, w, h, fill, stroke=None, body="", size=15, color=WHITE, radius=0.1,
        dash=None, sw=1.25):
    sp = sh.rect(s, t, x, y, w, h, fill=fill, stroke=stroke, sw=sw, dash=dash,
                 shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=radius)
    if body:
        sh._write(sp.text_frame, t, body, size, color, True, "center", "middle")
        tf = sp.text_frame
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = sh.U(2)
    return sp


def title(s, head, sub):
    label(s, 43, 26, 1190, 46, head, size=30, bold=True, align="left")
    label(s, 43, 72, 1190, 26, sub, size=16, color=MUTED, align="left")


def arr(s, x0, y0, x1, y1, color=WHITE, width=2):
    sh.arrow(s, t, x0, y0, x1, y1, color=color, width=width)


# ------------------------------------------------------------------ 1. tied to a kernel release
s = d.slide()
title(s.s, "A module is built for one kernel release",
      "Blue: the running kernel.  Green: the module.  Yellow: what goes wrong")
cases = [(43, "6.8.0-1057-aws", "vermagic 6.8.0-1057-aws ...", True),
         (665, "7.0.0-1012-aws", "vermagic 6.8.0-1057-aws ...  (example mismatch)", False)]
for x, kern, vm, ok in cases:
    box(s.s, x, 120, 572, 70, LINK_DK, stroke=LINK, body=f"running kernel  {kern}", size=17)
    box(s.s, x + 20, 230, 532, 70, GRN_DK, stroke=GRN, body=f"lustre.ko   {vm}", size=14)
    arr(s.s, x + 286, 228, x + 286, 194, WHITE if ok else YEL, 2.5)
    box(s.s, x + 160, 318, 252, 50, (LINK_DK if ok else YEL_DK), stroke=(LINK if ok else YEL),
        body=("loaded" if ok else "load refused"), size=16)
label(s.s, 43, 400, 1194, 30, "packages in the repository: one per kernel release", size=17, bold=True, align="left")
pk = ["...-6.8.0-1024-aws", "...", "...-6.8.0-1055-aws", "...-6.8.0-1057-aws", "...-7.0.0-1012-aws"]
x = 43
for p in pk:
    box(s.s, x, 440, 220, 52, CHIP, stroke=EDGE, body=p, size=13)
    x += 236
box(s.s, 43, 520, 1194, 56, YEL_DK, stroke=YEL, radius=0.3,
    body="not in the list: nothing to install", size=16)

# ------------------------------------------------------------------ 2. five nested levels
s = d.slide()
title(s.s, "Five levels, each one picks the next", "Left: the level.  Right: who decides it")
lv = [("Ubuntu release", "24.04 LTS", "you choose"),
      ("suite in the FSx repository", "dists/noble", "the codename"),
      ("kernel series", "6.8 / 6.14 / 6.17 / 7.0", "Ubuntu's kernel policy"),
      ("kernel release", "6.8.0-1057-aws", "the kernel update you applied"),
      ("module package", "lustre-client-modules-6.8.0-1057-aws", "whether AWS published it")]
y = 112
for i, (nm, ex, who) in enumerate(lv):
    box(s.s, 43 + i * 40, y, 640, 74, LINK_DK if i < 4 else GRN_DK, stroke=LINK if i < 4 else GRN,
        body=f"{nm}\n{ex}", size=15)
    label(s.s, 43 + i * 40 + 660, y, 1237 - (43 + i * 40 + 660), 74, who, size=15, color=MUTED, align="left")
    if i < 4:
        arr(s.s, 43 + i * 40 + 60, y + 76, 43 + (i + 1) * 40 + 60, y + 96)
    y += 98

# ------------------------------------------------------------------ 3. where the module is built
s = d.slide()
title(s.s, "Where the module gets built", "Green: you do it.  Blue: AWS or the OS does it.  Yellow: the cost")
cols = [("Pinned kernel + published module", "AWS builds it", LINK_DK, LINK,
         "apt-mark hold the kernel;\nraise it after AWS publishes", "kernel updates wait"),
        ("Baked into a custom AMI", "you build it once,\nwhen the image is made", GRN_DK, GRN,
         "replace nodes to update", "build tools stay\nout of the nodes"),
        ("DKMS on the host", "the host builds it\nwhen a kernel is installed", LINK_DK, LINK,
         "follows kernel updates", "build runs inside\nthe kernel install")]
for i, (head, who, fill, edge, how, cost) in enumerate(cols):
    x = 43 + i * 405
    label(s.s, x, 112, 380, 36, head, size=16, bold=True)
    box(s.s, x, 160, 380, 100, fill, stroke=edge, body=who, size=16)
    label(s.s, x, 272, 380, 60, how, size=15, color=WHITE)
    box(s.s, x, 350, 380, 80, YEL_DK if i == 2 else CHIP, stroke=YEL if i == 2 else EDGE, body=cost, size=15)

# ------------------------------------------------------------------ 4. when the build runs
s = d.slide()
title(s.s, "When the DKMS build runs", "Blue: packages and hooks.  Green: what you registered.  Yellow: the decision")
box(s.s, 43, 130, 170, 70, LINK_DK, stroke=LINK, body="apt install\nlinux-aws", size=15)
box(s.s, 260, 130, 150, 70, LINK_DK, stroke=LINK, body="dpkg", size=17)
arr(s.s, 213, 165, 256, 165)
box(s.s, 460, 112, 380, 54, LINK_DK, stroke=LINK, body="linux-image  ->  /etc/kernel/postinst.d/dkms", size=13)
box(s.s, 460, 180, 380, 54, LINK_DK, stroke=LINK, body="linux-headers  ->  /etc/kernel/header_postinst.d/dkms", size=13)
arr(s.s, 410, 150, 456, 139)
arr(s.s, 410, 180, 456, 207)
box(s.s, 900, 140, 220, 60, LINK_DK, stroke=LINK, body="dkms autoinstall", size=16)
arr(s.s, 840, 139, 896, 160)
arr(s.s, 840, 207, 896, 180)
box(s.s, 880, 250, 260, 60, YEL_DK, stroke=YEL, body="headers for this kernel?", size=15)
arr(s.s, 1010, 202, 1010, 246)
box(s.s, 600, 350, 260, 60, CHIP, stroke=EDGE, body="no: skip, log only", size=15)
box(s.s, 940, 350, 297, 60, GRN_DK, stroke=GRN, body="yes: build registered source", size=15)
arr(s.s, 930, 312, 760, 346, MUTED)
arr(s.s, 1060, 312, 1080, 346)
box(s.s, 940, 450, 297, 60, LINK_DK, stroke=LINK, body="result back to dpkg and apt", size=14)
arr(s.s, 1088, 412, 1088, 446)


# ------------------------------------------------------------------ 5. failure spreads to apt
s = d.slide()
title(s.s, "A failed build fails the kernel package", "Yellow: where the failure travels")
st = [("DKMS build fails", "source not ready for 7.0"), ("run-parts", "postinst.d/dkms exited 11"),
      ("dpkg", "error processing\nlinux-image-7.0.0-1012-aws"), ("apt", "exit code 100")]
x = 43
for i, (nm, sub) in enumerate(st):
    box(s.s, x, 140, 260, 70, YEL_DK, stroke=YEL, body=nm, size=17)
    label(s.s, x, 216, 260, 56, sub, size=14, color=MUTED)
    if i < 3:
        arr(s.s, x + 262, 175, x + 304, 175, YEL, 2.5)
    x += 306
box(s.s, 43, 320, 560, 80, CHIP, stroke=EDGE, body="package left half-configured\n(dpkg --audit lists it)", size=16)
box(s.s, 677, 320, 560, 80, CHIP, stroke=EDGE, body="every later apt run\nretries it and fails again", size=16)
arr(s.s, 603, 360, 673, 360)
box(s.s, 43, 440, 1194, 70, GRN_DK, stroke=GRN, radius=0.2,
    body="fix the cause or  dkms remove  the registration,  then  dpkg --configure -a", size=16)

# ------------------------------------------------------------------ 6. three ways to treat a kernel you cannot build for
s = d.slide()
title(s.s, "A kernel the source cannot build for", "Three settings, three outcomes (dkms 3.0.11)")
rows = [("no setting", "build runs and fails", "kernel package fails; apt stuck", YEL_DK, YEL),
        ("BUILD_EXCLUSIVE_KERNEL='^(6\\.8|...)\\.'", "no match: skipped (exit 77, ignored by autoinstall)",
         "apt fine; that kernel has no module", LINK_DK, LINK),
        ("/etc/dkms/no-autoinstall-errors", "build fails, reported as success", "apt fine; failure invisible", CHIP, EDGE)]
label(s.s, 43, 112, 380, 30, "setting", size=15, color=MUTED, align="left")
label(s.s, 443, 112, 400, 30, "what DKMS does", size=15, color=MUTED, align="left")
label(s.s, 863, 112, 374, 30, "what you get", size=15, color=MUTED, align="left")
y = 150
for nm, act, res, fill, edge in rows:
    box(s.s, 43, y, 380, 90, fill, stroke=edge, body=nm, size=14)
    box(s.s, 443, y, 400, 90, CHIP, stroke=EDGE, body=act, size=14)
    box(s.s, 863, y, 374, 90, CHIP, stroke=EDGE, body=res, size=14)
    y += 110
label(s.s, 43, 490, 1194, 30, "this article uses the second setting", size=15, color=MUTED, align="left")

# ------------------------------------------------------------------ 7. published series
s = d.slide()
title(s.s, "Oldest observed package per series", "Last-Modified seen in the list fetched 2026-09-12.  Bar: from that date to the fetch")
X0, X1 = 230, 1210
T0, T1 = 2023.0, 2026.75


def xt(y):
    return X0 + (X1 - X0) * (y - T0) / (T1 - T0)


for yr in (2023, 2024, 2025, 2026):
    sh.arrow(s.s, t, xt(yr), 120, xt(yr), 125, color=EDGE, width=1)
    label(s.s, xt(yr) - 40, 104, 80, 22, str(yr), size=13, color=MUTED)
ser = [("jammy 5.15", 2023 + 1.5 / 12, 27), ("jammy 6.2", 2023 + 11.5 / 12, 3), ("jammy 6.5", 2024 + 6.9 / 12, 4),
       ("jammy 6.8", 2024 + 10 / 12, 54), ("noble 6.8", 2025 + 2.6 / 12, 11), ("noble 6.14", 2025 + 8.7 / 12, 10),
       ("noble 6.17", 2026 + 1.75 / 12, 14), ("noble 7.0", 2026 + 8.3 / 12, 1)]
fetch = 2026 + 8.4 / 12
y = 140
for nm, t0, n in ser:
    label(s.s, 43, y, 180, 40, nm, size=15, align="left")
    noble = nm.startswith("noble")
    box(s.s, xt(t0), y + 6, max(10, xt(fetch) - xt(t0)), 28, LINK_DK if noble else CHIP,
        stroke=LINK if noble else EDGE, radius=0.3)
    label(s.s, xt(t0) - 80, y, 76, 40, f"{n}", size=13, color=MUTED, align="right")
    y += 48
label(s.s, 43, 538, 1194, 30, "number left of a bar: packages in the list.   not in the list: jammy 5.19, noble 6.11",
      size=14, color=MUTED, align="left")

# ------------------------------------------------------------------ 8. script flow
s = d.slide()
title(s.s, "What lustre_installer.sh does", "Blue: checks and setup.  Green: the mode you chose.  Yellow: stops with a hint")
flow = [("check arguments", 0), ("Ubuntu?", 1), ("take the lock", 0), ("suite exists?", 1),
        ("verify key fingerprint", 0), ("add the repository", 0), ("install lustre-client-utils", 0)]
y = 110
for i, (nm, q) in enumerate(flow):
    box(s.s, 43, y, 300, 50, LINK_DK, stroke=LINK, body=nm, size=14)
    if q:
        msg = "no: point to the Amazon Linux 2023 steps" if i == 1 else "no: point to an existing suite"
        box(s.s, 380, y, 400, 50, YEL_DK, stroke=YEL, body=msg, size=14)
        arr(s.s, 343, y + 25, 376, y + 25, YEL)
    if i < len(flow) - 1:
        arr(s.s, 193, y + 50, 193, y + 64)
    y += 64
modes = [("build", "build the source, install a deb"), ("dkms", "derive series, register with DKMS"),
         ("binary", "install the published module")]
for i, (m, w) in enumerate(modes):
    yy = 300 + i * 80
    box(s.s, 840, yy, 397, 64, GRN_DK, stroke=GRN, body=f"{m}:  {w}", size=14)
    arr(s.s, 343, y - 39, 836, yy + 32, MUTED, 1.5)
    arr(s.s, 1038, yy + 64, 1038, 524, MUTED, 1.5)
sh.rect(s.s, t, 1033, 524, 10, 10, fill=MUTED, stroke=None, shape=MSO_SHAPE.OVAL)
arr(s.s, 1038, 529, 1038, 558, MUTED, 1.5)
box(s.s, 840, 560, 190, 50, LINK_DK, stroke=LINK, body="verify", size=15)
box(s.s, 1047, 560, 190, 50, LINK_DK, stroke=LINK, body="print a summary", size=15)
arr(s.s, 1030, 585, 1043, 585)
arr(s.s, 1030, 585, 1043, 585)

# ------------------------------------------------------------------ 9. refresh-policy order
s = d.slide()
title(s.s, "Widening the policy only helps kernels installed later",
      "Green: built automatically.  Yellow: reported by name; build it with --kernel")
for i, (first, second, ok) in enumerate(((" --refresh-policy  adds 7.0", "7.0 kernel installed", True),
                                         ("7.0 kernel installed (skipped)", " --refresh-policy  adds 7.0", False))):
    y = 140 + i * 200
    box(s.s, 43, y, 380, 70, LINK_DK, stroke=LINK, body=f"1  {first.strip()}", size=15)
    box(s.s, 470, y, 380, 70, LINK_DK, stroke=LINK, body=f"2  {second.strip()}", size=15)
    arr(s.s, 423, y + 35, 466, y + 35)
    box(s.s, 900, y, 337, 70, GRN_DK if ok else YEL_DK, stroke=GRN if ok else YEL,
        body=("module built at install" if ok else "no rebuild; reported"), size=15)
    arr(s.s, 850, y + 35, 896, y + 35, WHITE if ok else YEL)

# ------------------------------------------------------------------ 10. AMI in two stages
s = d.slide()
title(s.s, "Baking an AMI in two stages", "Packer runs Ansible twice.  Green: stages.  Yellow: the choice")
box(s.s, 43, 220, 180, 80, CHIP, stroke=EDGE, body="parent AMI", size=17)
box(s.s, 270, 210, 230, 100, YEL_DK, stroke=YEL, body="pin the kernel\nseries?", size=16)
arr(s.s, 223, 260, 266, 260)
box(s.s, 560, 110, 300, 90, GRN_DK, stroke=GRN, body="stage 1\npin the series, reboot", size=15)
box(s.s, 560, 320, 300, 90, GRN_DK, stroke=GRN, body="stage 2\nbuild for the running kernel", size=15)
arr(s.s, 500, 240, 556, 160)
label(s.s, 480, 160, 80, 26, "yes", size=14, color=MUTED)
arr(s.s, 500, 280, 556, 360)
label(s.s, 480, 330, 80, 26, "no", size=14, color=MUTED)
arr(s.s, 710, 202, 710, 316)
box(s.s, 940, 320, 200, 90, LINK_DK, stroke=LINK, body="AMI", size=18)
arr(s.s, 860, 365, 936, 365)
label(s.s, 43, 470, 1194, 30, "reboot: proves the module loads on the kernel the image will boot", size=15, color=MUTED, align="left")

d.save(OUT)
# ZN-24: テンプレートのロゴ、著作権表示、ヘッダー、フッターを必ず消してから使う
sys.path.insert(0, "/Users/akazawt/works/data-science/claudecode/.claude/skills/review-html/packs/zenn-doc/templates")
import strip_branding  # noqa: E402
strip_branding.main(OUT, OUT)
print(OUT)
