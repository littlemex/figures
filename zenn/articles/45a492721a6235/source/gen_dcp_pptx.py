"""gen_dcp_pptx.py: DCP article static figures."""
import os
import sys

sys.path.insert(0, "/Users/akazawt/works/data-science/tools/pptxkit")
from pptx.dml.color import RGBColor  # noqa: E402
from pptx.enum.dml import MSO_LINE_DASH_STYLE  # noqa: E402
from pptx.enum.shapes import MSO_SHAPE  # noqa: E402

from pptxkit import Deck, shapes as sh  # noqa: E402

TPL = "/Users/akazawt/works/data-science/tools/pptxkit/templates/aws-dark-2024.pptx"
OUT = sys.argv[1] if len(sys.argv) > 1 else "looped-figures.pptx"

GRN = RGBColor(36, 160, 70)
GRN_HI = RGBColor(59, 209, 111)
GRN_DK = RGBColor(18, 64, 34)
CHIP = RGBColor(22, 26, 30)
EDGE = RGBColor(70, 74, 80)
GREY = RGBColor(52, 56, 64)
LINK = RGBColor(40, 150, 240)
LINK_DK = RGBColor(16, 44, 72)
YEL = RGBColor(251, 211, 50)
YEL_DK = RGBColor(80, 66, 14)
MUTED = RGBColor(150, 150, 158)
WHITE = RGBColor(245, 245, 248)
DASH = MSO_LINE_DASH_STYLE.DASH

d = Deck(TPL, check_layout=False)
t = d.theme


def label(s, x, y, w, h, body, size=16, color=WHITE, bold=False, align="center"):
    return sh.textbox(s, t, x, y, w, h, body, size=size, color=color, bold=bold, align=align,
                      anchor="middle")


def box(s, x, y, w, h, fill, stroke=None, body="", size=15, color=WHITE, radius=0.1,
        dash=None, sw=1.25):
    sp = sh.rect(s, t, x, y, w, h, fill=fill, stroke=stroke, sw=sw, dash=dash,
                 shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=radius)
    if body:
        sh._write(sp.text_frame, t, body, size, color, True, "center", "middle")
        tf = sp.text_frame
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = sh.U(2)
    return sp


def title(s, head, sub):
    label(s, 43, 26, 1190, 46, head, size=30, bold=True, align="left")
    label(s, 43, 72, 1190, 26, sub, size=16, color=MUTED, align="left")


def arrow(s, x1, y1, x2, y2, color=WHITE, width=2, dashed=False):
    sh.arrow(s, t, x1, y1, x2, y2, color=color, width=width, dashed=dashed)




ORG = RGBColor(236, 114, 17)
ORG_DK = RGBColor(74, 38, 8)

def gpu_col(s, x, y, w, h, name):
    box(s, x, y, w, h, CHIP, stroke=EDGE, radius=0.05)
    label(s, x, y + 4, w, 24, name, size=13, color=MUTED)

# ------------------------------------------------------------------ 1. what each GPU holds
s = d.slide().s
title(s, "What each of 4 GPUs holds", "Green: weights.  Blue: requests or tokens.  Yellow: experts")
rows = [("TP", "weights split by column/row"), ("DP", "full weights, different requests"),
        ("CP", "sequence split by position"), ("EP", "experts placed per GPU"),
        ("SP", "activations split, with TP")]
GX, GW, GH = 330, 220, 82
for r, (name, sub) in enumerate(rows):
    y = 122 + r * 100
    label(s, 43, y + 8, 90, 40, name, size=24, bold=True, align="left")
    label(s, 43, y + 46, 280, 30, sub, size=13, color=MUTED, align="left")
    for g in range(4):
        x = GX + g * (GW + 12)
        box(s, x, y, GW, GH, CHIP, stroke=EDGE, radius=0.06)
        label(s, x + 6, y + 2, 60, 22, f"GPU {g}", size=11, color=MUTED, align="left")
        if name == "TP":
            box(s, x + 20, y + 30, GW - 40, 40, GRN_DK, stroke=GRN, body=f"W[:, {g}/4]", size=14)
        elif name == "DP":
            box(s, x + 12, y + 30, 120, 40, GRN_DK, stroke=GRN, body="W (all)", size=13)
            box(s, x + 140, y + 30, 68, 40, LINK_DK, stroke=LINK, body=f"req {g}", size=13)
        elif name == "CP":
            box(s, x + 20, y + 30, GW - 40, 40, LINK_DK, stroke=LINK, body=f"tokens {g}/4", size=14)
        elif name == "EP":
            box(s, x + 20, y + 30, GW - 40, 40, YEL_DK, stroke=YEL, body=f"expert {g}", size=14)
        else:
            box(s, x + 12, y + 30, 92, 40, GRN_DK, stroke=GRN, body="W 1/4", size=13)
            box(s, x + 112, y + 30, 96, 40, LINK_DK, stroke=LINK, body="act 1/4", size=13)

# ------------------------------------------------------------------ 2. KV duplication under TP
s = d.slide().s
title(s, "KV cache on 4 GPUs: TP copies it, DCP splits it", "Blue: KV cache.  Dashed: a duplicate copy")
panels = [(43, "MLA, TP = 4", "one latent KV, 4 copies"),
          (448, "GQA (2 KV heads), TP = 4", "2 copies of each head"),
          (853, "MLA, TP = 4, DCP = 4", "1/4 of the tokens each (interleave = 1)")]
for px, head, sub in panels:
    label(s, px, 112, 384, 30, head, size=19, bold=True)
    label(s, px, 142, 384, 24, sub, size=14, color=MUTED)
    for g in range(4):
        y = 180 + g * 98
        box(s, px, y, 384, 86, CHIP, stroke=EDGE, radius=0.06)
        label(s, px + 8, y + 4, 70, 22, f"GPU {g}", size=11, color=MUTED, align="left")
        if px == 43:
            box(s, px + 90, y + 22, 270, 46, LINK_DK, stroke=LINK, body="latent KV, all tokens",
                size=14, dash=None if g == 0 else DASH)
        elif px == 448:
            h = g // 2
            box(s, px + 90, y + 22, 270, 46, LINK_DK, stroke=LINK, body=f"KV head {h}, all tokens",
                size=14, dash=None if g % 2 == 0 else DASH)
        else:
            box(s, px + 90, y + 22, 270, 46, LINK_DK, stroke=LINK, body=f"tokens {g}, {g + 4}, {g + 8}, ...",
                size=14)

# ------------------------------------------------------------------ 3. 8 GPUs with P/D and AFD
s = d.slide().s
title(s, "Eight GPUs split by phase (P/D) and by module (AFD)",
      "Green: prefill GPUs.  Blue: attention GPUs and KV.  Yellow: expert GPUs.  Arrows: data moving")
box(s, 43, 120, 420, 470, CHIP, stroke=EDGE, radius=0.04)
label(s, 43, 128, 420, 30, "Prefill pool", size=19, bold=True)
for g in range(4):
    box(s, 73 + (g % 2) * 190, 190 + (g // 2) * 140, 170, 110, GRN_DK, stroke=GRN,
        body=f"GPU {g}\nattention + FFN", size=14)
label(s, 43, 480, 420, 60, "whole prompt at once\n(compute-intensive)", size=14, color=MUTED)
box(s, 523, 120, 714, 470, CHIP, stroke=EDGE, radius=0.04)
label(s, 523, 128, 714, 30, "Decode pool", size=19, bold=True)
box(s, 553, 180, 310, 330, CHIP, stroke=LINK, radius=0.04)
label(s, 553, 186, 310, 26, "attention sub-pool (DCP)", size=15, bold=True, color=LINK)
for g in range(2):
    box(s, 573 + g * 150, 230, 130, 230, LINK_DK, stroke=LINK, body=f"GPU {4 + g}\nKV 1/2", size=14)
box(s, 897, 180, 310, 330, CHIP, stroke=YEL, radius=0.04)
label(s, 897, 186, 310, 26, "FFN sub-pool (EP)", size=15, bold=True, color=YEL)
for g in range(2):
    box(s, 917 + g * 150, 230, 130, 230, YEL_DK, stroke=YEL, body=f"GPU {6 + g}\nexperts {2 * g}-{2 * g + 1}", size=14)
arrow(s, 463, 300, 553, 300, color=LINK, width=3)
label(s, 440, 254, 140, 40, "KV once", size=14, color=LINK)
arrow(s, 863, 320, 897, 320, color=WHITE)
arrow(s, 897, 370, 863, 370, color=WHITE)
label(s, 553, 516, 654, 30, "one token per step (memory-intensive)", size=14, color=MUTED)
label(s, 553, 546, 654, 30, "activations go back and forth every layer", size=14, color=MUTED)

# ------------------------------------------------------------------ 4. vLLM result
s = d.slide().s
title(s, "vLLM blog: TP vs DCP, Kimi K2.6 on 8 × B200", "Values as reported in the vLLM blog (2026-08-07)")
mets = [("concurrency (reported)", 64, 512, "64", "512", "TP (KV full)", "DCP (KV 82%)"),
        ("tok/s per GPU", 1863, 6091, "≈1,863", "6,091", "TP, c64", "DCP, c512"),
        ("KV memory used", 100, 82, "100%", "82%", "TP, c64", "DCP, c512")]
for i, (name, a, b, la, lb, ta, tb) in enumerate(mets):
    x = 43 + i * 405
    label(s, x, 120, 380, 30, name, size=19, bold=True)
    m = max(a, b)
    for j, (v, lab, col, dk, tag) in enumerate([(a, la, MUTED, GREY, ta), (b, lb, LINK, LINK_DK, tb)]):
        h = 330 * v / m
        bx = x + 70 + j * 140
        box(s, bx, 520 - h, 100, h, dk, stroke=col, radius=0.02)
        label(s, bx - 20, 520 - h - 34, 140, 30, lab, size=17, bold=True, color=WHITE)
        label(s, bx - 30, 530, 160, 30, tag, size=14, color=col)
label(s, 43, 590, 1190, 30, "TP: KV full at 64.  DCP: 82% at 512", size=15, color=MUTED, align="left")

d.save(OUT)
# ZN-24: テンプレートのロゴ、著作権表示、ヘッダー、フッターを必ず消してから使う
sys.path.insert(0, "/Users/akazawt/works/data-science/claudecode/.claude/skills/review-html/packs/zenn-doc/templates")
import strip_branding  # noqa: E402
strip_branding.main(OUT, OUT)
print(OUT)
