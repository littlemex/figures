"""gen_gdr_pptx.py: GPUDirect RDMA 入門の静止画 4 枚（EFA のスタック、GPUDirect RDMA なし、あり、BAR1 の窓）。"""
import os
import sys

sys.path.insert(0, "/Users/akazawt/works/data-science/tools/pptxkit")
from pptx.dml.color import RGBColor  # noqa: E402
from pptx.enum.dml import MSO_LINE_DASH_STYLE  # noqa: E402
from pptx.enum.shapes import MSO_SHAPE  # noqa: E402

from pptxkit import Deck, shapes as sh  # noqa: E402

TPL = "/Users/akazawt/works/data-science/tools/pptxkit/templates/aws-dark-2024.pptx"
OUT = sys.argv[1] if len(sys.argv) > 1 else "gdr-figures.pptx"

GRN = RGBColor(36, 160, 70)
GRN_HI = RGBColor(59, 209, 111)
GRN_DK = RGBColor(18, 64, 34)
CHIP = RGBColor(22, 26, 30)
EDGE = RGBColor(70, 74, 80)
GREY = RGBColor(52, 56, 64)
LINK = RGBColor(40, 150, 240)
LINK_DK = RGBColor(16, 44, 72)
YEL = RGBColor(251, 211, 50)
YEL_DK = RGBColor(80, 66, 14)
MUTED = RGBColor(150, 150, 158)
WHITE = RGBColor(245, 245, 248)
DASH = MSO_LINE_DASH_STYLE.DASH

d = Deck(TPL, check_layout=False)
t = d.theme


def label(s, x, y, w, h, body, size=16, color=WHITE, bold=False, align="center"):
    return sh.textbox(s, t, x, y, w, h, body, size=size, color=color, bold=bold, align=align,
                      anchor="middle")


def box(s, x, y, w, h, fill, stroke=None, body="", size=15, color=WHITE, radius=0.1,
        dash=None, sw=1.25):
    sp = sh.rect(s, t, x, y, w, h, fill=fill, stroke=stroke, sw=sw, dash=dash,
                 shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=radius)
    if body:
        sh._write(sp.text_frame, t, body, size, color, True, "center", "middle")
        tf = sp.text_frame
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = sh.U(2)
    return sp


def title(s, head, sub):
    label(s, 43, 26, 1190, 46, head, size=30, bold=True, align="left")
    label(s, 43, 72, 1190, 26, sub, size=16, color=MUTED, align="left")


def arrow(s, x1, y1, x2, y2, color=WHITE, width=2, dashed=False):
    sh.arrow(s, t, x1, y1, x2, y2, color=color, width=width, dashed=dashed)




PNK = RGBColor(255, 120, 160)
PNK_DK = RGBColor(80, 30, 46)
# ------------------------------------------------------------------ 1. EFA stack
s = d.slide().s
title(s, "EFA software stack: from the app to the EFA device",
      "Blue: user space.  Grey: kernel.  Yellow: hardware.  Pink: SRD link.  Solid: data path.  Dashed: setup only")
layers = [(130, "App: NCCL / NVSHMEM / MPI", LINK_DK, LINK),
          (215, "libfabric (efa provider): reorders, splits large messages", LINK_DK, LINK),
          (300, "rdma-core + efadv: creates the SRD QP", LINK_DK, LINK)]
for y, b, f, st in layers:
    box(s, 43, y, 620, 62, f, stroke=st, body=b, size=15)
for y in (192, 277):
    arrow(s, 353, y, 353, y + 23)
box(s, 260, 400, 403, 62, GREY, stroke=EDGE, body="efa kernel driver", size=15)
box(s, 43, 500, 620, 70, YEL_DK, stroke=YEL, body="EFA device (Nitro hardware)", size=16)
arrow(s, 150, 362, 150, 500, color=WHITE, width=3)
label(s, 160, 372, 220, 26, "data path: kernel bypass", size=13, color=WHITE, align="left")
arrow(s, 560, 362, 560, 400, color=MUTED, dashed=True)
arrow(s, 560, 462, 560, 500, color=MUTED, dashed=True)
label(s, 400, 362, 150, 26, "setup", size=13, color=MUTED, align="right")
label(s, 690, 130, 160, 30, "user space", size=15, color=LINK, align="left")
label(s, 690, 418, 160, 30, "kernel", size=15, color=MUTED, align="left")
label(s, 690, 580, 160, 30, "hardware", size=15, color=YEL, align="left")
box(s, 980, 500, 257, 70, YEL_DK, stroke=YEL, body="peer EFA device", size=16)
arrow(s, 663, 535, 980, 535, color=PNK, width=3)
arrow(s, 980, 535, 663, 535, color=PNK, width=3)
label(s, 700, 470, 260, 26, "SRD: multipath, reliable, unordered", size=13, color=PNK)

# ------------------------------------------------------------------ 2. without GPUDirect RDMA
s = d.slide().s
title(s, "Without GPUDirect RDMA: a detour through system memory",
      "Green: GPU memory.  Blue: system memory.  Yellow: NIC and network.  Numbers: steps")
box(s, 43, 260, 300, 90, GRN_DK, stroke=GRN, body="GPU memory", size=17)
box(s, 463, 260, 300, 90, LINK_DK, stroke=LINK, body="system memory\n(bounce buffer)", size=17)
box(s, 883, 260, 300, 90, YEL_DK, stroke=YEL, body="NIC", size=17)
arrow(s, 343, 305, 463, 305, color=WHITE, width=3)
arrow(s, 763, 305, 883, 305, color=WHITE, width=3)
arrow(s, 1183, 305, 1237, 305, color=YEL, width=3)
label(s, 343, 230, 120, 26, "1. CUDA copy", size=14)
label(s, 733, 230, 180, 26, "2. NIC DMA read", size=14)
label(s, 1150, 360, 120, 26, "network", size=14, color=YEL)

# ------------------------------------------------------------------ 3. with GPUDirect RDMA
s = d.slide().s
title(s, "With GPUDirect RDMA: NIC accesses GPU memory",
      "Green: GPU memory.  Yellow: NIC and network.  Grey box: not on the data path")
for x, name in ((43, "Machine A"), (693, "Machine B")):
    label(s, x, 112, 544, 30, name, size=18, bold=True)
    box(s, x, 150, 544, 330, CHIP, stroke=EDGE, radius=0.04)
box(s, 73, 180, 220, 80, GRN_DK, stroke=GRN, body="GPU memory", size=16)
box(s, 333, 330, 220, 80, YEL_DK, stroke=YEL, body="NIC", size=16)
box(s, 73, 370, 220, 80, GREY, stroke=EDGE, body="system memory", size=15)
box(s, 723, 330, 220, 80, YEL_DK, stroke=YEL, body="NIC", size=16)
box(s, 983, 180, 220, 80, GRN_DK, stroke=GRN, body="GPU memory", size=16)
box(s, 983, 370, 220, 80, GREY, stroke=EDGE, body="system memory", size=15)
arrow(s, 293, 230, 400, 330, color=WHITE, width=3)
arrow(s, 553, 370, 723, 370, color=YEL, width=3)
arrow(s, 880, 330, 983, 230, color=WHITE, width=3)
label(s, 553, 330, 170, 30, "network", size=14, color=YEL)
label(s, 43, 510, 1190, 30, "GPU → NIC → NIC → GPU", size=15, color=MUTED, align="left")

# ------------------------------------------------------------------ 4. BAR1
s = d.slide().s
title(s, "BAR1: a window onto GPU memory",
      "Grey: PCIe address space.  Green: GPU memory.  Yellow: NIC.  Pink: BAR1 window.  Dashed: maps to")
box(s, 43, 150, 1194, 70, CHIP, stroke=EDGE, body="", radius=0.06)
label(s, 63, 150, 600, 70, "PCIe address space (shared by CPU, NIC, GPU)", size=15, align="left")
box(s, 700, 160, 260, 50, PNK_DK, stroke=PNK, body="BAR1 window", size=15)
box(s, 600, 330, 460, 120, GRN_DK, stroke=GRN, body="", radius=0.06)
label(s, 615, 335, 240, 30, "GPU memory (VRAM)", size=16, bold=True, align="left")
box(s, 780, 385, 240, 50, PNK_DK, stroke=PNK, body="mapped region", size=14)
arrow(s, 900, 210, 900, 385, color=PNK, width=3, dashed=True)
box(s, 103, 330, 300, 120, YEL_DK, stroke=YEL, body="NIC", size=17)
arrow(s, 403, 360, 720, 210, color=WHITE, width=3)
label(s, 120, 280, 320, 30, "DMA to the window address", size=14, align="left")
d.save(OUT)
# ZN-24: テンプレートのロゴ、著作権表示、ヘッダー、フッターを必ず消してから使う
sys.path.insert(0, "/Users/akazawt/works/data-science/claudecode/.claude/skills/review-html/packs/zenn-doc/templates")
import strip_branding  # noqa: E402
strip_branding.main(OUT, OUT)
print(OUT)
