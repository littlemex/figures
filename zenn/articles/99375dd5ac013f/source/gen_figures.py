"""gen_figures.py -- post-training 時代のスケーリング則と台数見積もりの記事の図 5 枚を、ネイティブ shape の pptx で組む。

    /usr/bin/python3 gen_figures.py <out.pptx>

見た目は review-html の zenn-doc テンプレート (nxd_interfaces_pptx.py) にそろえる。数値は calc.py の出力と同じ。
"""
import os
import sys

sys.path.insert(0, "/Users/akazawt/works/data-science/tools/pptxkit")
sys.path.insert(0, "/Users/akazawt/works/data-science/claudecode/.claude/skills/review-html/packs/zenn-doc/templates")
from pptx.dml.color import RGBColor  # noqa: E402
from pptx.enum.shapes import MSO_SHAPE  # noqa: E402

from pptxkit import Deck, shapes as sh  # noqa: E402

TPL = "/Users/akazawt/works/data-science/tools/pptxkit/templates/aws-dark-2024.pptx"
OUT = sys.argv[1] if len(sys.argv) > 1 else "figures.pptx"

GRN = RGBColor(36, 160, 70)
GRN_HI = RGBColor(59, 209, 111)
GRN_DK = RGBColor(18, 64, 34)
CHIP = RGBColor(22, 26, 30)
EDGE = RGBColor(70, 74, 80)
GREY = RGBColor(52, 56, 64)
LINK = RGBColor(40, 150, 240)
LINK_DK = RGBColor(16, 44, 72)
YEL = RGBColor(251, 211, 50)
YEL_DK = RGBColor(80, 66, 14)
MUTED = RGBColor(150, 150, 158)
WHITE = RGBColor(245, 245, 248)

d = Deck(TPL, check_layout=False)
t = d.theme


def label(s, x, y, w, h, body, size=16, color=WHITE, bold=False, align="center"):
    return sh.textbox(s, t, x, y, w, h, body, size=size, color=color, bold=bold, align=align,
                      anchor="middle")


def box(s, x, y, w, h, fill, stroke=None, body="", size=15, color=WHITE, radius=0.1, dash=None, sw=1.25):
    sp = sh.rect(s, t, x, y, w, h, fill=fill, stroke=stroke, sw=sw, dash=dash,
                 shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=radius)
    if body:
        sh._write(sp.text_frame, t, body, size, color, True, "center", "middle")
        tf = sp.text_frame
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = sh.U(2)
    return sp


def title(s, head, sub):
    label(s, 43, 26, 1190, 46, head, size=30, bold=True, align="left")
    label(s, 43, 72, 1190, 26, sub, size=16, color=MUTED, align="left")


def arrow(s, x1, y1, x2, y2, color=WHITE, width=2):
    sh.arrow(s, t, x1, y1, x2, y2, color=color, width=width)


# ------------------------------------------------------------------ 1. three scaling laws
s = d.slide().s
title(s, "From one to three scaling laws",
      "Green: what scales.  Blue: what sets the throughput.  Yellow: the shared base")
cols = [("Pre-training", "parameters · tokens", "parallelism config\nfixed per stage"),
        ("Post-training", "SFT data · RL compute", "rollout lengths · KV seats\nenv CPU time"),
        ("Test-time compute", "tokens · samples per answer", "request mix\n+ KV seats")]
label(s, 43, 200, 150, 80, "scales", size=15, color=GRN_HI, align="left")
label(s, 43, 300, 150, 100, "throughput", size=15, color=LINK, align="left")
label(s, 43, 452, 150, 70, "base", size=15, color=YEL, align="left")
x = 193
CW = 330
for i, (head, what, sets) in enumerate(cols):
    box(s, x, 120, CW, 60, CHIP, stroke=EDGE, body=head, size=21)
    box(s, x, 200, CW, 80, GRN_DK, stroke=GRN, body=what, size=16)
    box(s, x, 300, CW, 100, LINK_DK, stroke=LINK, body=sets, size=16)
    arrow(s, x + CW // 2, 400, x + CW // 2, 452, color=MUTED)
    if i < 2:
        arrow(s, x + CW, 150, x + CW + 22, 150)
    x += CW + 22
box(s, 193, 452, 1044, 70, YEL_DK, stroke=YEL, radius=0.25,
    body="tightly coupled GPUs  ·  high-bandwidth low-latency network  ·  scalable distributed storage", size=16)
label(s, 193, 532, 1044, 34, "on this cloud: EFA  ·  FSx for Lustre  ·  S3", size=15, color=MUTED)

# ------------------------------------------------------------------ 2. pre-training cost loop
s = d.slide().s
title(s, "Pre-training: the config fixes the throughput",
      "Green: you choose it.  Blue: it follows.  Yellow: watch it every step")
steps = [("Config", "TP / PP / DP / EP\ndtype, recompute", GRN_DK, GRN),
         ("Step time", "stable\nwithin a stage", LINK_DK, LINK),
         ("Tokens / s", "per GPU", LINK_DK, LINK),
         ("GPU-hours", "C / (3600 · peak · MFU · A)", LINK_DK, LINK)]
x = 43
W = 270
for i, (head, sub, fill, edge) in enumerate(steps):
    box(s, x, 130, W, 60, fill, stroke=edge, body=head, size=20)
    label(s, x, 196, W, 56, sub, size=15, color=WHITE)
    if i < 3:
        arrow(s, x + W, 160, x + W + 38, 160)
    x += W + 38
box(s, 43, 300, 1194, 150, CHIP, stroke=EDGE, radius=0.05)
label(s, 63, 310, 700, 30, "Monitor (DCGM, Prometheus, Grafana)", size=17, bold=True, color=YEL, align="left")
for j, m in enumerate(["step time", "MFU", "Tensor Core active", "exposed comm", "restarts · ckpt"]):
    box(s, 63 + j * 232, 360, 210, 60, YEL_DK, stroke=YEL, body=m, size=16)
arrow(s, 1100, 300, 1100, 252, color=YEL)
label(s, 1110, 262, 120, 30, "A", size=16, color=YEL, bold=True, align="left")
box(s, 43, 490, 1194, 60, CHIP, stroke=EDGE, radius=0.25,
    body="model: −1 pt MFU at 35%  →  +2.9% GPU-hours", size=18)

# ------------------------------------------------------------------ 3. H200 -> B200 transfer
s = d.slide().s
title(s, "H200 to B200: each part speeds up by its own ratio",
      "Ratios from a simple model.  Blue: spec sheet.  Yellow: about ×1")
rows = [("Compute (bf16 dense)", 2.27, LINK), ("HBM bandwidth", 1.60, LINK), ("NVLink per GPU *", 2.0, LINK),
        ("EFA per node", 1.0, YEL), ("Waiting (data load)", 1.0, YEL)]
y = 130
SCALE = 330
for name, r, col in rows:
    label(s, 43, y, 300, 50, name, size=17, align="left")
    sh.rect(s, t, 350, y + 8, int(r * SCALE), 34, fill=col, stroke=None)
    label(s, 360 + int(r * SCALE), y, 420, 50, "× %.2f" % r if r != 1.0 else "× 1.0 (EFAv4: +18% coll.)" if name.startswith("EFA") else "× 1.0",
          size=16, align="left")
    y += 62
sh.rect(s, t, 350 + SCALE, 128, 2, 312, fill=MUTED, stroke=None)
label(s, 350 + SCALE - 60, 440, 120, 26, "× 1", size=13, color=MUTED)
label(s, 43, 452, 1194, 26, "* not split out in the four cases", size=13, color=MUTED, align="left")
box(s, 43, 492, 1194, 70, CHIP, stroke=EDGE, radius=0.2,
    body="8 × p5en  →  4 to 6 × p6-b200, set by the step-time breakdown", size=19)

# ------------------------------------------------------------------ 4. RL step
s = d.slide().s
title(s, "RL: rollout takes the time, the trainer takes the FLOPs",
      "Green: generates.  Blue: learns.  Yellow: outside the GPU.  Gray: weight sync")
box(s, 43, 120, 330, 80, GRN_DK, stroke=GRN, body="Rollout\nSGLang / vLLM", size=18)
box(s, 473, 120, 330, 80, YEL_DK, stroke=YEL, body="Reward · environment\nverifier, sandbox, tools", size=17)
box(s, 903, 120, 334, 80, LINK_DK, stroke=LINK, body="Trainer\nMegatron / FSDP", size=18)
arrow(s, 373, 160, 473, 160)
arrow(s, 803, 160, 903, 160)
label(s, 373, 124, 100, 30, "samples", size=13, color=MUTED)
sh.arrow(s, t, 1070, 200, 1070, 236, color=WHITE, width=2)
sh.rect(s, t, 208, 236, 864, 2, fill=WHITE, stroke=None)
arrow(s, 208, 236, 208, 200)
label(s, 500, 238, 300, 28, "weight sync", size=13, color=MUTED)
# timeline of one synchronous step
label(s, 43, 290, 600, 30, "One synchronous step (illustration)", size=16, bold=True, align="left")
segs = [("rollout: most requests", 0.45, GRN_DK, GRN), ("long tail", 0.25, GRN_DK, GRN_HI),
        ("env", 0.05, YEL_DK, YEL), ("logprob + train", 0.20, LINK_DK, LINK), ("sync", 0.05, GREY, EDGE)]
x = 43
for name, f, fill, edge in segs:
    w = int(1194 * f)
    box(s, x, 328, w, 56, fill, stroke=edge, body=name if f > 0.06 else "", size=15, radius=0.05)
    if f <= 0.06:
        label(s, x - 20, 388, w + 40, 24, name, size=13, color=MUTED)
    x += w
box(s, 43, 430, 580, 70, CHIP, stroke=EDGE, body="FLOPs: rollout ≈ 20% (model)", size=18)
box(s, 657, 430, 580, 70, CHIP, stroke=EDGE, body="time: rollout 63 – 87% (Seer, Table 1)", size=18)
box(s, 43, 520, 1194, 56, CHIP, stroke=EDGE, radius=0.3,
    body="the breakdown drifts: answers grow longer, KV seats and the tail change", size=17)

# ------------------------------------------------------------------ 5. two ways to build it
s = d.slide().s
title(s, "Two ways to run RL post-training",
      "Green: the RL framework.  Blue: the cluster layer.  Yellow: the shared base")
for x, head, orch, sched, note in ((43, "Slurm", "ParallelCluster · PCS\nHyperPod (Slurm)", "Ray started\ninside the Slurm job", "Ray on Slurm"),
                                   (665, "Kubernetes", "EKS · HyperPod (EKS)", "KubeRay\n+ Kueue / KAI", "reference examples on GitHub")):
    label(s, x, 112, 572, 36, head, size=22, bold=True)
    box(s, x, 156, 572, 300, CHIP, stroke=EDGE, radius=0.04)
    box(s, x + 30, 176, 512, 70, LINK_DK, stroke=LINK, body=orch, size=17)
    box(s, x + 30, 266, 512, 70, LINK_DK, stroke=LINK, body=sched, size=17)
    box(s, x + 30, 356, 512, 80, GRN_DK, stroke=GRN, body="slime · miles · NeMo RL · verl\n" + note, size=17)
    arrow(s, x + 286, 246, x + 286, 266)
    arrow(s, x + 286, 336, x + 286, 356)
box(s, 43, 476, 1194, 50, YEL_DK, stroke=YEL, radius=0.3,
    body="p5en · p6-b200 · p6-b300  ·  EFA  ·  FSx for Lustre  ·  S3", size=17)
label(s, 43, 536, 1194, 36, "rollout and trainer can sit on different node groups", size=16, color=MUTED)

d.save(OUT)
import strip_branding  # noqa: E402
strip_branding.main(OUT, OUT)
print(OUT)
