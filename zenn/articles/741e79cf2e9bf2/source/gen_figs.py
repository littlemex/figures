"""Figures for the Zenn article on SGLang same-node KV transfer through NIXL and libfabric.

    /usr/bin/python3 gen_figs.py <out.pptx>

Figure text is English. Sources: SGLang 5a15cde8, NIXL 1.1.0 05e4243f / 1.4.0 c0a1102b, libfabric 5a13558c.
"""
import os
import sys

sys.path.insert(0, "/Users/akazawt/works/data-science/tools/pptxkit")
from pptx.dml.color import RGBColor  # noqa: E402
from pptx.enum.shapes import MSO_SHAPE  # noqa: E402

from pptxkit import Deck, shapes as sh  # noqa: E402

TPL = "/Users/akazawt/works/data-science/tools/pptxkit/templates/aws-dark-2024.pptx"
OUT = sys.argv[1] if len(sys.argv) > 1 else "figs.pptx"

GRN = RGBColor(36, 160, 70)
GRN_HI = RGBColor(59, 209, 111)
GRN_DK = RGBColor(18, 64, 34)
CHIP = RGBColor(22, 26, 30)
EDGE = RGBColor(70, 74, 80)
GREY = RGBColor(52, 56, 64)
LINK = RGBColor(40, 150, 240)
LINK_DK = RGBColor(16, 44, 72)
YEL = RGBColor(251, 211, 50)
YEL_DK = RGBColor(80, 66, 14)
RED = RGBColor(230, 80, 70)
RED_DK = RGBColor(80, 24, 22)
MUTED = RGBColor(150, 150, 158)
WHITE = RGBColor(245, 245, 248)

d = Deck(TPL, check_layout=False)
t = d.theme


def label(s, x, y, w, h, body, size=16, color=WHITE, bold=False, align="center"):
    return sh.textbox(s, t, x, y, w, h, body, size=size, color=color, bold=bold, align=align,
                      anchor="middle")


def box(s, x, y, w, h, fill, stroke=None, body="", size=15, color=WHITE, radius=0.1, dash=None, sw=1.25):
    sp = sh.rect(s, t, x, y, w, h, fill=fill, stroke=stroke, sw=sw, dash=dash,
                 shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=radius)
    if body:
        sh._write(sp.text_frame, t, body, size, color, True, "center", "middle")
        tf = sp.text_frame
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = sh.U(2)
    return sp


def arrow(s, x1, y1, x2, y2, color=WHITE, width=2):
    sh.arrow(s, t, x1, y1, x2, y2, color=color, width=width)


def title(s, head, sub):
    label(s, 43, 26, 1190, 46, head, size=30, bold=True, align="left")
    label(s, 43, 72, 1190, 26, sub, size=16, color=MUTED, align="left")


# 1. map
s = d.slide().s
title(s, "Who moves the KV cache", "Prefill and decode on one node, NIXL with its LIBFABRIC backend, FI_PROVIDER=efa")
for x, role in ((43, "SGLang prefill"), (683, "SGLang decode")):
    box(s, x, 120, 554, 150, CHIP, stroke=EDGE, radius=0.05)
    label(s, x + 14, 126, 300, 30, role, size=18, bold=True, align="left")
    box(s, x + 20, 166, 245, 44, GRN_DK, stroke=GRN, body="TP rank 0  (device 0)", size=14)
    box(s, x + 289, 166, 245, 44, GRN_DK, stroke=GRN, body="TP rank 1  (device 1)", size=14)
    label(s, x + 20, 218, 514, 40, "one process per rank; device numbers as seen by each process", size=13, color=MUTED)
box(s, 43, 300, 1194, 56, LINK_DK, stroke=LINK, body="NIXL  (LIBFABRIC backend)", size=18)
box(s, 43, 386, 1194, 56, LINK_DK, stroke=LINK, body="libfabric  efa provider", size=18)
box(s, 43, 486, 560, 90, YEL_DK, stroke=YEL, body="peer on the same node\nshm provider (shared memory, CUDA IPC)", size=16)
box(s, 677, 486, 560, 90, GREY, stroke=EDGE, body="peer on another node\nEFA device (RDMA)", size=16)
arrow(s, 323, 442, 323, 486, YEL, 2.5)
arrow(s, 957, 442, 957, 486)
label(s, 43, 590, 1194, 30, "By default, same-node data does not go through the EFA device", size=15, color=YEL)

# 2. efa hands off to shm: who owns which step
s = d.slide().s
title(s, "Who does what when NIXL writes to a same-node peer",
      "FI_REMOTE_CQ_DATA is a libfabric flag; NIXL sets it, the libfabric providers act on it")
lanes = [("SGLang", GRN_DK, GRN), ("NIXL\nLIBFABRIC backend", LINK_DK, LINK),
         ("libfabric\nefa provider", YEL_DK, YEL), ("libfabric\nshm provider", YEL_DK, YEL)]
for i, (name, fill, stroke) in enumerate(lanes):
    y = 118 + i * 112
    box(s, 43, y, 1194, 102, CHIP, stroke=EDGE, radius=0.04)
    box(s, 55, y + 12, 200, 78, fill, stroke=stroke, body=name, size=15)
box(s, 290, 139, 320, 60, GRN_DK, stroke=GRN, body="send the KV cache", size=15)
box(s, 290, 251, 320, 60, LINK_DK, stroke=LINK, body="post an RDMA write\n1.1.0 fi_writedata / 1.4.0 fi_writemsg", size=13)
arrow(s, 450, 199, 450, 251)
box(s, 660, 263, 250, 36, YEL_DK, stroke=YEL, body="+ FI_REMOTE_CQ_DATA", size=14, radius=0.3)
arrow(s, 610, 281, 660, 281, YEL, 2)
box(s, 660, 363, 250, 60, YEL_DK, stroke=YEL, body="peer on this node?", size=15)
arrow(s, 785, 299, 785, 363, YEL, 2)
box(s, 975, 363, 250, 60, GREY, stroke=GRN_HI, body="EFA device\nRDMA write", size=14)
arrow(s, 910, 393, 975, 393, GRN_HI, 2.5)
label(s, 908, 352, 70, 36, "no, or\nshm off", size=11, color=GRN_HI)
box(s, 660, 475, 250, 60, YEL_DK, stroke=YEL, body="flag set:\nskip the CMA fast path", size=14)
arrow(s, 785, 423, 785, 475, YEL, 2.5)
label(s, 795, 432, 160, 36, "yes, default (=1)", size=12, color=YEL, align="left")
box(s, 975, 463, 250, 38, LINK_DK, stroke=LINK, body="GPU memory: CUDA IPC", size=13)
box(s, 975, 509, 250, 38, RED_DK, stroke=RED, body="IPC setup fails: SAR", size=13)
arrow(s, 910, 495, 975, 482, YEL, 2)
arrow(s, 910, 515, 975, 528, YEL, 2)
legend = [("SGLang", GRN_DK, GRN), ("NIXL", LINK_DK, LINK), ("libfabric", YEL_DK, YEL)]
for i, (name, fill, stroke) in enumerate(legend):
    x = 43 + i * 170
    box(s, x, 580, 26, 20, fill, stroke=stroke, radius=0.2)
    label(s, x + 32, 576, 130, 28, name, size=13, align="left")

# 3. CUDA IPC on shm
s = d.slide().s
title(s, "shm with GPU memory: CUDA IPC", "The sender hands over a handle; the receiver opens it and copies")
label(s, 43, 112, 560, 30, "sender (posting thread)", size=17, bold=True, color=GRN_HI)
label(s, 677, 112, 560, 30, "receiver (progress loop)", size=17, bold=True, color=LINK)
sy = [150, 240, 330]
for y, body in zip(sy, ["find the buffer's base\ncuMemGetAddressRange", "export an IPC handle\ncudaIpcGetMemHandle", "put the handle in a\ncommand in shared memory"]):
    box(s, 93, y, 460, 70, GRN_DK, stroke=GRN, body=body, size=15)
for a, b in zip(sy, sy[1:]):
    arrow(s, 323, a + 70, 323, b)
ry = [240, 330, 420]
for y, body in zip(ry, ["open the handle\ncudaIpcOpenMemHandle", "copy into its own buffer\ncudaMemcpy on its own thread", "mark done after the copy returns"]):
    box(s, 727, y, 460, 70, LINK_DK, stroke=LINK, body=body, size=15)
for a, b in zip(ry, ry[1:]):
    arrow(s, 957, a + 70, 957, b)
arrow(s, 553, 365, 727, 275, YEL, 2.5)
box(s, 93, 430, 460, 60, CHIP, stroke=EDGE, body="the sender sees completion only after that", size=14, color=MUTED)
arrow(s, 727, 455, 553, 460, MUTED, 2)
box(s, 43, 540, 1194, 60, YEL_DK, stroke=YEL, radius=0.3, body="NIXL writes carry FI_REMOTE_CQ_DATA, so shm does not take its sender-side fast path", size=15)

# 4. SAR
s = d.slide().s
title(s, "SAR: segmentation and reassembly", "The fallback when preparing CUDA IPC fails: copy through shared host memory")
box(s, 43, 150, 250, 90, GRN_DK, stroke=GRN, body="sender's\nGPU buffer", size=16)
box(s, 987, 150, 250, 90, LINK_DK, stroke=LINK, body="receiver's\nGPU buffer", size=16)
box(s, 380, 120, 520, 150, CHIP, stroke=EDGE, radius=0.05)
label(s, 390, 126, 500, 30, "receiver's shared memory: bounce buffers", size=15, color=MUTED)
for i in range(6):
    box(s, 400 + i * 82, 172, 70, 50, GREY, stroke=EDGE, body="32K", size=13)
label(s, 390, 228, 500, 30, "up to 64 buffers of 32 KiB (also capped per peer), reused", size=14, color=MUTED)
arrow(s, 293, 195, 380, 195, YEL, 2.5)
arrow(s, 900, 195, 987, 195, YEL, 2.5)
label(s, 293, 245, 87, 30, "copy in", size=13, color=YEL)
label(s, 900, 245, 87, 30, "copy out", size=13, color=YEL)
box(s, 43, 330, 1194, 200, CHIP, stroke=EDGE, radius=0.04)
steps = [("sender fills buffers", "flag: FULL"), ("receiver drains them", "flag: EMPTY"), ("repeat until the whole\nmessage is across", "")]
x = 73
for i, (a, b) in enumerate(steps):
    box(s, x, 360, 340, 80, YEL_DK, stroke=YEL, body=a, size=15)
    label(s, x, 450, 340, 30, b, size=14, color=MUTED)
    if i < 2:
        arrow(s, x + 340, 400, x + 380, 400)
    x += 380
label(s, 43, 560, 1194, 30, "Two copies per piece, relayed through shared host memory", size=16, color=YEL)

# 5. per-thread device
s = d.slide().s
title(s, "The current CUDA device is per thread", "SGLang's NIXL transfer threads never select their GPU")
box(s, 43, 120, 1194, 400, CHIP, stroke=EDGE, radius=0.04)
label(s, 63, 128, 700, 30, "prefill process, TP rank 1  (its KV cache is on device 1)", size=17, bold=True, align="left")
box(s, 83, 180, 340, 70, GRN_DK, stroke=GRN, body="main thread\ncurrent device: 1", size=15)
box(s, 83, 300, 340, 90, RED_DK, stroke=RED, body="transfer threads\ncontext: device 0\n(never set)", size=15)
box(s, 520, 300, 300, 90, LINK_DK, stroke=LINK, body="shm: where is this\nbuffer?", size=15)
box(s, 900, 300, 300, 90, RED_DK, stroke=RED, body="device 0 context:\naddress not found\nCUDA_ERROR_NOT_FOUND", size=14)
arrow(s, 423, 345, 520, 345)
arrow(s, 820, 345, 900, 345, RED, 2.5)
box(s, 520, 420, 680, 70, YEL_DK, stroke=YEL, body="no IPC handle: warn once, fall back to SAR", size=16)
arrow(s, 1050, 390, 1050, 420, YEL, 2.5)
label(s, 43, 540, 1194, 30, "TP rank 0 is unaffected only because its KV cache is on the default device", size=15, color=MUTED)
label(s, 43, 580, 1194, 30, "Tested patch (not upstream): set the device at the start of each transfer thread", size=16, color=GRN_HI)

# 6. NIXL 1.4.0
s = d.slide().s
title(s, "NIXL 1.4.0 restores the context before posting", "In the measured setup, the SAR fallback warnings disappear without changing SGLang")
for x, head, col, dk, rows in (
        (43, "NIXL 1.1.0", RED, RED_DK, ["transfer thread posts", "context: device 0", "IPC handle fails", "SAR"]),
        (665, "NIXL 1.4.0", GRN_HI, GRN_DK, ["transfer thread posts", "NIXL sets the registration\ncontext first", "IPC handle succeeds", "CUDA IPC"])):
    label(s, x, 112, 572, 34, head, size=22, bold=True, color=col)
    y = 160
    for i, r in enumerate(rows):
        last = i == len(rows) - 1
        box(s, x + 60, y, 452, 70, dk if (last or i in (1, 2)) else CHIP, stroke=col if (last or i in (1, 2)) else EDGE, body=r, size=15)
        if not last:
            arrow(s, x + 286, y + 70, x + 286, y + 100)
        y += 100
label(s, 43, 590, 1194, 30, "This fixes the sender side only; the receiver side is the next figure", size=15, color=MUTED)

# 7. progress thread
s = d.slide().s
title(s, "NIXL's progress thread", "A background thread that polls completions; in the measured setup the shm receive copy ran in it")
box(s, 43, 120, 1194, 450, CHIP, stroke=EDGE, radius=0.04)
label(s, 63, 128, 700, 30, "decode process (one per TP rank)", size=17, bold=True, align="left")
box(s, 83, 180, 330, 70, GRN_DK, stroke=GRN, body="scheduler thread\nruns the model", size=15)
box(s, 83, 300, 330, 90, LINK_DK, stroke=LINK, body="progress thread (when enabled)\nloop: progressActiveRails()", size=15)
box(s, 500, 300, 330, 90, LINK_DK, stroke=LINK, body="fi_cq_read()\nholding ep_mutex_", size=15)
box(s, 917, 300, 290, 90, YEL_DK, stroke=YEL, body="shm progress\nreceive copy (cudaMemcpy)", size=14)
arrow(s, 413, 345, 500, 345)
arrow(s, 830, 345, 917, 345, YEL, 2.5)
box(s, 500, 440, 707, 90, CHIP, stroke=EDGE, body="the same lock guards posting writes;\ncalls that need it would wait (not measured)", size=14, color=MUTED)
arrow(s, 665, 390, 665, 440, MUTED, 2)
label(s, 83, 450, 330, 60, "started by NIXL;\nit sets no CUDA device", size=14, color=RED)

# 8. progress thread device
s = d.slide().s
title(s, "Decode TP rank 1: progress thread on the wrong device", "Measured on the shm path with NIXL 1.4.0")
for x, head, col, dk, a, b, c in (
        (43, "as shipped", RED, RED_DK, "progress thread\ncontext: device 0", "destination: device 1\n(decode TP rank 1)", "host-staged copy\nin two steps"),
        (665, "creator's device (test patch)", GRN_HI, GRN_DK, "progress thread\ncontext: device 1", "destination: device 1", "one direct copy")):
    label(s, x, 112, 572, 34, head, size=20, bold=True, color=col)
    box(s, x + 40, 170, 492, 80, dk, stroke=col, body=a, size=16)
    box(s, x + 40, 290, 492, 80, CHIP, stroke=EDGE, body=b, size=16)
    box(s, x + 40, 410, 492, 80, dk, stroke=col, body=c, size=16)
    arrow(s, x + 286, 250, x + 286, 290)
    arrow(s, x + 286, 370, x + 286, 410)
label(s, 43, 520, 1194, 30, "decode TP rank 0 is unaffected: its device is the default one", size=15, color=MUTED)
label(s, 43, 560, 1194, 30, "why fixing the thread's device shortens decode steps: not known", size=15, color=YEL)

# 9. options
s = d.slide().s
title(s, "Three setups on one node", "Where each problem remains, as measured")
cols = [("NIXL 1.1.0, shm", ["sender: SAR on TP rank 1", "receiver: not measured", "receiving thread:\nnot measured"], RED, RED_DK),
        ("NIXL 1.4.0, shm", ["sender: CUDA IPC", "receiver: wrong-GPU copy\non decode TP rank 1", "receiver copies in the\nprogress thread"], YEL, YEL_DK),
        ("FI_EFA_ENABLE_SHM_TRANSFER=0", ["no shm path", "RDMA write through\nthe EFA device", "no shm receive copy"], GRN_HI, GRN_DK)]
x = 43
for head, rows, col, dk in cols:
    label(s, x, 112, 380, 34, head, size=17, bold=True, color=col)
    y = 160
    for r in rows:
        box(s, x, y, 380, 80, dk, stroke=col, body=r, size=15)
        y += 100
    x += 407
label(s, 43, 470, 1194, 30, "One transfer alone is slower through the device; under load the KV transfer stage was shorter", size=15, color=MUTED)

# 10. overview: two nodes vs one node
s = d.slide().s
title(s, "Where the bytes go", "libfabric with FI_PROVIDER=efa: another node vs the same node")
box(s, 43, 120, 560, 470, CHIP, stroke=EDGE, radius=0.04)
label(s, 63, 128, 520, 30, "peer on another node", size=18, bold=True, align="left")
box(s, 73, 175, 220, 60, GRN_DK, stroke=GRN, body="GPU (node A)", size=15)
box(s, 353, 175, 220, 60, GRN_DK, stroke=GRN, body="GPU (node B)", size=15)
box(s, 73, 290, 220, 60, LINK_DK, stroke=LINK, body="efa provider", size=15)
box(s, 353, 290, 220, 60, LINK_DK, stroke=LINK, body="efa provider", size=15)
box(s, 73, 405, 220, 60, GREY, stroke=EDGE, body="EFA device", size=15)
box(s, 353, 405, 220, 60, GREY, stroke=EDGE, body="EFA device", size=15)
arrow(s, 183, 235, 183, 290)
arrow(s, 183, 350, 183, 405)
arrow(s, 293, 435, 353, 435, GRN_HI, 3)
arrow(s, 463, 405, 463, 350)
arrow(s, 463, 290, 463, 235)
label(s, 73, 490, 500, 60, "through the network", size=15, color=GRN_HI)
box(s, 677, 120, 560, 470, CHIP, stroke=EDGE, radius=0.04)
label(s, 697, 128, 520, 30, "peer on the same node", size=18, bold=True, align="left")
box(s, 707, 175, 220, 60, GRN_DK, stroke=GRN, body="GPU of process A", size=15)
box(s, 987, 175, 220, 60, GRN_DK, stroke=GRN, body="GPU of process B", size=15)
box(s, 707, 270, 500, 54, LINK_DK, stroke=LINK, body="efa provider: the peer is local", size=15)
box(s, 707, 370, 240, 110, YEL_DK, stroke=YEL, body="default (=1)\nshm provider\nCUDA IPC or SAR", size=15)
box(s, 967, 370, 240, 110, GRN_DK, stroke=GRN_HI, body="=0\nEFA device\n(measured: RDMA write)", size=15)
arrow(s, 817, 235, 817, 270)
arrow(s, 827, 324, 827, 370, YEL, 2.5)
arrow(s, 1087, 324, 1087, 370, GRN_HI, 2.5)
label(s, 707, 495, 500, 60, "FI_EFA_ENABLE_SHM_TRANSFER\ndecides which one", size=15, color=YEL)

# 11. shm provider
s = d.slide().s
title(s, "shm provider", "Commands and responses go through regions of host memory; how the data moves depends on the protocol")
box(s, 43, 150, 300, 110, GRN_DK, stroke=GRN, body="process A\n(sender)", size=17)
box(s, 937, 150, 300, 110, LINK_DK, stroke=LINK, body="process B\n(receiver)", size=17)
box(s, 403, 120, 474, 200, CHIP, stroke=YEL, radius=0.05)
label(s, 413, 128, 454, 30, "receiver's region in host memory", size=15, color=YEL)
box(s, 433, 165, 414, 60, GREY, stroke=EDGE, body="command queue", size=15)
box(s, 433, 240, 414, 60, GREY, stroke=EDGE, body="SAR bounce buffers", size=15)
arrow(s, 343, 195, 433, 195, YEL, 2.5)
arrow(s, 847, 195, 937, 195, YEL, 2.5)
box(s, 403, 370, 474, 130, CHIP, stroke=EDGE, radius=0.05)
label(s, 413, 378, 454, 30, "sender's region in host memory", size=15, color=MUTED)
box(s, 433, 415, 414, 60, GREY, stroke=EDGE, body="response queue", size=15)
arrow(s, 960, 260, 847, 430, MUTED, 2)
arrow(s, 433, 430, 320, 260, MUTED, 2)
label(s, 43, 280, 260, 60, "opens the receiver's region\n(shm_open + mmap)", size=14, color=MUTED, align="left")
label(s, 43, 540, 1194, 30, "Each process creates its own region and maps its peers' regions", size=15, color=MUTED)

# 12. CUDA IPC (basic)
s = d.slide().s
title(s, "CUDA IPC", "A send into a GPU buffer: the receiver copies with a handle from the sender")
box(s, 43, 150, 330, 110, GRN_DK, stroke=GRN, body="sender's GPU buffer", size=16)
box(s, 907, 150, 330, 110, LINK_DK, stroke=LINK, body="receiver's GPU buffer", size=16)
box(s, 473, 110, 334, 70, YEL_DK, stroke=YEL, body="IPC handle\n(in a shm command)", size=15)
arrow(s, 373, 170, 473, 145, YEL, 2.5)
arrow(s, 807, 145, 907, 170, YEL, 2.5)
arrow(s, 373, 235, 907, 235, GRN_HI, 3.5)
label(s, 473, 245, 334, 30, "1 copy, done by the receiver", size=15, color=GRN_HI)
box(s, 43, 340, 1194, 120, CHIP, stroke=EDGE, radius=0.05)
label(s, 63, 350, 1154, 30, "precondition on the sender", size=16, bold=True, color=YEL, align="left")
label(s, 63, 385, 1154, 60, "cuMemGetAddressRange must find the buffer in the sending thread's\ncurrent CUDA context, or no handle can be exported", size=15, align="left")

# 13. SAR (basic)
s = d.slide().s
title(s, "SAR", "Segmentation and reassembly through bounce buffers; here, after the IPC handle fails")
box(s, 43, 150, 270, 110, GRN_DK, stroke=GRN, body="sender's GPU buffer", size=16)
box(s, 967, 150, 270, 110, LINK_DK, stroke=LINK, body="receiver's GPU buffer", size=16)
box(s, 433, 140, 414, 130, CHIP, stroke=YEL, radius=0.05)
label(s, 443, 146, 394, 26, "receiver's shared memory", size=14, color=YEL)
for i in range(4):
    box(s, 463 + i * 92, 185, 76, 56, GREY, stroke=EDGE, body="32 KiB", size=13)
arrow(s, 313, 205, 433, 205, YEL, 2.5)
arrow(s, 847, 205, 967, 205, YEL, 2.5)
label(s, 313, 215, 120, 50, "copy 1\n(sender)", size=14, color=YEL)
label(s, 847, 215, 120, 50, "copy 2\n(receiver)", size=14, color=YEL)
box(s, 43, 330, 1194, 130, CHIP, stroke=EDGE, radius=0.05)
label(s, 63, 345, 1154, 30, "per piece: sender copies in, receiver copies out, buffer reused", size=16, align="left")
label(s, 63, 390, 1154, 60, "repeated until the whole message is across\n(2 copies per piece, versus 1 copy with CUDA IPC)", size=15, color=MUTED, align="left")

# 14. what the EFA counter measures, and where it was read
s = d.slide().s
title(s, "What was measured, and where", "A hardware counter of the EFA device, read through sysfs from a probe Pod on the same node")
box(s, 43, 110, 1194, 70, YEL_DK, stroke=YEL, radius=0.15)
sh.textbox(s, t, 63, 118, 1154, 32, "/sys/class/infiniband/<device>/ports/1/hw_counters/rdma_write_bytes",
           size=21, color=WHITE, bold=True, align="center", anchor="middle", mono=True)
label(s, 63, 150, 1154, 26, "reading this file returns the device's counter at that moment (fetched again if the last fetch is older than about 10 ms)", size=13, color=YEL)
box(s, 43, 200, 820, 450, CHIP, stroke=EDGE, radius=0.03)
label(s, 58, 206, 400, 26, "one node", size=14, color=MUTED, align="left")
chain = [(63, "probe Pod\n(hostNetwork)\nreads the file", CHIP, WHITE),
         (263, "kernel\nRDMA core\n(sysfs read)", LINK_DK, LINK),
         (463, "kernel\nefa driver\nGET_STATS", LINK_DK, LINK)]
for x, body, fill, stroke in chain:
    box(s, x, 250, 170, 110, fill, stroke=stroke, body=body, size=14)
box(s, 663, 250, 180, 380, YEL_DK, stroke=YEL, radius=0.06)
label(s, 673, 258, 160, 30, "EFA device", size=16, bold=True, color=YEL)
box(s, 673, 300, 160, 60, CHIP, stroke=YEL, body="RDMA write\nbyte counter", size=14)
for x in (233, 433, 633):
    arrow(s, x, 285, x + 30, 285, WHITE, 2)
    arrow(s, x + 30, 330, x, 330, YEL, 2)
label(s, 63, 366, 600, 26, "white: read request      yellow: counter value returned", size=13, color=MUTED, align="left")
box(s, 63, 430, 560, 200, GRN_DK, stroke=GRN, radius=0.06)
label(s, 73, 436, 540, 26, "engine Pod in the measured setup (no hostNetwork)", size=15, bold=True)
box(s, 93, 472, 230, 64, GREY, stroke=EDGE, body="prefill\nTP0  TP1", size=15)
box(s, 363, 472, 230, 64, GREY, stroke=EDGE, body="decode\nTP0  TP1", size=15)
label(s, 73, 552, 540, 50, "hw_counters is not visible inside this Pod", size=14, color=RED)
arrow(s, 623, 490, 663, 490, LINK, 3)
arrow(s, 663, 560, 623, 560, LINK, 3)
label(s, 673, 470, 165, 110, "KV cache\nas RDMA writes\n(shm off)", size=13, color=LINK)
box(s, 893, 200, 344, 450, CHIP, stroke=EDGE, radius=0.03)
label(s, 903, 208, 324, 26, "rdma_write_bytes, 17 requests", size=15, bold=True)
box(s, 913, 245, 304, 80, GREY, stroke=EDGE, radius=0.06)
label(s, 923, 250, 284, 26, "shm on (default)", size=14, color=MUTED)
label(s, 923, 278, 284, 36, "measured: 0", size=18, bold=True)
box(s, 913, 340, 304, 210, GREY, stroke=YEL, radius=0.06)
label(s, 923, 346, 284, 26, "shm off", size=14, color=YEL)
label(s, 923, 374, 284, 50, "measured\n8,663,389,312", size=17, bold=True)
label(s, 923, 432, 284, 50, "computed KV volume\n8,663,334,912", size=15, color=MUTED)
label(s, 923, 494, 284, 40, "ratio 1.0000063", size=17, bold=True, color=YEL)
label(s, 903, 565, 324, 70, "computed: 31,104 B/token\nx 8,192 tokens x 2 ranks x 17", size=13, color=MUTED)

d.save(OUT)
sys.path.insert(0, "/Users/akazawt/works/data-science/claudecode/.claude/skills/review-html/packs/zenn-doc/templates")
import strip_branding  # noqa: E402
strip_branding.main(OUT, OUT)
print(OUT)
