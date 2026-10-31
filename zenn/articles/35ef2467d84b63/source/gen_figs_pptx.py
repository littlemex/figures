"""gen_figs_pptx.py: the structure figures of the Engram hands-on article, as native pptx shapes.

    /usr/bin/python3 gen_figs_pptx.py <out.pptx>

Slides, in order: engram-why, overview, compress, hash, gate, placement, training-step, comparison.
Roles (ZN-25, ZN-29), each with a colour and a faint pattern so that it reads without colour:
  blue + "x" grid   the GPT backbone (embedding, blocks, output head)
  green + "/"       Engram (the table, key and value, the gate)
  yellow + dotted   computed from the token ids alone (compressed ids, n-grams, row ids)
  grey, no pattern  data and values (text, hidden state)
The numbers of the trained model come from zenmai-lab 7d12d41 (engram.py, model.py) and the published
trace traces/nanogpt-hastings.json (seed 1337).
"""
import sys

sys.path.insert(0, "/Users/akazawt/works/data-science/tools/pptxkit")
TEMPLATES = "/Users/akazawt/works/data-science/claudecode/.claude/skills/review-html/packs/zenn-doc/templates"
sys.path.insert(0, TEMPLATES)
from pptx.dml.color import RGBColor  # noqa: E402
from pptx.enum.dml import MSO_LINE_DASH_STYLE  # noqa: E402
from pptx.enum.shapes import MSO_SHAPE  # noqa: E402

from pptxkit import Deck, shapes as sh  # noqa: E402
from patterns import pattern_fill  # noqa: E402

TPL = "/Users/akazawt/works/data-science/tools/pptxkit/templates/aws-dark-2024.pptx"
OUT = sys.argv[1] if len(sys.argv) > 1 else "figures.pptx"

GRN, GRN_DK = RGBColor(36, 160, 70), RGBColor(18, 64, 34)
BLU, BLU_DK = RGBColor(40, 150, 240), RGBColor(16, 44, 72)
YEL, YEL_DK = RGBColor(251, 211, 50), RGBColor(80, 66, 14)
GREY, EDGE, CHIP = RGBColor(52, 56, 64), RGBColor(110, 116, 124), RGBColor(22, 26, 30)
MUTED, WHITE = RGBColor(160, 160, 168), RGBColor(245, 245, 248)
ROLE = {"backbone": (BLU_DK, BLU, "cross"), "engram": (GRN_DK, GRN, "diagonal"),
        "ids": (YEL_DK, YEL, "grid"), "data": (GREY, EDGE, None)}

d = Deck(TPL, check_layout=False)
t = d.theme


def label(s, x, y, w, h, body, size=16, color=WHITE, bold=False, align="center"):
    return sh.textbox(s, t, x, y, w, h, body, size=size, color=color, bold=bold, align=align, anchor="middle")


def box(s, x, y, w, h, role, body="", size=15, radius=0.1, dash=None, sw=1.5, color=WHITE):
    fill, stroke, pat = ROLE[role]
    sp = sh.rect(s, t, x, y, w, h, fill=fill, stroke=stroke, sw=sw, dash=dash,
                 shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=radius)
    pattern_fill(sp, pat, fill, stroke)
    if body:
        sh._write(sp.text_frame, t, body, size, color, True, "center", "middle")
        tf = sp.text_frame
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = sh.U(2)
    return sp


def frame(s, x, y, w, h):
    return sh.rect(s, t, x, y, w, h, fill=CHIP, stroke=EDGE, sw=1, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.03)


def title(s, head, roles):
    label(s, 43, 22, 1190, 46, head, size=30, bold=True, align="left")
    names = {"backbone": "GPT backbone", "engram": "Engram", "ids": "from token ids only", "data": "data"}
    x = 43
    for r in roles:
        box(s, x, 76, 34, 22, r)
        label(s, x + 42, 74, 220, 26, names[r], size=15, color=MUTED, align="left")
        x += 74 + 10 * len(names[r])


def arrow(s, x1, y1, x2, y2, color=WHITE, width=2):
    sh.arrow(s, t, x1, y1, x2, y2, color=color, width=width)


def seg(s, pts, color=WHITE, width=2):
    """A polyline with one arrowhead, at the last point."""
    from pptx.enum.shapes import MSO_CONNECTOR
    from pptx.util import Pt
    for (x1, y1), (x2, y2) in zip(pts[:-2], pts[1:-1]):
        c = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, sh.U(x1), sh.U(y1), sh.U(x2), sh.U(y2))
        c.line.color.rgb = color
        c.line.width = Pt(width)
    (x1, y1), (x2, y2) = pts[-2], pts[-1]
    arrow(s, x1, y1, x2, y2, color=color, width=width)


def num(s, x, y, n):
    c = sh.rect(s, t, x, y, 30, 30, fill=WHITE, stroke=WHITE, sw=1, shape=MSO_SHAPE.OVAL)
    sh._write(c.text_frame, t, str(n), 15, RGBColor(10, 10, 12), True, "center", "middle")
    tf = c.text_frame
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0


# ---------------------------------------------------------------- 1. engram-why (as-is / to-be)
s = d.slide().s
title(s, "What Engram may free up (the paper's view)", ["data", "backbone", "engram"])
for x, head_ in [(43, "Without Engram"), (653, "With Engram")]:
    label(s, x, 118, 584, 30, head_, size=19, bold=True)
    frame(s, x, 156, 584, 446)
ys = [252, 300, 348, 396, 444, 492]
for x0, n_build in [(83, 3), (693, 2)]:
    lw = 504 if x0 == 83 else 330
    lx = x0 if x0 == 83 else x0 + 174
    box(s, x0, 176, 504, 44, "data", body="tokens  \"Alexander the Great\"", size=15)
    arrow(s, lx + lw // 2, 220, lx + lw // 2, ys[0])
    for i, y in enumerate(ys):
        build = i < n_build
        box(s, lx, y, lw, 40, "backbone", body="build the name" if build else "read the context",
            size=14, dash=MSO_LINE_DASH_STYLE.DASH if build else None)
box(s, 693, 300, 150, 40, "engram", body="table lookup", size=14)
arrow(s, 768, 220, 768, 300, color=GRN, width=2)
arrow(s, 843, 320, 867, 320, color=GRN, width=3)
label(s, 83, 548, 504, 30, "6 blocks", size=14, color=MUTED)
label(s, 867, 548, 330, 30, "same 6 blocks", size=14, color=MUTED)
label(s, 43, 612, 1194, 30, "illustrative", size=14, color=MUTED)

# ---------------------------------------------------------------- 2. overview of one Engram layer
s = d.slide().s
title(s, "Inside one Engram layer (this article's model)", ["data", "ids", "engram", "backbone"])
top = [(1, "compress ids", "50,257 → 32,962", "ids"), (2, "n-grams", "2-gram, 3-gram", "ids"),
       (3, "hash", "8 + 8 = 16 row ids", "ids"), (4, "read rows", "16 × 64 = 1,024", "engram")]
for i, (n, head_, sub, role) in enumerate(top):
    x = 60 + i * 300
    box(s, x, 150, 250, 86, role, body=f"{head_}\n{sub}", size=16)
    num(s, x - 12, 138, n)
    if i < 3:
        arrow(s, x + 250, 193, x + 300, 193)
box(s, 60, 290, 250, 60, "data", body="token ids", size=16)
arrow(s, 185, 290, 185, 236)
box(s, 960, 290, 250, 70, "engram", body="table\n1,049,422 rows × 64", size=15)
arrow(s, 1085, 290, 1085, 236, color=GRN)
low = [(5, "key, value", "512 + 512", 960), (6, "gate α", "from h and key", 660), (7, "add", "h + α · value", 360)]
for n, head_, sub, x in low:
    box(s, x, 420, 250, 86, "engram", body=f"{head_}\n{sub}", size=16)
    num(s, x - 12, 408, n)
seg(s, [(1210, 193), (1245, 193), (1245, 463), (1210, 463)], color=GRN)
arrow(s, 960, 463, 910, 463)
arrow(s, 660, 463, 610, 463)
box(s, 60, 430, 250, 66, "backbone", body="block 1", size=16)
arrow(s, 360, 463, 310, 463)
box(s, 60, 570, 250, 60, "backbone", body="block 0", size=16)
box(s, 660, 570, 250, 60, "data", body="hidden state h (512)", size=15)
arrow(s, 310, 600, 660, 600)
arrow(s, 785, 570, 785, 506)
seg(s, [(700, 570), (700, 540), (485, 540), (485, 506)])
label(s, 60, 650, 1150, 30, "1-3: token ids only, before block 0", size=15, color=MUTED)

# ---------------------------------------------------------------- 3. compress the token ids
s = d.slide().s
title(s, "Step 1: one id for case and leading-space variants", ["data", "ids"])
toks = [("\"␣The\"", 383), ("\"␣the\"", 262), ("\"The\"", 464), ("\"the\"", 1169)]
for r, (w, i) in enumerate(toks):
    y = 160 + r * 92
    box(s, 80, y, 340, 64, "data", body=f"{w}   token {i}", size=18)
    seg(s, [(420, y + 32), (560, y + 32), (560, 330), (640, 330)])
box(s, 640, 298, 320, 64, "ids", body="compressed id 231", size=19)
label(s, 80, 560, 1100, 30, "␣ = leading space.  50,257 tokens → 32,962 ids", size=16, color=MUTED, align="left")

# ---------------------------------------------------------------- 4. hash: n-grams to row ids
s = d.slide().s
title(s, "Steps 2-4: the n-grams ending at \" the\" pick 16 rows", ["data", "ids", "engram"])
words = [("Only", 556), ("Alexander", 7565), ("the", 231), ("Great", None)]
for i, (w, c) in enumerate(words):
    x = 80 + i * 160
    if c is None:
        box(s, x, 140, 146, 54, "data", body="next: ?", size=17, dash=MSO_LINE_DASH_STYLE.DASH)
        continue
    box(s, x, 140, 146, 54, "data", body=w, size=17, sw=3 if i == 2 else 1.5)
    label(s, x, 196, 146, 24, f"id {c}", size=13, color=YEL)
label(s, 80 + 2 * 160, 112, 146, 26, "t", size=17, bold=True)
box(s, 80, 236, 466, 50, "ids", body="3-gram  Only Alexander the", size=16)
box(s, 80 + 1 * 160, 304, 306, 50, "ids", body="2-gram  Alexander the", size=16)
label(s, 80, 380, 600, 30, "XOR of (id × odd multiplier) → one number", size=15, color=MUTED, align="left")
frame(s, 760, 150, 470, 450)
label(s, 760, 552, 470, 28, "16 ranges, prime sizes 65,537 ... 65,651", size=14, bold=True)
box(s, 780, 176, 210, 30, "ids", body="2-gram: 8 heads", size=13, radius=0.2, sw=1)
box(s, 1000, 176, 215, 30, "ids", body="3-gram: 8 heads", size=13, radius=0.2, sw=1)
rows2 = [28481, 44042, 30922, 61096, 35305, 12384, 34790, 25524]
rows3 = [50523, 43514, 48875, 13539, 35234, 52272, 1325, 64750]
for k in range(8):
    y = 214 + k * 26
    box(s, 780, y, 210, 22, "engram", body=f"h{k}: row {rows2[k]:,}", size=11, radius=0.2, sw=1)
    box(s, 1000, y, 215, 22, "engram", body=f"h{k}: row {rows3[k]:,}", size=11, radius=0.2, sw=1)
seg(s, [(546, 261), (720, 261), (720, 118), (1107, 118), (1107, 176)], color=YEL)
seg(s, [(546, 329), (740, 329), (740, 191), (780, 191)], color=YEL)
label(s, 770, 428, 450, 30, "each range: its own prime", size=14, color=MUTED)
box(s, 800, 476, 390, 56, "engram", body="4. join 16 rows: e_t, 1,024 numbers", size=15)
arrow(s, 995, 458, 995, 476, color=GRN)
label(s, 80, 612, 600, 30, "seed 1337, row within each range", size=14, color=MUTED, align="left")

# ---------------------------------------------------------------- 5. key, value and the gate
s = d.slide().s
title(s, "Steps 5-7: how much of the lookup to add", ["data", "engram", "backbone"])
box(s, 60, 140, 200, 64, "engram", body="e_t\n1,024 numbers", size=16)
box(s, 300, 140, 220, 64, "engram", body="W_K, W_V\n1,024 → 512 + 512", size=15)
arrow(s, 260, 172, 300, 172)
box(s, 580, 100, 220, 50, "engram", body="value v_t (512)", size=16)
box(s, 580, 200, 220, 50, "engram", body="key k_t (512)", size=16)
arrow(s, 520, 160, 580, 125)
arrow(s, 520, 184, 580, 225)
box(s, 60, 330, 200, 64, "data", body="h_t from block 0\n512 numbers", size=15)
box(s, 580, 330, 300, 64, "engram", body="score s_t", size=17)
arrow(s, 690, 250, 690, 330)
arrow(s, 260, 362, 580, 362)
box(s, 580, 440, 300, 56, "engram", body="gate α_t (0 to 1)", size=17)
arrow(s, 730, 394, 730, 440)
box(s, 580, 560, 300, 64, "engram", body="h_t + α_t · v_t", size=19)
arrow(s, 730, 496, 730, 560)
seg(s, [(800, 125), (1180, 125), (1180, 592), (880, 592)])
seg(s, [(160, 394), (160, 592), (580, 592)])
box(s, 960, 650, 260, 50, "backbone", body="block 1", size=16)
seg(s, [(730, 624), (730, 675), (960, 675)])
label(s, 200, 450, 360, 40, "at \" the\": α = 0.88", size=18, bold=True, align="left")

# ---------------------------------------------------------------- 6. where it sits in the GPT
s = d.slide().s
title(s, "Where Engram sits in the 12-block GPT", ["data", "ids", "engram", "backbone"])
box(s, 80, 140, 240, 50, "data", body="token ids", size=16)
stack = ["embedding", "block 0", "Engram", "block 1", "block 2", "...", "block 11", "output head"]
y, prev = 220, 190
for name in stack:
    if name == "...":
        label(s, 80, y, 240, 26, "...", size=16, color=MUTED)
        y += 30
        prev = None
        continue
    if prev is not None:
        arrow(s, 200, prev, 200, y)
    box(s, 80, y, 240, 40, "engram" if name == "Engram" else "backbone", body=name, size=15)
    prev = y + 40
    y += 56
box(s, 520, 140, 300, 50, "ids", body="hash: 16 row ids", size=15)
arrow(s, 320, 165, 520, 165, color=YEL)
seg(s, [(670, 190), (670, 352), (320, 352)], color=YEL)
label(s, 700, 240, 300, 30, "before block 0", size=16, color=MUTED, align="left")
label(s, 340, 380, 400, 30, "before block 1 (2nd block)", size=16, color=MUTED, align="left")

# ---------------------------------------------------------------- 7. one training step
s = d.slide().s
title(s, "One training step", [])
row = [("windows", "128 × 512\ntokens"), ("model", "next-token\nprobabilities"), ("loss", "vs. the true\nnext token"),
       ("gradient", "for each\nweight"), ("update", "all weights")]
for i, (head_, body) in enumerate(row):
    x = 43 + i * 245
    label(s, x, 170, 210, 30, head_, size=18, bold=True)
    box(s, x, 210, 210, 110, "data", body=body, size=16)
    if i < 4:
        arrow(s, x + 210, 265, x + 245, 265)
seg(s, [(43 + 4 * 245 + 105, 320), (43 + 4 * 245 + 105, 400), (43 + 245 + 105, 400), (43 + 245 + 105, 320)])
label(s, 400, 410, 600, 30, "× 3,600 steps", size=16, color=MUTED)

# ---------------------------------------------------------------- 8. how the two runs are compared (as-is / to-be)
s = d.slide().s
title(s, "Same backbone, with and without Engram", ["backbone", "engram"])
for x, head_, extra in [(43, "Trained without Engram", False), (653, "With Engram", True)]:
    label(s, x, 116, 584, 30, head_, size=19, bold=True)
    frame(s, x, 154, 584, 430)
    for i, text in enumerate(["same starting weights", "same batches, same order", "236M tokens, AdamW"]):
        box(s, x + 30, 178 + i * 80, 524, 56, "backbone", body=text, size=16)
    if extra:
        box(s, x + 30, 418, 524, 56, "engram", body="+ Engram: 68.2M weights (table 67.2M)", size=16)
        box(s, x + 30, 498, 524, 56, "engram", body="table: Adam, 5 × backbone LR, no decay", size=16)
    label(s, x, 596, 584, 30, "about 132.0M weights" if extra else "about 63.8M weights", size=17, bold=True)

d.save(OUT)
import strip_branding  # noqa: E402

strip_branding.main(OUT, OUT)
print(OUT)
