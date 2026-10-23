"""gen_part1_pptx.py: 「Engram を自分で作って学習させる」の静止画 4 枚（モデルの全体、比べ方、作る順番、学習の 1 歩）。"""
import sys

sys.path.insert(0, "/Users/akazawt/works/data-science/tools/pptxkit")
from pptx.dml.color import RGBColor  # noqa: E402
from pptx.enum.shapes import MSO_SHAPE  # noqa: E402

from pptxkit import Deck, shapes as sh  # noqa: E402

TPL = "/Users/akazawt/works/data-science/tools/pptxkit/templates/aws-dark-2024.pptx"
OUT = sys.argv[1] if len(sys.argv) > 1 else "part1-figures.pptx"

GRN, GRN_DK = RGBColor(36, 160, 70), RGBColor(18, 64, 34)
CHIP, EDGE, GREY = RGBColor(22, 26, 30), RGBColor(70, 74, 80), RGBColor(52, 56, 64)
LINK, LINK_DK = RGBColor(40, 150, 240), RGBColor(16, 44, 72)
YEL, YEL_DK = RGBColor(251, 211, 50), RGBColor(80, 66, 14)
PNK, PNK_DK = RGBColor(255, 120, 160), RGBColor(80, 30, 46)
MUTED, WHITE = RGBColor(150, 150, 158), RGBColor(245, 245, 248)

d = Deck(TPL, check_layout=False)
t = d.theme


def label(s, x, y, w, h, body, size=16, color=WHITE, bold=False, align="center"):
    return sh.textbox(s, t, x, y, w, h, body, size=size, color=color, bold=bold, align=align, anchor="middle")


def box(s, x, y, w, h, fill, stroke=None, body="", size=15, color=WHITE, radius=0.1, sw=1.25):
    sp = sh.rect(s, t, x, y, w, h, fill=fill, stroke=stroke, sw=sw, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=radius)
    if body:
        sh._write(sp.text_frame, t, body, size, color, True, "center", "middle")
        tf = sp.text_frame
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = sh.U(2)
    return sp


def title(s, head, sub):
    label(s, 43, 26, 1190, 46, head, size=30, bold=True, align="left")
    label(s, 43, 72, 1190, 26, sub, size=16, color=MUTED, align="left")


def arrow(s, x1, y1, x2, y2, color=WHITE, width=2):
    sh.arrow(s, t, x1, y1, x2, y2, color=color, width=width)


# ---------------------------------------------------------------- 1. the model
s = d.slide().s
title(s, "A 12-block GPT with one Engram layer",
      "Grey: tokens.  Green: GPU compute.  Pink: row ids from the hash.  Yellow: the Engram table")
y = 150
box(s, 60, y, 200, 56, GREY, stroke=EDGE, body="token ids", size=15)
arrow(s, 260, y + 28, 300, y + 28)
box(s, 300, y, 180, 56, GRN_DK, stroke=GRN, body="embedding", size=15)
arrow(s, 480, y + 28, 520, y + 28)
box(s, 520, y, 150, 56, GRN_DK, stroke=GRN, body="block 0", size=15)
arrow(s, 670, y + 28, 710, y + 28)
box(s, 710, y, 170, 56, YEL_DK, stroke=YEL, body="+ Engram", size=15)
arrow(s, 880, y + 28, 920, y + 28)
box(s, 920, y, 150, 56, GRN_DK, stroke=GRN, body="block 1 ... 11", size=15)
arrow(s, 1070, y + 28, 1110, y + 28)
box(s, 1110, y, 120, 56, GRN_DK, stroke=GRN, body="next-token\nprobabilities", size=13)
# the Engram detail
box(s, 60, 260, 1170, 340, CHIP, stroke=EDGE, radius=0.03)
label(s, 80, 270, 900, 30, "Engram detail (after block 0, before the Attention of block 1)", size=16, bold=True, align="left")
box(s, 90, 320, 230, 70, GREY, stroke=EDGE, body="normalized ids\n50,257 -> 32,962", size=14)
arrow(s, 320, 355, 360, 355)
box(s, 360, 320, 270, 70, PNK_DK, stroke=PNK, body="2-gram, 3-gram\n× 8 hashes = 16 row ids", size=14)
arrow(s, 630, 355, 670, 355)
box(s, 670, 310, 250, 90, YEL_DK, stroke=YEL, body="table\n1,049,422 rows × 64", size=15)
arrow(s, 920, 355, 960, 355)
box(s, 960, 320, 240, 70, GRN_DK, stroke=GRN, body="e_t -> key, value\nW_K, W_V", size=14)
arrow(s, 1080, 390, 1080, 440)
box(s, 860, 440, 340, 70, GRN_DK, stroke=GRN, body="gate α from h, key\nout = h + α · value", size=14)
box(s, 90, 440, 520, 70, GREY, stroke=EDGE, body="h: hidden state after block 0", size=14)
arrow(s, 610, 475, 860, 475)

# ---------------------------------------------------------------- 2. how they are compared
s = d.slide().s
title(s, "Same base model, with and without Engram",
      "Green: shared.  Yellow: only in the Engram run.  One GPU per run, three seeds")
for x, head_, extra in [(43, "Engram run", True), (653, "baseline run", False)]:
    label(s, x, 112, 584, 30, head_, size=18, bold=True)
    box(s, x, 150, 584, 430, CHIP, stroke=EDGE, radius=0.04)
    rows = [("same backbone weights before training", True), ("same batches, same order", True),
            ("236M tokens, AdamW", True)]
    for i, (text, _) in enumerate(rows):
        box(s, x + 30, 175 + i * 80, 524, 56, GRN_DK, stroke=GRN, body=text, size=15)
    if extra:
        box(s, x + 30, 415, 524, 56, YEL_DK, stroke=YEL, body="+ Engram 68.2M (table 67.2M)", size=15)
        box(s, x + 30, 495, 524, 56, YEL_DK, stroke=YEL, body="table: Adam, 5× LR, no decay", size=15)
    label(s, x, 592, 584, 30, "total about 132.0M weights" if extra else "total about 63.8M weights", size=16, bold=True)

# ---------------------------------------------------------------- 3. the build steps
s = d.slide().s
title(s, "Four steps from text to a comparison",
      "Grey: data.  Green: what we build and run.  Yellow: the Engram part")
steps = [("1. data", "WikiText-103\n-> GPT-2 token ids", GREY, EDGE),
         ("2. model", "12-block GPT backbone\n+ Engram after block 0", YEL_DK, YEL),
         ("3. train", "backbone alone and\nbackbone + Engram", GRN_DK, GRN),
         ("4. evaluate", "text not used in training\nperplexity, Engram off, LogitLens", GRN_DK, GRN)]
for i, (head_, body, f, st) in enumerate(steps):
    x = 43 + i * 302
    label(s, x, 150, 270, 30, head_, size=18, bold=True)
    box(s, x, 190, 270, 120, f, stroke=st, body=body, size=15)
    if i < 3:
        arrow(s, x + 270, 250, x + 302, 250)
label(s, 43, 360, 1194, 30, "three seeds for each run; one GPU per run", size=15, color=MUTED)

# ---------------------------------------------------------------- 4. one training step
s = d.slide().s
title(s, "One training step",
      "Grey: data.  Green: GPU compute.  Pink: computed from the prediction.  Yellow: what the step changes")
row = [("windows", "128 windows of\n512 tokens", GREY, EDGE), ("model", "next-token\nprobabilities", GRN_DK, GRN),
       ("loss", "how far from\nthe true next token", PNK_DK, PNK), ("gradient", "how the loss changes\nwith each weight", PNK_DK, PNK),
       ("update", "AdamW or Adam\n(learning rate)", YEL_DK, YEL)]
for i, (head_, body, f, st) in enumerate(row):
    x = 43 + i * 242
    box(s, x, 180, 210, 120, f, stroke=st, body=f"{head_}\n{body}", size=15)
    if i < 4:
        arrow(s, x + 210, 240, x + 242, 240)
from pptx.enum.shapes import MSO_CONNECTOR  # noqa: E402
from pptx.util import Pt  # noqa: E402
for x1, y1, x2, y2 in [(1116, 300, 1116, 380), (1116, 380, 148, 380)]:
    c = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, sh.U(x1), sh.U(y1), sh.U(x2), sh.U(y2))
    c.line.color.rgb = MUTED
    c.line.width = Pt(2)
arrow(s, 148, 380, 148, 300, color=MUTED)
label(s, 300, 392, 680, 30, "repeat: 3,600 steps, 65,536 tokens per step, 236M tokens", size=15, color=MUTED)
d.save(OUT)
sys.path.insert(0, "/Users/akazawt/works/data-science/claudecode/.claude/skills/review-html/packs/zenn-doc/templates")
import strip_branding  # noqa: E402
strip_branding.main(OUT, OUT)
print(OUT)
