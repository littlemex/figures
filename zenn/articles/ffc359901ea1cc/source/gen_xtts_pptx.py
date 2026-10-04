"""gen_xtts_pptx.py — XTTS v2 を NxD Inference で動かす記事の図を、ネイティブ shape の pptx 2 枚で組む。

    /usr/bin/python3 gen_xtts_pptx.py <out.pptx>

1. XTTS v2 の構造と CPU / Neuron の分担、1081 位置の系列の内訳
2. ブリッジ (NeuronGPT2InferenceModel) が Coqui TTS / Hugging Face と NxD Inference の間に入る位置
図の文字は英語。根拠は記事本文と littlemex/samples の xttsv2-nxd-inference (commit b6334ca)。
"""
import sys

sys.path.insert(0, "/Users/akazawt/works/data-science/tools/pptxkit")
from pptx.dml.color import RGBColor  # noqa: E402
from pptx.enum.dml import MSO_LINE_DASH_STYLE  # noqa: E402
from pptx.enum.shapes import MSO_SHAPE  # noqa: E402

from pptxkit import Deck, shapes as sh  # noqa: E402

TPL = "/Users/akazawt/works/data-science/tools/pptxkit/templates/aws-dark-2024.pptx"
OUT = sys.argv[1] if len(sys.argv) > 1 else "xtts-figures.pptx"

GRN = RGBColor(36, 160, 70)
GRN_HI = RGBColor(59, 209, 111)
GRN_DK = RGBColor(18, 64, 34)
CHIP = RGBColor(22, 26, 30)
EDGE = RGBColor(70, 74, 80)
GREY = RGBColor(52, 56, 64)
LINK = RGBColor(40, 150, 240)
LINK_DK = RGBColor(16, 44, 72)
YEL = RGBColor(251, 211, 50)
YEL_DK = RGBColor(80, 66, 14)
MUTED = RGBColor(150, 150, 158)
WHITE = RGBColor(245, 245, 248)

d = Deck(TPL, check_layout=False)
t = d.theme


def label(s, x, y, w, h, body, size=16, color=WHITE, bold=False, align="center"):
    return sh.textbox(s, t, x, y, w, h, body, size=size, color=color, bold=bold, align=align,
                      anchor="middle")


def box(s, x, y, w, h, fill, stroke=None, body="", size=15, color=WHITE, radius=0.1,
        dash=None, sw=1.25):
    sp = sh.rect(s, t, x, y, w, h, fill=fill, stroke=stroke, sw=sw, dash=dash,
                 shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=radius)
    if body:
        sh._write(sp.text_frame, t, body, size, color, True, "center", "middle")
        tf = sp.text_frame
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = sh.U(2)
    return sp


def arrow(s, x1, y1, x2, y2, color=WHITE, w=2.0, head=False):
    sh.arrow(s, t, x1, y1, x2, y2, color=color, width=w, head=head)


def title(s, head, sub):
    label(s, 43, 26, 1190, 46, head, size=30, bold=True, align="left")
    label(s, 43, 72, 1190, 26, sub, size=16, color=MUTED, align="left")


CPU = (GREY, EDGE)
NEU = (GRN_DK, GRN_HI)

# ------------------------------------------------------------------ 1. XTTS v2 structure
s = d.slide()
title(s.s, "XTTS v2 at a glance", "Only the GPT decoder moves to Neuron")
# inputs
box(s.s, 43, 140, 170, 56, CHIP, stroke=EDGE, body="Reference audio", size=15)
box(s.s, 43, 300, 170, 56, CHIP, stroke=EDGE, body="Text", size=15)
box(s.s, 250, 120, 220, 44, *CPU, body="ConditioningEncoder", size=14)
box(s.s, 250, 176, 220, 44, *CPU, body="PerceiverResampler", size=14)
box(s.s, 250, 300, 220, 56, *CPU, body="Tokenizer + embeddings", size=14)
arrow(s.s, 213, 168, 250, 142)
arrow(s.s, 360, 164, 360, 176)
arrow(s.s, 213, 328, 250, 328)
label(s.s, 250, 226, 220, 22, "speaker prompt", size=13, color=MUTED)
# GPT
box(s.s, 520, 150, 330, 190, *NEU, body="GPT-2 decoder\n30 layers", size=22, sw=2)
label(s.s, 520, 346, 330, 24, "NeuronCores  (TP=2 on trn1.2xlarge)", size=13, color=GRN_HI)
arrow(s.s, 470, 198, 520, 210)
arrow(s.s, 470, 328, 520, 290)
# autoregressive loop: Neuron returns hidden states; CPU lm_head + HF sampling pick the token
sh.arrow(s.s, t, 850, 300, 880, 300, color=YEL, width=2)
box(s.s, 880, 278, 150, 44, *CPU, body="lm_head + sampling", size=12)
sh.arrow(s.s, t, 955, 278, 955, 196, color=YEL, width=2)
sh.arrow(s.s, t, 955, 196, 850, 196, color=YEL, width=2)
label(s.s, 862, 160, 190, 22, "next token embedding", size=12, color=YEL)
label(s.s, 862, 326, 190, 22, "repeat until EOS", size=12, color=YEL)
# decoder
box(s.s, 1060, 150, 177, 56, *CPU, body="HiFi-GAN decoder", size=14)
box(s.s, 1060, 250, 177, 56, CHIP, stroke=EDGE, body="Waveform", size=15)
arrow(s.s, 1030, 290, 1060, 180)
arrow(s.s, 1148, 206, 1148, 250)
# legend
box(s.s, 43, 392, 24, 18, *CPU)
label(s.s, 74, 388, 80, 26, "CPU", size=13, color=MUTED, align="left")
box(s.s, 150, 392, 24, 18, *NEU)
label(s.s, 181, 388, 140, 26, "NeuronCores", size=13, color=MUTED, align="left")

# sequence layout: 1081 static positions
SX, SW = 43, 1194
segs = [("<=70", 70, GREY), ("text <= 402 + BOS/EOS", 404, LINK_DK), ("audio <= 605 + BOS/EOS", 607, YEL_DK)]
label(s.s, SX, 440, 700, 28, "Static capacity: up to 1081 positions", size=17, bold=True, align="left")
x = SX
for name, n, col in segs:
    w = SW * n / 1081
    box(s.s, x, 476, w - 4, 46, col, stroke=EDGE, body=name, size=12 if n < 100 else 14, radius=0.06)
    x += w
pre_w = SW * (70 + 404) / 1081
box(s.s, SX, 532, pre_w - 4, 34, None, stroke=GRN_HI, dash=MSO_LINE_DASH_STYLE.DASH,
    body="prefill: once, all at once", size=13, color=GRN_HI, radius=0.2)
box(s.s, SX + pre_w, 532, SW - pre_w - 4, 34, None, stroke=YEL, dash=MSO_LINE_DASH_STYLE.DASH,
    body="decode: one token per step", size=13, color=YEL, radius=0.2)

# ------------------------------------------------------------------ 2. the bridge
s = d.slide()
title(s.s, "One new class bridges Coqui TTS and NxD Inference", "Swap gpt_inference; nothing above it changes")
L = 43
rows = [
    (130, "Coqui TTS   synthesize()", "unchanged", CHIP, EDGE),
    (220, "Hugging Face   generate()", "unchanged", CHIP, EDGE),
]
for y, name, tag, f, st in rows:
    gx0 = L if y == 130 else L + 100
    box(s.s, gx0, y, 760 - (gx0 - L), 62, f, stroke=st, body=name, size=19)
    box(s.s, L + 790, y + 14, 160, 34, GREY, stroke=EDGE, body=tag, size=14, radius=0.3)
arrow(s.s, L + 380, 192, L + 380, 220)
sh.arrow(s.s, t, L + 40, 192, L + 40, 316, color=GRN_HI, width=2)
label(s.s, L + 42, 254, 58, 40, "before\ngenerate()", size=9, color=GRN_HI)
# bridge
box(s.s, L, 316, 760, 112, LINK_DK, stroke=LINK, sw=2.5, radius=0.06)
label(s.s, L + 16, 322, 728, 34, "NeuronGPT2InferenceModel  (bridge)", size=19, bold=True)
box(s.s, L + 24, 364, 220, 48, *CPU, body="embeddings", size=14)
box(s.s, L + 270, 364, 220, 48, YEL_DK, stroke=YEL, body="state (write position)", size=14)
box(s.s, L + 516, 364, 220, 48, *CPU, body="lm_head", size=14)
box(s.s, L + 790, 352, 160, 40, GRN_DK, stroke=GRN_HI, body="new", size=15, radius=0.3)
sh.arrow(s.s, t, L + 300, 282, L + 300, 316, color=WHITE, width=2)
sh.arrow(s.s, t, L + 460, 316, L + 460, 282, color=WHITE, width=2)
label(s.s, L + 90, 284, 200, 30, "forward()", size=13, color=MUTED)
label(s.s, L + 470, 284, 220, 30, "logits, dummy KV", size=13, color=MUTED)
# NxD side
box(s.s, L, 466, 760, 150, CHIP, stroke=GRN, radius=0.04)
label(s.s, L + 200, 470, 360, 30, "NxD Inference application", size=15, color=GRN_HI)
box(s.s, L + 24, 504, 300, 46, GRN_DK, stroke=GRN, body="prefill_app", size=16)
box(s.s, L + 436, 504, 300, 46, GRN_DK, stroke=GRN, body="decode_app", size=16)
sh.arrow(s.s, t, L + 324, 527, L + 436, 527, color=YEL, width=2)
label(s.s, L + 324, 500, 112, 24, "KV copy", size=12, color=YEL)
box(s.s, L + 24, 560, 712, 44, GRN, stroke=GRN_HI, body="NeuronGPTTransformer on NeuronCores  (KV cache on device)",
    size=15)
sh.arrow(s.s, t, L + 174, 428, L + 174, 504, color=GRN_HI, width=2)
sh.arrow(s.s, t, L + 586, 428, L + 586, 504, color=YEL, width=2)
label(s.s, L + 184, 432, 230, 26, "store_prefix_emb()", size=12, color=GRN_HI, align="left")
label(s.s, L + 596, 436, 160, 26, "forward(), every step", size=12, color=YEL, align="left")
box(s.s, L + 790, 520, 160, 34, GREY, stroke=EDGE, body="earlier article", size=13, radius=0.3)
# swap hook
box(s.s, 1010, 316, 227, 112, None, stroke=EDGE, dash=MSO_LINE_DASH_STYLE.DASH, radius=0.08,
    body="gpt.gpt_inference\n=\nbridge", size=15, color=MUTED)

d.save(OUT)
sys.path.insert(0, "/Users/akazawt/works/data-science/claudecode/.claude/skills/review-html/packs/zenn-doc/templates")
import strip_branding  # noqa: E402
strip_branding.main(OUT, OUT)
print(OUT)
