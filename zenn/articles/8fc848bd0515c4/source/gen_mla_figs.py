"""gen_mla_figs.py: figures for the Zenn article on SGLang's two procedures for DeepSeek-V2 MLA.

    /usr/bin/python3 gen_mla_figs.py <out.pptx>

Values: DeepSeek-V2-Lite-Chat config (16 heads, kv_lora_rank 512, rope 64, nope 128, v 128);
SGLang 5a15cde8 attention_backend_handler.py L172-178, schedule_policy.py L602/L922,
extend_attention.py L609; batch times are the article's measured values.
"""
import sys
sys.path.insert(0, "/Users/akazawt/works/data-science/tools/pptxkit")
from pptx.dml.color import RGBColor
from pptx.enum.dml import MSO_LINE_DASH_STYLE
from pptx.enum.shapes import MSO_SHAPE
from pptxkit import Deck, shapes as sh

TPL = "/Users/akazawt/works/data-science/tools/pptxkit/templates/aws-dark-2024.pptx"
OUT = sys.argv[1]
GRN = RGBColor(36, 160, 70); GRN_HI = RGBColor(59, 209, 111); GRN_DK = RGBColor(18, 64, 34)
BLU = RGBColor(40, 150, 240); BLU_DK = RGBColor(16, 48, 84)
ORG = RGBColor(240, 140, 40); ORG_DK = RGBColor(84, 46, 12)
RED = RGBColor(230, 70, 70); RED_DK = RGBColor(80, 20, 20)
GRY = RGBColor(52, 56, 64); EDGE = RGBColor(90, 94, 100); MUTED = RGBColor(160, 160, 168)
WHITE = RGBColor(245, 245, 248)
d = Deck(TPL, check_layout=False); t = d.theme

def label(s, x, y, w, h, body, size=18, color=WHITE, bold=False, align="center", anchor="middle"):
    return sh.textbox(s, t, x, y, w, h, body, size=size, color=color, bold=bold, align=align, anchor=anchor)

def box(s, x, y, w, h, fill, stroke=None, body="", size=18, color=WHITE, dash=None, sw=1.5):
    sp = sh.rect(s, t, x, y, w, h, fill=fill, stroke=stroke, sw=sw, dash=dash, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.08)
    if body:
        sh._write(sp.text_frame, t, body, size, color, True, "center", "middle")
        tf = sp.text_frame; tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = sh.U(2)
    return sp

def title(s, head, sub=""):
    label(s, 43, 22, 1190, 46, head, size=30, bold=True, align="left")
    if sub:
        label(s, 43, 68, 1190, 28, sub, size=17, color=MUTED, align="left")

def arrow(s, x1, y1, x2, y2, color=MUTED, w=2):
    sh.arrow(s, t, x1, y1, x2, y2, color=color, width=w)

# 1. what the KV cache keeps per token per layer (to scale)
s = d.slide(); title(s.s, "What the KV cache keeps per token, per layer", "DeepSeek-V2-Lite sizes; top row: a non-MLA layer with the same head sizes, for comparison; to scale")
x0, wmax = 330, 880
label(s.s, 43, 170, 280, 60, "Ordinary multi-head\nattention (K and V\nfor every head)", size=17, align="left")
kw = wmax * 3072 / 5120; vw = wmax * 2048 / 5120
box(s.s, x0, 175, kw, 56, BLU_DK, BLU, "K: 16 heads × (128 + 64) = 3,072", size=16)
box(s.s, x0 + kw, 175, vw, 56, ORG_DK, ORG, "V: 16 × 128 = 2,048", size=16)
label(s.s, x0, 236, wmax, 28, "5,120 values", size=18, bold=True)
label(s.s, 43, 330, 280, 60, "MLA (DeepSeek-V2)\nkeeps one compressed\nvector", size=17, align="left")
lw = wmax * 512 / 5120; rw = wmax * 64 / 5120
box(s.s, x0, 335, lw, 56, GRN_DK, GRN_HI, "latent 512", size=16)
box(s.s, x0 + lw, 335, rw, 56, GRY, EDGE, "", size=12)
label(s.s, x0 + lw + rw + 10, 335, 600, 56, "+ rope key 64  =  576 values", size=18, bold=True, align="left")
label(s.s, 43, 470, 1190, 80, "Both SGLang procedures below compute this same MLA layer.", size=18, color=MUTED, align="left")

# 2. two procedures for the same MLA layer
s = d.slide(); title(s.s, "Two procedures for the same MLA layer", "SGLang names them MHA and MLA (absorbed); mathematically the same attention")
box(s.s, 60, 300, 210, 90, GRN_DK, GRN_HI, "Cached latent\n512 + 64", size=18)
label(s.s, 360, 132, 700, 34, "MHA: decompress, then ordinary attention", size=20, bold=True, color=BLU, align="left")
box(s.s, 360, 180, 230, 70, GRY, EDGE, "Up-project\nfor every token", size=17)
box(s.s, 640, 180, 230, 70, BLU_DK, BLU, "K, V per head\n16 heads", size=17)
box(s.s, 920, 180, 260, 70, BLU_DK, BLU, "Multi-head\nattention", size=17)
arrow(s.s, 270, 320, 360, 215, BLU); arrow(s.s, 590, 215, 640, 215, BLU); arrow(s.s, 870, 215, 920, 215, BLU)
label(s.s, 360, 432, 700, 34, "MLA (absorbed): compute on the latent itself", size=20, bold=True, color=ORG, align="left")
box(s.s, 360, 480, 230, 70, GRY, EDGE, "Fold W_UK\ninto the query", size=17)
box(s.s, 640, 480, 230, 70, ORG_DK, ORG, "Attend over\nthe latent", size=17)
box(s.s, 920, 480, 260, 70, ORG_DK, ORG, "Fold W_UV\ninto the output", size=17)
arrow(s.s, 270, 370, 360, 515, ORG); arrow(s.s, 590, 515, 640, 515, ORG); arrow(s.s, 870, 515, 920, 515, ORG)
label(s.s, 360, 260, 820, 30, "cost grows with every cached token decompressed", size=16, color=MUTED, align="left")
label(s.s, 360, 560, 820, 30, "nothing decompressed; heavier per query token", size=16, color=MUTED, align="left")

# 3. prefill vs decode
s = d.slide(); title(s.s, "Prefill and decode have opposite shapes", "Which procedure is cheaper depends on new tokens versus cached tokens")
label(s.s, 60, 140, 560, 34, "Prefill, no cached prefix", size=21, bold=True, align="left")
box(s.s, 60, 190, 560, 60, BLU_DK, BLU, "8,192 new tokens", size=18)
label(s.s, 60, 260, 560, 30, "cached: 0", size=17, color=MUTED, align="left")
box(s.s, 60, 320, 560, 70, BLU_DK, BLU, "MHA measured faster here", size=20)
label(s.s, 680, 140, 560, 34, "Decode", size=21, bold=True, align="left")
box(s.s, 680, 190, 26, 60, ORG_DK, ORG, "", size=12)
label(s.s, 716, 190, 520, 60, "1 new token", size=18, align="left")
box(s.s, 680, 260, 560, 40, GRY, EDGE, "long cached history", size=17)
box(s.s, 680, 320, 560, 70, ORG_DK, ORG, "SGLang always uses absorbed MLA", size=20)
label(s.s, 60, 440, 1180, 120, "8,192-token prefill, prefix 0, triton: MHA ~365 ms, forced absorbed MLA ~671 ms (median)", size=19, align="left")

# 4. SGLang's rule
s = d.slide(); title(s.s, "SGLang's rule for the triton backend", "Ordinary case (no piecewise CUDA graph, no deterministic inference); decided once per batch")
box(s.s, 60, 280, 200, 80, GRY, EDGE, "A batch", size=20)
box(s.s, 330, 280, 260, 80, GRY, EDGE, "Extend, non-\nspeculative?", size=19)
box(s.s, 670, 160, 300, 80, GRY, EDGE, "Sum of cached\nprefix == 0 ?", size=19)
box(s.s, 1040, 160, 200, 80, BLU_DK, BLU, "MHA", size=24)
box(s.s, 1040, 400, 200, 90, ORG_DK, ORG, "absorbed\nMLA", size=20)
arrow(s.s, 260, 320, 330, 320); arrow(s.s, 590, 300, 670, 200); label(s.s, 560, 215, 60, 30, "yes", size=16, color=MUTED)
arrow(s.s, 970, 200, 1040, 200); label(s.s, 975, 160, 60, 30, "yes", size=16, color=MUTED)
arrow(s.s, 820, 240, 1060, 400); label(s.s, 960, 280, 60, 30, "no", size=16, color=MUTED)
arrow(s.s, 590, 340, 1040, 445); label(s.s, 620, 420, 280, 30, "no (decode, speculative)", size=16, color=MUTED)
label(s.s, 60, 560, 1180, 60, "One cached token in any request: whole batch to absorbed MLA", size=19, align="left")

# 5. page size as the reuse unit
s = d.slide(); title(s.s, "Page size is the unit of reuse", "A 5-token shared prefix, 8,192-token prompt")
label(s.s, 60, 140, 400, 34, "page size 1", size=21, bold=True, align="left")
for i in range(5):
    box(s.s, 60 + i * 38, 190, 34, 50, GRN_DK, GRN_HI, "", size=10)
box(s.s, 60 + 5 * 38, 190, 960, 50, BLU_DK, BLU, "8,187 new tokens", size=17)
label(s.s, 60, 250, 1180, 30, "5 tokens reused  →  summed prefix 5  →  absorbed MLA", size=19, color=ORG, align="left")
label(s.s, 60, 296, 400, 34, "page size 64", size=21, bold=True, align="left")
box(s.s, 60, 380, 1150, 50, BLU_DK, BLU, "8,192 new tokens", size=17)
box(s.s, 60, 372, 160, 66, None, ORG, "", size=12, dash=MSO_LINE_DASH_STYLE.DASH, sw=2)
label(s.s, 232, 338, 400, 30, "first page: only 5 of its 64 tokens match", size=15, color=ORG, align="left")
label(s.s, 60, 440, 1180, 30, "5 of 64 match: the page is not reused  →  summed prefix 0  →  MHA", size=19, color=BLU, align="left")
label(s.s, 60, 520, 1180, 60, "Cost: anything shorter than a page is never reused.", size=19, color=MUTED, align="left")

# 6. /health split
s = d.slide(); title(s.s, "A one-token /health request can split a prompt", "Budget 8,192 tokens per prefill batch, page size 64")
label(s.s, 60, 140, 300, 34, "batch n", size=20, bold=True, align="left")
box(s.s, 60, 185, 40, 60, RED_DK, RED, "1", size=15)
box(s.s, 100, 185, 1040, 60, BLU_DK, BLU, "prompt A: first 8,128 tokens (rounded down to a page)", size=17)
label(s.s, 60, 250, 1180, 30, "health uses 1 token  →  8,191 left  →  8,128  →  64 tokens wait", size=17, color=MUTED, align="left")
label(s.s, 60, 320, 400, 34, "a later batch", size=20, bold=True, align="left")
box(s.s, 60, 365, 150, 60, ORG_DK, ORG, "A: last 64", size=16)
box(s.s, 210, 365, 930, 60, BLU_DK, BLU, "prompt B: first 8,128 tokens", size=17)
label(s.s, 60, 430, 1180, 30, "A's 64 tokens have 8,128 cached tokens  →  summed prefix > 0  →  absorbed MLA for the whole batch", size=17, color=ORG, align="left")
label(s.s, 60, 500, 1180, 60, "B is cut to 8,128 the same way: the chain continues", size=19, align="left")

# 7. kernel waste
s = d.slide(); title(s.s, "Launched query blocks in the batch with A's 64 tokens", "triton extend attention, grid = (batch, heads, ceil(longest extend / 32)), measured version")
label(s.s, 60, 150, 260, 40, "A: 64 tokens", size=18, align="left")
box(s.s, 330, 150, 60, 40, ORG_DK, ORG, "2", size=14)
box(s.s, 390, 150, 840, 40, RED_DK, RED, "252 masked blocks, each still loops over 8,128 cached tokens", size=15, dash=MSO_LINE_DASH_STYLE.DASH)
label(s.s, 60, 210, 260, 40, "B: 8,128 tokens", size=18, align="left")
box(s.s, 330, 210, 900, 40, BLU_DK, BLU, "254 blocks with real work", size=15)
label(s.s, 60, 300, 1180, 40, "One-line early return for blocks past their own sequence:", size=19, align="left")
box(s.s, 60, 360, 1100 * 1533 / 1533, 50, RED_DK, RED, "measured version: ~1,533 ms per batch", size=16)
box(s.s, 60, 430, 1100 * 704 / 1533, 50, GRN_DK, GRN_HI, "with early return: ~704 ms", size=16)
label(s.s, 60, 510, 1180, 60, "Bars to scale. Fixed upstream in sgl-project/sglang#32109 (merged 2026-07-23).", size=17, color=MUTED, align="left")

d.save(OUT)
sys.path.insert(0, "/Users/akazawt/works/data-science/claudecode/.claude/skills/review-html/packs/zenn-doc/templates")
import strip_branding
strip_branding.main(OUT, OUT)
print(OUT)
