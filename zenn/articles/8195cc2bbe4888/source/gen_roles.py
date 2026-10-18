"""gen_roles.py -- NCCL4Py と nccl-extensions の各部品の受け持ちを 1 枚の pptx にする。

    /usr/bin/python3 gen_roles.py <out.pptx>

根拠は nccl4py v0.6.0 (893470119) と nccl-extensions v0.1.0 (e57f0da) のソース、NCCL 2.30.7 の配布物。
"""
import os
import sys

sys.path.insert(0, "/Users/akazawt/works/data-science/tools/pptxkit")
from pptx.dml.color import RGBColor  # noqa: E402
from pptx.enum.dml import MSO_LINE_DASH_STYLE  # noqa: E402
from pptx.enum.shapes import MSO_SHAPE  # noqa: E402

from pptxkit import Deck, shapes as sh  # noqa: E402

TPL = "/Users/akazawt/works/data-science/tools/pptxkit/templates/aws-dark-2024.pptx"
OUT = sys.argv[1] if len(sys.argv) > 1 else "roles.pptx"
TEMPLATES = "/Users/akazawt/works/data-science/claudecode/.claude/skills/review-html/packs/zenn-doc/templates"

GRN = RGBColor(36, 160, 70)
GRN_DK = RGBColor(18, 64, 34)
CHIP = RGBColor(22, 26, 30)
EDGE = RGBColor(70, 74, 80)
GREY = RGBColor(52, 56, 64)
LINK = RGBColor(40, 150, 240)
LINK_DK = RGBColor(16, 44, 72)
YEL = RGBColor(251, 211, 50)
YEL_DK = RGBColor(80, 66, 14)
MUTED = RGBColor(150, 150, 158)
WHITE = RGBColor(245, 245, 248)

d = Deck(TPL, check_layout=False)
t = d.theme


def label(s, x, y, w, h, body, size=16, color=WHITE, bold=False, align="center"):
    return sh.textbox(s, t, x, y, w, h, body, size=size, color=color, bold=bold, align=align, anchor="middle")


def box(s, x, y, w, h, fill, stroke=None, body="", size=15, color=WHITE, radius=0.1, dash=None):
    sp = sh.rect(s, t, x, y, w, h, fill=fill, stroke=stroke, sw=1.25, dash=dash,
                 shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=radius)
    if body:
        sh._write(sp.text_frame, t, body, size, color, True, "center", "middle")
        tf = sp.text_frame
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = sh.U(2)
    return sp


s = d.slide().s
label(s, 43, 26, 1190, 46, "What each piece does", size=30, bold=True, align="left")
label(s, 43, 72, 1190, 26, "Green: your code.  Blue: Python API and native libraries.  Yellow: bindings.  Gray: GPU driver.  Dashed: opened at run time",
      size=16, color=MUTED, align="left")

LX, RX, CW = 230, 790, 430
label(s, LX, 108, CW, 30, "nccl4py + nvidia-nccl", size=17, bold=True)
label(s, RX, 108, CW, 30, "nccl-extensions", size=17, bold=True)

rows = [("You", 146), ("Python API", 230), ("Bindings", 336), ("Native", 442), ("Driver", 560)]
for name, y in rows:
    label(s, 43, y, 170, 70, name, size=16, color=MUTED, align="left")

box(s, LX, 146, RX + CW - LX, 64, GRN_DK, stroke=GRN, body="your Python program", size=17)

box(s, LX, 230, CW, 86, LINK_DK, stroke=LINK, size=15,
    body="nccl.core\nCommunicator, Buffer, group, interop")
box(s, RX, 230, CW, 86, LINK_DK, stroke=LINK, size=15,
    body="nccl.ep\nGroup, Handle, Tensor, interop.torch")

box(s, LX, 336, CW, 86, YEL_DK, stroke=YEL, size=15,
    body="nccl.bindings\nnccl / cynccl / _internal")
box(s, RX, 336, CW, 86, YEL_DK, stroke=YEL, size=15,
    body="nccl._extensions\nnccl_ep / cynccl_ep / _internal, _runtime")

box(s, LX, 442, CW, 86, LINK_DK, stroke=LINK, size=15,
    body="libnccl.so.2\ncommunicators, collectives, device API")
box(s, RX, 442, CW, 86, LINK_DK, stroke=LINK, size=15,
    body="libnccl_ep.so\ndispatch, combine, JIT kernels")

box(s, LX, 560, RX + CW - LX, 56, GREY, stroke=EDGE, body="libcuda.so.1 (GPU driver)", size=16)

cx_l, cx_r = LX + CW / 2, RX + CW / 2
for x in (cx_l, cx_r):
    sh.arrow(s, t, x, 210, x, 230, color=WHITE, width=2)
    sh.arrow(s, t, x, 316, x, 336, color=WHITE, width=2)
    sh.arrow(s, t, x, 422, x, 442, color=MUTED, width=2, dashed=True)
sh.arrow(s, t, RX, 395, LX + CW, 395, color=YEL, width=2)
label(s, LX + CW, 397, RX - LX - CW, 40, "loads libnccl\nfirst", size=12, color=YEL)
sh.arrow(s, t, LX + CW, 283, RX, 283, color=YEL, width=2.5)
label(s, LX + CW, 248, RX - LX - CW, 30, "comm.ptr", size=14, color=YEL)
sh.arrow(s, t, RX, 495, LX + CW, 495, color=WHITE, width=2)
label(s, LX + CW, 460, RX - LX - CW, 30, "links", size=14, color=MUTED)
sh.arrow(s, t, cx_r, 528, cx_r, 560, color=WHITE, width=2)
sh.arrow(s, t, cx_l, 528, cx_l, 560, color=MUTED, width=2, dashed=True)
label(s, cx_l + 8, 530, 240, 28, "at run time", size=14, color=MUTED, align="left")
label(s, cx_r + 8, 530, 200, 28, "links", size=14, color=MUTED, align="left")

d.save(OUT)
sys.path.insert(0, TEMPLATES)
import strip_branding  # noqa: E402
strip_branding.main(OUT, OUT)
print(OUT)
