"""AI エージェントの MCP 接続と SaaS の鍵を AWS で中央管理する構成案と、SaaS 監査ログの検索。
python3 gen.py -> devin-keys-and-saas-audit.pptx
図の見た目は zenn-doc テンプレ (黒地、色枠の角丸、題の下に色の凡例) に合わせる。
gen.py は parts/p1.py から順に連結して作る。直すときは parts を直して連結し直す。
"""
import os
import sys

sys.path.insert(0, os.environ.get("PPTXKIT", "pptxkit"))
from pptx.dml.color import RGBColor  # noqa: E402
from pptx.enum.shapes import MSO_SHAPE  # noqa: E402
from pptxkit import Deck, shapes as sh  # noqa: E402

HERE = os.path.dirname(os.path.abspath(__file__))
TPL = os.environ.get("PPTXKIT", "pptxkit") + "/templates/aws-dark-2024.pptx"

# 色の役割: 青 = AWS が提供、緑 = お客様が設定・作る、灰 = 他社の製品、黄 = 注意、赤 = 今の問題
BLU, BLU_DK = RGBColor(40, 150, 240), RGBColor(16, 44, 72)
GRN, GRN_DK = RGBColor(36, 160, 70), RGBColor(18, 64, 34)
YEL, YEL_DK = RGBColor(251, 211, 50), RGBColor(80, 66, 14)
RED, RED_DK = RGBColor(220, 60, 90), RGBColor(80, 18, 34)
GRY, GRY_DK = RGBColor(110, 114, 122), RGBColor(52, 56, 64)
CHIP, EDGE = RGBColor(22, 26, 30), RGBColor(70, 74, 80)
MUTED, WHITE = RGBColor(160, 160, 168), RGBColor(245, 245, 248)

AC = "https://docs.aws.amazon.com/bedrock-agentcore/latest/devguide/"
GH = "https://docs.github.com/en/enterprise-cloud@latest/admin/monitoring-activity-in-your-enterprise/reviewing-audit-logs-for-your-enterprise/"
U = {
    "registry": AC + "registry.html",
    "regions": AC + "agentcore-regions.html",
    "jwt": AC + "inbound-jwt-authorizer.html",
    "entra": AC + "identity-idp-microsoft.html",
    "outbound": AC + "gateway-outbound-auth.html",
    "providers": AC + "resource-providers.html",
    "limits": AC + "bedrock-agentcore-limits.html",
    "vault": AC + "key-features-and-benefits.html",
    "policy": AC + "policy.html",
    "gwtrail": AC + "enabling-cloudtrail-data-event-logging.html",
    "acprice": "https://aws.amazon.com/bedrock/agentcore/pricing/",
    "patpol": "https://docs.github.com/en/organizations/managing-programmatic-access-to-your-organization/setting-a-personal-access-token-policy-for-your-organization",
    "ghapp": "https://docs.github.com/en/apps/oauth-apps/building-oauth-apps/differences-between-github-apps-and-oauth-apps",
    "ghmcp": "https://github.com/github/github-mcp-server/blob/main/docs/remote-server.md",
    "smprice": "https://aws.amazon.com/secrets-manager/pricing/",
    "smrot": "https://docs.aws.amazon.com/secretsmanager/latest/userguide/rotating-secrets.html",
    "afapps": "https://docs.aws.amazon.com/appfabric/latest/adminguide/supported-applications.html",
    "afgh": "https://docs.aws.amazon.com/appfabric/latest/adminguide/github.html",
    "afprice": "https://aws.amazon.com/appfabric/pricing/",
    "ghstream": GH + "streaming-the-audit-log-for-your-enterprise",
    "ghevents": GH + "audit-log-events-for-your-enterprise",
    "ghorglog": "https://docs.github.com/en/organizations/keeping-your-organization-secure/managing-security-settings-for-your-organization/reviewing-the-audit-log-for-your-organization",
    "freee": "https://developer.freee.co.jp/",
    "athena": "https://docs.aws.amazon.com/athena/latest/ug/what-is.html",
    "athpart": "https://docs.aws.amazon.com/athena/latest/ug/partitions.html",
    "athprice": "https://aws.amazon.com/athena/pricing/",
    "sfnathena": "https://docs.aws.amazon.com/step-functions/latest/dg/connect-athena.html",
    "ctprice": "https://aws.amazon.com/cloudtrail/pricing/",
    "objlock": "https://docs.aws.amazon.com/AmazonS3/latest/userguide/object-lock.html",
    "kmsdel": "https://docs.aws.amazon.com/kms/latest/developerguide/deleting-keys.html",
    "osa": "https://docs.aws.amazon.com/opensearch-service/latest/developerguide/security-analytics.html",
    "seclake": "https://docs.aws.amazon.com/security-lake/latest/userguide/custom-sources.html",
    "mcpauth": "https://modelcontextprotocol.io/specification/2026-07-28/basic/authorization",
    "rfc9728": "https://www.rfc-editor.org/rfc/rfc9728",
    "rfc8414": "https://www.rfc-editor.org/rfc/rfc8414",
    "rfc7591": "https://www.rfc-editor.org/rfc/rfc7591",
    "cimd": "https://datatracker.ietf.org/doc/draft-ietf-oauth-client-id-metadata-document/",
    "oauth21": "https://datatracker.ietf.org/doc/draft-ietf-oauth-v2-1/",
    "rfc7636": "https://www.rfc-editor.org/rfc/rfc7636",
    "rfc8707": "https://www.rfc-editor.org/rfc/rfc8707",
    "rfc6750": "https://www.rfc-editor.org/rfc/rfc6750",
    "regsync": "https://docs.aws.amazon.com/bedrock-agentcore/latest/devguide/registry-sync-records.html",
    "reglife": "https://docs.aws.amazon.com/bedrock-agentcore/latest/devguide/registry-record-lifecycle.html",
    "regeb": "https://docs.aws.amazon.com/bedrock-agentcore/latest/devguide/registry-eventbridge.html",
    "gwobs": "https://docs.aws.amazon.com/bedrock-agentcore/latest/devguide/observability-gateway-metrics.html",
    "s_entra": "https://github.com/awslabs/agentcore-samples/tree/abab3d8beb684c79fc827a95a2d015e44e76dbbc/06-workshops/02-AgentCore-gateway/04-integration/05-entraid-3lo-gateway",
    "s_proxy": "https://github.com/awslabs/agentcore-samples/blob/abab3d8beb684c79fc827a95a2d015e44e76dbbc/06-workshops/02-AgentCore-gateway/04-integration/05-entraid-3lo-gateway/lambda/mcp_proxy_lambda.py",
    "s_ide": "https://github.com/awslabs/agentcore-samples/tree/abab3d8beb684c79fc827a95a2d015e44e76dbbc/06-workshops/02-AgentCore-gateway/04-integration/03-ide-gateway-tool",
    "z_gw": "https://zenn.dev/aws_japan/articles/amazon-quick-agentcore-gateway-vs-direct-mcp",
    "z_kiro": "https://zenn.dev/aws_japan/articles/agent-registry-kiro-governance",
    "z_cog": "https://zenn.dev/aws_japan/articles/e86655a8ebd114",
    "z_dcr": "https://zenn.dev/aws_japan/articles/agentcore-gateway-remote-mcp",
}


def L(text_, key):
    return f"[{text_}]({U[key]})"


d = Deck(TPL, footer="page", check_layout=False)
t = d.theme


def box(s, x, y, w, h, body="", fill=CHIP, stroke=EDGE, size=18, color=WHITE, bold=True, dash=None,
        radius=0.12, align="center"):
    sp = sh.rect(s, t, x, y, w, h, fill=fill, stroke=stroke, sw=1.5, dash=dash,
                 shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=radius)
    if body:
        sh._write(sp.text_frame, t, body, size, color, bold, align, "middle")
        tf = sp.text_frame
        tf.margin_left = tf.margin_right = sh.U(8)
        tf.margin_top = tf.margin_bottom = sh.U(2)
    return sp


def aws(s, x, y, w, h, body, size=18):
    return box(s, x, y, w, h, body, fill=BLU_DK, stroke=BLU, size=size)


def mine(s, x, y, w, h, body, size=18):
    return box(s, x, y, w, h, body, fill=GRN_DK, stroke=GRN, size=size)


def other(s, x, y, w, h, body, size=18):
    return box(s, x, y, w, h, body, fill=GRY_DK, stroke=GRY, size=size)


def warn(s, x, y, w, h, body, size=18):
    return box(s, x, y, w, h, body, fill=YEL_DK, stroke=YEL, size=size)


def bad(s, x, y, w, h, body, size=18):
    return box(s, x, y, w, h, body, fill=RED_DK, stroke=RED, size=size)


def text(s, x, y, w, h, body, size=18, color=WHITE, bold=False, align="left"):
    return sh.textbox(s, t, x, y, w, h, body, size=size, color=color, bold=bold, align=align, anchor="middle")


def line(s, a, ai, b, bi, color=WHITE, dashed=False, width=2.0, head=False):
    sh.connect(sh.arrow(s, t, 0, 0, 1, 1, color=color, dashed=dashed, width=width, head=head), a, ai, b, bi)


def legend(s, y, items):
    x = 43
    for col, lab in items:
        sh.rect(s, t, x, y + 6, 18, 18, fill=col, stroke=col, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.2)
        text(s, x + 26, y, max(260, 17 * len(lab) + 20), 30, lab, size=15, color=MUTED)
        x += 70 + 17 * len(lab)


LEG = [(BLU, "AWS が提供"), (GRN, "お客様が設定・作る"), (GRY, "他社の製品"), (YEL, "注意")]


def figslide(title_, eyebrow, sub=None, leg=LEG):
    s_ = d.slide(title_, eyebrow=eyebrow, sub=sub)
    if leg:
        legend(s_.s, s_.y, leg)
        s_.spacer(38)
    return s_
