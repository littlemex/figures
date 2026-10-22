"""gen_pptx.py -- figures for the Zenn article e47ba6e81153c2 (scheduler comparison) and the
eks-distributed-ai book chapter Advanced03, as native-shape pptx in the zenn-doc (ZN-25) look.

    /usr/bin/python3 gen_pptx.py <out.pptx>

Text in the figures is English. Colour roles, stated under every title:
  green  = the job-level scheduling piece being explained
  blue   = stock Kubernetes / EKS, unchanged
  yellow = where it goes wrong, or the caveat
  purple = Slurm
Sources: Kubernetes docs (kube-scheduler, Job suspend, ResourceQuota, PriorityClass),
Kueue v0.20 docs, Volcano architecture + gang plugin source, KAI-Scheduler README/docs, Slurm overview, sched_config, preempt,
priority_multifactor (all read 2026-10-22 JST).
"""
import os
import sys

sys.path.insert(0, "/Users/akazawt/works/data-science/tools/pptxkit")
from pptx.dml.color import RGBColor  # noqa: E402
from pptx.enum.dml import MSO_LINE_DASH_STYLE  # noqa: E402
from pptx.enum.shapes import MSO_SHAPE  # noqa: E402

from pptxkit import Deck, shapes as sh  # noqa: E402

TPL = "/Users/akazawt/works/data-science/tools/pptxkit/templates/aws-dark-2024.pptx"
OUT = sys.argv[1] if len(sys.argv) > 1 else "schedulers.pptx"

GRN = RGBColor(36, 160, 70)
GRN_HI = RGBColor(59, 209, 111)
GRN_DK = RGBColor(18, 64, 34)
CHIP = RGBColor(22, 26, 30)
EDGE = RGBColor(70, 74, 80)
GREY = RGBColor(52, 56, 64)
LINK = RGBColor(40, 150, 240)
LINK_DK = RGBColor(16, 44, 72)
YEL = RGBColor(251, 211, 50)
YEL_DK = RGBColor(80, 66, 14)
PUR = RGBColor(170, 120, 240)
PUR_DK = RGBColor(52, 34, 84)
MUTED = RGBColor(150, 150, 158)
WHITE = RGBColor(245, 245, 248)
BLACK = RGBColor(0, 0, 0)

d = Deck(TPL, check_layout=False)
t = d.theme
DASH = MSO_LINE_DASH_STYLE.DASH


def label(s, x, y, w, h, body, size=16, color=WHITE, bold=False, align="center",
          anchor="middle"):
    return sh.textbox(s, t, x, y, w, h, body, size=size, color=color, bold=bold, align=align,
                      anchor=anchor)


def box(s, x, y, w, h, fill, stroke=None, body="", size=15, color=WHITE, radius=0.1,
        dash=None, sw=1.25, bold=True, align="center"):
    sp = sh.rect(s, t, x, y, w, h, fill=fill, stroke=stroke, sw=sw, dash=dash,
                 shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=radius)
    if body:
        sh._write(sp.text_frame, t, body, size, color, bold, align, "middle")
        tf = sp.text_frame
        tf.margin_left = tf.margin_right = sh.U(8)
        tf.margin_top = tf.margin_bottom = sh.U(2)
    return sp


def panel(s, x, y, w, h, head, color, size=16):
    box(s, x, y, w, h, CHIP, stroke=color, radius=0.04, sw=1.5)
    if head:
        label(s, x + 16, y + 8, w - 32, 26, head, size=size, bold=True, color=color, align="left")


def arrow(s, x1, y1, x2, y2, color=WHITE, width=2, dashed=False):
    return sh.arrow(s, t, x1, y1, x2, y2, color=color, width=width, dashed=dashed)


def elbow(s, pts, color=WHITE, width=2, dashed=False):
    """A polyline of straight segments, arrowhead on the last one only."""
    for i in range(len(pts) - 1):
        (x1, y1), (x2, y2) = pts[i], pts[i + 1]
        if i == len(pts) - 2:
            arrow(s, x1, y1, x2, y2, color, width, dashed)
        else:
            c = s.shapes.add_connector(1, sh.U(x1), sh.U(y1), sh.U(x2), sh.U(y2))
            c.shadow.inherit = False
            c.line.color.rgb = color
            c.line.width = sh.Pt(width)
            if dashed:
                c.line.dash_style = DASH


def num(s, x, y, n, color=YEL):
    sh.badge(s, t, x, y, 30, str(n), color=BLACK, size=15, fill=color)


def title(s, head, legend):
    label(s, 43, 24, 1194, 46, head, size=30, bold=True, align="left")
    label(s, 43, 70, 1194, 26, legend, size=15, color=MUTED, align="left")


def takeaway(s, text):
    sh.rect(s, t, 43, 640, 1194, 1, fill=EDGE)
    label(s, 43, 652, 1194, 40, text, size=19, bold=True, align="left")


LEG = "Green: job-level scheduling.  Blue: stock Kubernetes.  Yellow: the catch."

# ============================================================== 1. intro: Slurm vs vanilla
s = d.slide().s
title(s, "What Slurm does that vanilla Kubernetes does not",
      "Purple: Slurm, as configured.  Yellow: what vanilla Kubernetes does instead")
rows = [
    ("Start a multi-node job\nas one unit",
     "starts only when the requested\nnodes and GPUs are all granted",
     "Pods are bound one by one;\nsome run, some stay Pending"),
    ("Hold a job in a queue\nunder a limit",
     "group limits such as GrpTRES\nkeep the job PENDING",
     "ResourceQuota rejects the Pod;\nthe Job controller just retries"),
    ("Run short jobs early,\nbig job not delayed",
     "backfill: only if the big job's\nplanned start does not move",
     "no planned start time\nfor a whole job"),
    ("Pause a running job,\nresume, no restart",
     "suspend stops the processes;\nresume continues them",
     "suspending a Job\ndeletes its running Pods"),
    ("Priority from\npast usage",
     "multifactor priority with\nfair-share from accounting",
     "PriorityClass gives each Pod\na fixed value"),
]
label(s, 43, 108, 330, 28, "ID and capability", size=15, color=MUTED, align="left")
label(s, 390, 108, 400, 28, "Slurm", size=15, color=PUR, bold=True, align="left")
label(s, 808, 108, 429, 28, "Vanilla Kubernetes", size=15, color=YEL, bold=True, align="left")
y = 140
for k, (cap, slurm, k8s) in enumerate(rows):
    box(s, 43, y, 64, 88, GRN_DK, stroke=GRN_HI, body=f"D{k+1}", size=20, radius=0.12)
    box(s, 115, y, 258, 88, CHIP, stroke=EDGE, body=cap, size=15, radius=0.08, align="left")
    box(s, 390, y, 400, 88, PUR_DK, stroke=PUR, body=slurm, size=15, radius=0.16)
    box(s, 808, y, 429, 88, YEL_DK, stroke=YEL, body=k8s, size=15, radius=0.16)
    y += 98
takeaway(s, "Each row decides about a whole job or account; vanilla Kubernetes decides per Pod")

# ============================================================== 2. landscape
s = d.slide().s
title(s, "Where each scheduler sits", "Green: job-level scheduling.  Blue: stock Kubernetes.  Purple: Slurm")
panel(s, 43, 110, 820, 520, "Kubernetes: the steps are separate", LINK)
box(s, 73, 152, 230, 96, LINK_DK, stroke=LINK, body="1  Job / JobSet /\nTrainJob object\nis created", size=14)
box(s, 333, 152, 230, 96, GRN_DK, stroke=GRN_HI, body="2  Kueue\nholds it suspended\nuntil admitted", size=14)
box(s, 593, 152, 240, 96, LINK_DK, stroke=LINK, body="3  its controller\ncreates the Pods", size=14)
arrow(s, 303, 200, 333, 200)
arrow(s, 563, 200, 593, 200)
box(s, 73, 290, 760, 190, None, stroke=EDGE, radius=0.06, dash=DASH)
label(s, 89, 298, 730, 26, "4  Scheduler: which node?  (picked per Pod by schedulerName)",
      size=15, bold=True, color=WHITE, align="left")
box(s, 93, 336, 230, 76, LINK_DK, stroke=LINK, body="kube-scheduler\nper Pod", size=14)
box(s, 339, 336, 230, 76, GRN_DK, stroke=GRN_HI, body="Volcano\nPodGroup", size=14)
box(s, 585, 336, 228, 76, GRN_DK, stroke=GRN_HI, body="KAI-Scheduler\nPodGroup", size=14)
label(s, 93, 420, 720, 50, "Volcano and KAI also bring their own queues,\nso for their Pods they cover step 2 as well.",
      size=13, color=MUTED, align="left")
elbow(s, [(713, 248), (713, 290)])
box(s, 73, 510, 330, 100, LINK_DK, stroke=LINK, body="5  kubelet on the node\nruns the Pod", size=15)
box(s, 433, 510, 400, 100, LINK_DK, stroke=LINK, body="Karpenter (EKS): adds EC2 nodes\nfor Pods marked Unschedulable;\nit does not order jobs", size=14)
arrow(s, 238, 480, 238, 510)
panel(s, 893, 110, 344, 520, "Slurm: one controller", PUR)
box(s, 913, 152, 304, 70, PUR_DK, stroke=PUR, body="sbatch / salloc\nsubmit the job", size=14)
box(s, 913, 262, 304, 160, PUR_DK, stroke=PUR, body="slurmctld\nqueue + priority + backfill\n+ the allocation,\nin one decision", size=15)
box(s, 913, 462, 304, 70, PUR_DK, stroke=PUR, body="slurmd on each node\nlaunches the tasks", size=14)
arrow(s, 1065, 222, 1065, 262)
arrow(s, 1065, 422, 1065, 462)
label(s, 913, 548, 304, 60, "Steps 2 and 4 on the left\nare one decision here.", size=13, color=MUTED)
takeaway(s, "Kueue fills step 2; Volcano and KAI take step 4 for the Pods they own")

# ============================================================== 3. partial placement
s = d.slide().s
title(s, "Why a 4-Pod job can get stuck half placed", "Blue: stock Kubernetes.  Yellow: the catch")
panel(s, 43, 110, 380, 510, "The job's 4 Pods", LINK)
for i in range(4):
    st = "bound" if i < 2 else "Pending"
    box(s, 73, 152 + i * 88, 320, 70, LINK_DK if i < 2 else YEL_DK, stroke=LINK if i < 2 else YEL,
        body=f"worker-{i}   ({st})", size=17)
label(s, 73, 520, 320, 60, "kube-scheduler takes them\none at a time", size=14, color=MUTED)
panel(s, 463, 110, 774, 510, "Cluster: 2 GPUs in total", LINK)
for j, nm in enumerate(("node-a", "node-b")):
    x = 493 + j * 370
    box(s, x, 152, 340, 230, GREY, stroke=EDGE, radius=0.05)
    label(s, x + 14, 160, 300, 26, nm, size=15, bold=True, align="left")
    box(s, x + 20, 200, 300, 156, LINK_DK, stroke=LINK, body=f"GPU  <-  worker-{j}\n\nwaiting at rendezvous\nfor its peers", size=16)
box(s, 493, 404, 714, 196, YEL_DK, stroke=YEL,
    body="worker-2 and worker-3 have nowhere to go.\n\nworker-0 and worker-1 keep their GPUs and wait.\nTwo such jobs can block each other for good.", size=17)
takeaway(s, "kube-scheduler never asks whether all four fit before binding the first")

# ============================================================== 4. kube-scheduler architecture
s = d.slide().s
title(s, "Vanilla Kubernetes: who touches a Pod",
      "Blue: stock Kubernetes.  Arrows show which way the data goes; all of it through the API server")
box(s, 503, 270, 274, 120, LINK_DK, stroke=LINK, body="API server", size=22)
box(s, 563, 440, 154, 64, GREY, stroke=EDGE, body="etcd", size=17)
arrow(s, 640, 390, 640, 440, color=MUTED)
box(s, 73, 140, 320, 110, LINK_DK, stroke=LINK, body="kube-controller-manager\nJob controller", size=16)
box(s, 887, 140, 320, 110, LINK_DK, stroke=LINK, body="kube-scheduler", size=18)
box(s, 73, 430, 320, 110, LINK_DK, stroke=LINK, body="kubelet\non every node", size=17)
box(s, 887, 430, 320, 110, YEL_DK, stroke=YEL, body="the Job object exists,\nbut the scheduler never\nreads it", size=15)
elbow(s, [(393, 195), (450, 195), (450, 300), (503, 300)])
label(s, 400, 150, 200, 40, "1  create N Pods", size=14, align="left")
elbow(s, [(777, 300), (830, 300), (830, 180), (887, 180)])
label(s, 650, 140, 230, 40, "2  unscheduled Pods", size=14, align="right")
elbow(s, [(887, 220), (860, 220), (860, 350), (777, 350)])
label(s, 790, 360, 240, 50, "3  Binding,\none Pod per cycle", size=14, align="left")
elbow(s, [(503, 360), (450, 360), (450, 485), (393, 485)])
label(s, 400, 494, 220, 50, "4  Pods bound\nto this node", size=14, align="left")
takeaway(s, "Every component works on Pods; nothing places the job as a whole")

# ============================================================== 5. Kueue behaviour
s = d.slide().s
title(s, "Kueue: hold whole jobs until the quota has room",
      "Green: Kueue.  Blue: stock Kubernetes.  Quota 2 GPUs, each job needs 2, same priority, first in first out")
lanes = [("job-a", [(0, 3, "admitted, running", GRN_DK, GRN_HI)]),
         ("job-b", [(0, 3, "suspended: no Pods", GREY, EDGE), (3, 6, "admitted, running", GRN_DK, GRN_HI)]),
         ("job-c", [(0, 6, "suspended: no Pods", GREY, EDGE), (6, 9, "admitted, running", GRN_DK, GRN_HI)])]
X0, XW = 220, 1000 / 9
for i, (nm, segs) in enumerate(lanes):
    y = 150 + i * 96
    label(s, 43, y, 160, 70, nm, size=20, bold=True, align="left")
    for a, b, txt, fill, stroke in segs:
        box(s, X0 + a * XW + 3, y, (b - a) * XW - 6, 70, fill, stroke=stroke, body=txt, size=15,
            color=WHITE if fill != GREY else MUTED, radius=0.15)
y = 150 + 3 * 96
sh.rect(s, t, X0, y + 4, 1000, 2, fill=MUTED)
for k, tx in ((0, "submit all 3"), (3, "job-a done"), (6, "job-b done")):
    sh.rect(s, t, X0 + k * XW - 1, y - 6, 2, 22, fill=MUTED)
    label(s, X0 + k * XW - 80, y + 18, 160, 24, tx, size=13, color=MUTED)
box(s, 43, 500, 590, 120, LINK_DK, stroke=LINK,
    body="Nodes appear only for the admitted job:\nKarpenter never sees job-b or job-c\nwhile they are suspended", size=15)
box(s, 653, 500, 584, 120, YEL_DK, stroke=YEL,
    body="Admitted does not mean placed: if spot capacity\nis short, the Pods still wait. waitForPodsReady\ncan evict and requeue the job", size=15)
takeaway(s, "Kueue decides when a job may create Pods; kube-scheduler still picks the node")

# ============================================================== 6. Kueue on EKS architecture
s = d.slide().s
title(s, "Kueue on EKS: what happens to one TrainJob",
      "Green: Kueue (one controller-manager process and its objects).  Blue: unchanged EKS components")
box(s, 43, 120, 230, 96, CHIP, stroke=EDGE, body="1  kubectl apply TrainJob\nlabel kueue.x-k8s.io/\nqueue-name: gpu-ddp", size=13)
box(s, 43, 252, 230, 96, LINK_DK, stroke=LINK, body="EKS API server", size=17)
arrow(s, 158, 216, 158, 252)
panel(s, 313, 110, 924, 262, "Kueue  (kueue-controller-manager in kueue-system)", GRN_HI)
box(s, 333, 150, 210, 90, GRN_DK, stroke=GRN, body="2  webhook\nsets suspend: true", size=14)
box(s, 563, 150, 210, 90, GRN_DK, stroke=GRN, body="3  job reconciler\ncreates a Workload", size=14)
box(s, 793, 150, 200, 90, GRN_DK, stroke=GRN, body="4  LocalQueue\n(namespaced)", size=14)
box(s, 1013, 150, 204, 90, GRN_DK, stroke=GRN_HI, body="ClusterQueue\nGPU quota\n(cluster-scoped)", size=13)
box(s, 1013, 262, 204, 94, GRN_DK, stroke=GRN, body="ResourceFlavor\nnodeLabels\n(cluster-scoped)", size=13)
arrow(s, 543, 195, 563, 195)
arrow(s, 773, 195, 793, 195)
arrow(s, 993, 195, 1013, 195)
arrow(s, 1115, 262, 1115, 240, color=GRN)
elbow(s, [(273, 300), (295, 300), (295, 195), (333, 195)])
box(s, 333, 262, 660, 94, GRN_DK, stroke=GRN_HI,
    body="5  admitted: add the flavor's nodeSelector + tolerations, clear suspend\nnot admitted: the job stays suspended and no Pods exist", size=14)
panel(s, 43, 396, 1194, 170, "", LINK)
label(s, 843, 402, 378, 26, "Unchanged EKS components", size=16, bold=True, color=LINK, align="right")
box(s, 73, 440, 330, 100, LINK_DK, stroke=LINK, body="6  Trainer creates the JobSet;\nJobSet and Job controllers\ncreate the Pods", size=14)
box(s, 463, 440, 300, 100, LINK_DK, stroke=LINK, body="7  kube-scheduler\nbinds each Pod\nonce a node fits", size=14)
box(s, 853, 440, 354, 100, LINK_DK, stroke=LINK, body="8  Karpenter\nlaunches EC2 nodes for\nUnschedulable Pods", size=14)
arrow(s, 403, 490, 463, 490)
arrow(s, 763, 470, 853, 470)
arrow(s, 853, 512, 763, 512)
label(s, 768, 516, 84, 22, "node Ready", size=11, color=MUTED)
elbow(s, [(663, 356), (663, 384), (238, 384), (238, 440)], color=GRN_HI)
label(s, 43, 576, 1194, 56, "LocalQueue and Workload live in the job's namespace; ClusterQueue and ResourceFlavor are cluster-wide.\nIn this book's kueue.tf the webhook acts only in namespaces labelled managed-by-kueue=true.",
      size=13, color=MUTED, align="left")
takeaway(s, "Kueue holds and labels the job; Trainer, kube-scheduler and Karpenter are unchanged")

# ============================================================== 8. Volcano architecture
s = d.slide().s
title(s, "Volcano: its own controllers and its own scheduler", "Green: Volcano.  Blue: stock Kubernetes")
box(s, 43, 130, 240, 90, CHIP, stroke=EDGE, body="kubectl apply\nvcjob", size=15)
box(s, 43, 280, 240, 110, LINK_DK, stroke=LINK, body="API server", size=19)
arrow(s, 163, 220, 163, 280)
panel(s, 323, 110, 914, 510, "Volcano", GRN_HI)
box(s, 353, 150, 420, 84, GRN_DK, stroke=GRN, body="admission webhook\nvalidates and defaults\nJob, PodGroup, Queue, Pod", size=13)
box(s, 353, 250, 420, 84, GRN_DK, stroke=GRN, body="controller-manager\ncreates the PodGroup\nand the Job's Pods", size=13)
box(s, 803, 150, 404, 184, CHIP, stroke=GRN, dash=DASH, body="Queue  (an object, not a process)\nweight and capability per tenant,\nread by the scheduler", size=14)
panel(s, 353, 352, 854, 248, "volcano-scheduler session  (actions run in the configured order)", GRN_HI)
acts = [("enqueue", "room in the\nQueue?"), ("allocate", "try each\ntask"), ("preempt", "by\npriority"),
        ("reclaim", "back to the\nQueue's share"), ("backfill", "best-effort\nPods")]
for k, (a, b) in enumerate(acts):
    x = 373 + k * 166
    box(s, x, 396, 150, 96, GRN_DK, stroke=GRN_HI, body=f"{a}\n{b}", size=13)
    if k < 4:
        arrow(s, x + 150, 444, x + 166, 444)
box(s, 373, 512, 814, 64, GRN_DK, stroke=GRN, body="plugins: gang, drf, proportion, priority ...", size=14)
elbow(s, [(283, 320), (330, 320), (330, 192), (353, 192)], color=LINK)
elbow(s, [(353, 470), (300, 470), (300, 390)], color=LINK)
label(s, 43, 410, 250, 80, "the scheduler binds\ncommitted tasks\nthrough the API server", size=13, color=MUTED)
takeaway(s, "Volcano covers queueing, gang and placement in one system for its Pods")

# ============================================================== 10. KAI architecture
s = d.slide().s
title(s, "KAI-Scheduler: grouping, scheduling and binding are separate",
      "Green: KAI-Scheduler.  Blue: stock Kubernetes.  Arrows are objects written and watched via the API server")
box(s, 43, 270, 230, 120, LINK_DK, stroke=LINK, body="API server", size=20)
box(s, 323, 130, 280, 100, GRN_DK, stroke=GRN, body="pod-grouper\nwrites a PodGroup for\nPyTorchJob, JobSet, ...", size=14)
box(s, 643, 130, 280, 100, GRN_DK, stroke=GRN, body="queue controller\nhierarchical queues", size=14)
box(s, 323, 280, 600, 120, GRN_DK, stroke=GRN_HI, body="scheduler\nsession + actions: allocate, consolidate, reclaim, preempt\nwrites a BindRequest", size=15)
box(s, 973, 280, 264, 120, GRN_DK, stroke=GRN, body="binder\n(own Deployment)\nreads it, binds the Pod", size=15)
arrow(s, 923, 340, 973, 340)
label(s, 918, 300, 60, 24, "", size=11)
arrow(s, 463, 230, 463, 280)
arrow(s, 783, 230, 783, 280)
elbow(s, [(1105, 400), (1105, 440), (158, 440), (158, 390)], color=LINK)
arrow(s, 273, 300, 323, 300, color=LINK)
box(s, 323, 480, 914, 120, YEL_DK, stroke=YEL,
    body="GPU sharing: a Pod asks for a fraction with the gpu-fraction annotation;\na reservation Pod holds the physical GPU. Memory is not isolated between the sharing Pods by default.", size=14)
takeaway(s, "Placement is decided in one Deployment and applied by another")

# ============================================================== 11. Slurm job unit
s = d.slide().s
title(s, "What \"job\" means on each side", "Blue: Kubernetes.  Purple: Slurm.  Green: what had to be added")
panel(s, 43, 110, 584, 510, "Kubernetes  (simplified)", LINK)
box(s, 73, 152, 524, 70, LINK_DK, stroke=LINK, body="Job / JobSet / TrainJob", size=17)
for k in range(4):
    box(s, 73 + k * 132, 270, 116, 70, LINK_DK, stroke=LINK, body=f"Pod {k}", size=15)
    arrow(s, 335, 222, 131 + k * 132, 270, color=LINK)
box(s, 73, 392, 524, 90, GRN_DK, stroke=GRN_HI, body="PodGroup / Kueue Workload\nadded so a scheduler or a queue sees the group", size=15)
label(s, 73, 500, 524, 90, "Each Pod is scheduled as its own object.\nThe group is an extra object on top.", size=15, color=MUTED)
panel(s, 653, 110, 584, 510, "Slurm", PUR)
box(s, 683, 152, 524, 70, PUR_DK, stroke=PUR, body="job = one allocation of nodes, CPUs, GPUs", size=16)
box(s, 683, 270, 524, 70, PUR_DK, stroke=PUR, body="job step (srun) inside the allocation", size=16)
for k in range(4):
    box(s, 683 + k * 132, 392, 116, 70, PUR_DK, stroke=PUR, body=f"task {k}", size=15)
arrow(s, 945, 222, 945, 270, color=PUR)
for k in range(4):
    arrow(s, 945, 340, 741 + k * 132, 392, color=PUR)
label(s, 683, 500, 524, 90, "Resources are granted to the job first;\ntasks start inside what was granted.", size=15, color=MUTED)
takeaway(s, "Slurm grants resources per job; Kubernetes adds a group object to get there")

# ============================================================== 12. Slurm architecture
s = d.slide().s
title(s, "Slurm: allocation first, then the tasks", "Purple: Slurm daemons.  Grey: commands and optional parts")
box(s, 43, 130, 280, 100, GREY, stroke=EDGE, body="sbatch / salloc\nrequest resources;\na batch script runs on node-0", size=14)
box(s, 383, 120, 470, 200, PUR_DK, stroke=PUR, body="slurmctld\nqueue, multifactor priority, backfill\n-> one allocation for the whole job", size=16)
box(s, 933, 140, 304, 90, GREY, stroke=EDGE, body="slurmdbd (optional)\naccounting, fair-share history", size=14)
arrow(s, 323, 180, 383, 180, color=PUR)
arrow(s, 933, 185, 853, 185, color=MUTED)
box(s, 43, 330, 280, 100, GREY, stroke=EDGE, body="srun\nruns inside the batch script\nor the salloc shell", size=14)
arrow(s, 183, 230, 183, 330, color=MUTED, dashed=True)
elbow(s, [(323, 360), (520, 360), (520, 320)], color=PUR)
label(s, 530, 330, 330, 30, "1  create a step, get its credential", size=14, align="left")
for k in range(3):
    elbow(s, [(323, 410), (513 + k * 290, 410), (513 + k * 290, 500)], color=PUR)
label(s, 540, 376, 680, 30, "2  ask the slurmd of each node in the step to launch", size=14, align="left")
for k in range(3):
    x = 383 + k * 290
    box(s, x, 500, 260, 120, PUR_DK, stroke=PUR, body=f"node-{k}\nslurmd -> slurmstepd\n-> task processes", size=14)
takeaway(s, "No container layer by default: slurmstepd runs the tasks as processes")

# ============================================================== 13. intro: the stuck job, as an architecture
s = d.slide().s
title(s, "Vanilla Kubernetes: a 4-worker job stuck half placed",
      "Blue: stock Kubernetes, each part working correctly.  Yellow: the outcome no part notices")
box(s, 43, 130, 230, 90, CHIP, stroke=EDGE, body="TrainJob\n4 workers x 1 GPU", size=15)
box(s, 43, 270, 230, 90, LINK_DK, stroke=LINK, body="Job controller\ncreates 4 Pods", size=15)
arrow(s, 158, 220, 158, 270)
panel(s, 313, 110, 420, 380, "kube-scheduler", LINK)
label(s, 333, 146, 380, 24, "scheduling queue: one Pod per cycle", size=14, color=MUTED, align="left")
for i in range(4):
    ok = i < 2
    box(s, 343, 182 + i * 72, 220, 58, LINK_DK if ok else YEL_DK, stroke=LINK if ok else YEL,
        body=f"worker-{i}", size=16)
    box(s, 583, 182 + i * 72, 130, 58, None, stroke=LINK if ok else YEL,
        body="bound" if ok else "no fit:\nPending", size=13, color=LINK if ok else YEL)
arrow(s, 273, 315, 343, 315)
panel(s, 773, 110, 464, 380, "Cluster: 2 GPUs in total", LINK)
for j in range(2):
    y = 150 + j * 150
    box(s, 793, y, 424, 130, GREY, stroke=EDGE, radius=0.05)
    label(s, 807, y + 6, 300, 24, f"node-{'ab'[j]}  (1 GPU)", size=14, bold=True, align="left")
    box(s, 813, y + 38, 384, 78, LINK_DK, stroke=LINK, body=f"worker-{j} holds the GPU,\nwaits at rendezvous for the other 3", size=14)
elbow(s, [(713, 211), (750, 211), (750, 230), (813, 230)], color=LINK)
elbow(s, [(713, 283), (760, 283), (760, 380), (813, 380)], color=LINK)
box(s, 793, 448, 424, 34, None, stroke=YEL, dash=DASH, body="no GPU left for worker-2 and worker-3", size=13, color=YEL, radius=0.3)
box(s, 43, 520, 1194, 100, YEL_DK, stroke=YEL,
    body="Result: 2 GPUs are held and the job makes no progress, possibly for good.\nNo part is broken: the scheduler handled each Pod correctly; nothing asked whether all 4 fit together.", size=16)
takeaway(s, "The default scheduler decides per Pod, so a job can be stuck half placed")

d.save(OUT)
sys.path.insert(0, "/Users/akazawt/works/data-science/claudecode/.claude/skills/review-html/packs/zenn-doc/templates")
import strip_branding  # noqa: E402
strip_branding.main(OUT, OUT)
print(OUT)
