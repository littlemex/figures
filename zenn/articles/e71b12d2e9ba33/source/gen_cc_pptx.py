"""gen_cc_pptx.py — Trainium の Collective Communication の記事の図を、ネイティブ shape の pptx 3 枚で組む。

    /usr/bin/python3 gen_cc_pptx.py <out.pptx>

1. scale-up トポロジーの変遷 (Trn1 → Trn2 → Trn2 UltraServer → Trn3 UltraServer)
2. Collective Communication を動かすソフトウェアスタック (Neuron SDK 2.32.0)
3. 計算と通信の重ね方: GPU (NCCL の既定の経路) と Trainium
図の文字は英語。出典 (2026-10-04 取得):
- trn1-arch / trn2-arch / trn3-arch.html、neuron-runtime/about/collectives.html
- setup/pytorch/manual.html (パッケージと版)、release-notes/index.html (2.32.0, 2026-08-17)
- neuron-runtime/explore/compute-comm-overlap.html、NCCL env.html と 2.28.3 / 2.30.7 のリリースノート
"""
import sys

sys.path.insert(0, "/Users/akazawt/works/data-science/tools/pptxkit")
from pptx.dml.color import RGBColor  # noqa: E402
from pptx.enum.dml import MSO_LINE_DASH_STYLE  # noqa: E402
from pptx.enum.shapes import MSO_SHAPE  # noqa: E402

from pptxkit import Deck, shapes as sh  # noqa: E402

TPL = "/Users/akazawt/works/data-science/tools/pptxkit/templates/aws-dark-2024.pptx"
OUT = sys.argv[1] if len(sys.argv) > 1 else "cc-figures.pptx"

GRN = RGBColor(36, 160, 70)
GRN_HI = RGBColor(59, 209, 111)
GRN_DK = RGBColor(18, 64, 34)
CHIP = RGBColor(22, 26, 30)
EDGE = RGBColor(70, 74, 80)
GREY = RGBColor(52, 56, 64)
LINK = RGBColor(40, 150, 240)
LINK_DK = RGBColor(16, 44, 72)
YEL = RGBColor(251, 211, 50)
YEL_DK = RGBColor(80, 66, 14)
MUTED = RGBColor(150, 150, 158)
WHITE = RGBColor(245, 245, 248)

d = Deck(TPL, check_layout=False)
t = d.theme


def label(s, x, y, w, h, body, size=16, color=WHITE, bold=False, align="center",
          anchor="middle"):
    return sh.textbox(s, t, x, y, w, h, body, size=size, color=color, bold=bold, align=align,
                      anchor=anchor)


def box(s, x, y, w, h, fill, stroke=None, body="", size=14, color=WHITE, radius=0.08,
        dash=None, sw=1.0):
    sp = sh.rect(s, t, x, y, w, h, fill=fill, stroke=stroke, sw=sw, dash=dash,
                 shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=radius)
    if body:
        sh._write(sp.text_frame, t, body, size, color, True, "center", "middle")
        tf = sp.text_frame
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = sh.U(2)
    return sp


def line(s, x1, y1, x2, y2, color=LINK, w=1.5):
    return sh.rect(s, t, min(x1, x2) - (w / 2 if x1 == x2 else 0),
                   min(y1, y2) - (w / 2 if y1 == y2 else 0),
                   abs(x2 - x1) or w, abs(y2 - y1) or w, fill=color)


def title(s, head, sub):
    label(s, 43, 26, 1190, 46, head, size=30, bold=True, align="left")
    label(s, 43, 72, 1190, 26, sub, size=16, color=MUTED, align="left")


def torus(s, x, y, n=4, p=34, node=16, color=GRN):
    for r in range(n):
        line(s, x - 14, y + r * p, x + (n - 1) * p + 14, y + r * p)
        line(s, x + r * p, y - 14, x + r * p, y + (n - 1) * p + 14)
    for r in range(n):
        for c in range(n):
            box(s, x + c * p - node / 2, y + r * p - node / 2, node, node, color, radius=0.2)


def tags(s, x, y, w, items):
    for j, tag in enumerate(items):
        box(s, x, y + j * 36, w, 30, GRN_DK, stroke=GRN_HI, body=tag, size=14, radius=0.3)


# ------------------------------------------------------------------ 1. topology
s = d.slide()
title(s.s, "Scale-up topology, generation by generation", "Green tags: what changed")
COLS = [43, 352, 661, 970]
CW = 268
heads = [("Trn1", "trn1.32xlarge"), ("Trn2", "trn2.48xlarge"), ("Trn2", "UltraServer"),
         ("Trn3", "UltraServer")]
for x, (a, b) in zip(COLS, heads):
    label(s.s, x, 112, CW, 32, a, size=24, bold=True)
    label(s.s, x, 142, CW, 22, b, size=14, color=MUTED)
    box(s.s, x, 172, CW, 250, CHIP, stroke=EDGE, radius=0.04)

torus(s.s, COLS[0] + 83, 222)
label(s.s, COLS[0], 372, CW, 40, "16 chips  |  NeuronLink-v2", size=14, color=MUTED)
torus(s.s, COLS[1] + 83, 222)
label(s.s, COLS[1], 372, CW, 40, "16 chips  |  NeuronLink-v3", size=14, color=MUTED)
# UltraServer: four 4x4 layers stacked in depth, chips with the same coordinates on a Z ring
for k in range(3, -1, -1):
    ox, oy = COLS[2] + 66 + k * 22, 206 + k * 26
    torus(s.s, ox, oy, p=26, node=12, color=GRN if k == 0 else RGBColor(28, 110, 52))
for c in (0, 3):
    for r in (0, 3):
        x0, y0 = COLS[2] + 66 + c * 26, 206 + r * 26
        sh.arrow(s.s, t, x0, y0, x0 + 66, y0 + 78, color=YEL, width=1.5, head=True)
label(s.s, COLS[2], 372, CW, 40, "64 chips  |  4 x 16, Z ring", size=14, color=MUTED)
# Trn3 Gen2: chips -> level-1 switch per server -> level-2 switches across servers
x3 = COLS[3]
for k in range(2):
    box(s.s, x3 + 34 + k * 104, 186, 96, 30, LINK_DK, stroke=LINK, body="L2", size=12)
for i in range(4):
    sx = x3 + 18 + i * 60
    for k in range(2):
        sh.arrow(s.s, t, sx + 18 + k * 16, 262, x3 + 82 + k * 104, 216, color=LINK, width=1.2)
    box(s.s, sx, 262, 52, 26, LINK_DK, stroke=LINK, body="L1", size=11)
    for j in range(4):
        cx = sx + j * 13 + 1
        line(s.s, cx + 5, 288, cx + 5, 316, w=1)
        box(s.s, cx, 316, 11, 11, GRN, radius=0.2)
    box(s.s, sx - 3, 254, 58, 82, None, stroke=EDGE, dash=MSO_LINE_DASH_STYLE.DASH, radius=0.08)
label(s.s, x3, 340, CW, 22, "Gen2: 4 of 36 servers shown", size=12, color=MUTED)
label(s.s, x3, 366, CW, 50, "Gen1: 64 chips\nGen2: 144 chips", size=14, color=MUTED)

tags(s.s, COLS[0], 440, CW, ["2D torus", "4 neighbours per chip"])
tags(s.s, COLS[1], 440, CW, ["Same 4x4 2D torus", "Faster NeuronLink-v3"])
tags(s.s, COLS[2], 440, CW, ["+ Z ring across 4 servers", "3D torus (4x4x4)"])
tags(s.s, COLS[3], 440, CW, ["Torus -> switch", "All-to-all"])
for i in range(3):
    ax = COLS[i] + CW + 4
    sh.arrow(s.s, t, ax, 297, COLS[i + 1] - 4, 297, color=GRN_HI, width=2.5)

# ------------------------------------------------------------------ 2. software stack
s = d.slide()
title(s.s, "From all_reduce() to the wire", "What each layer does, and when it works")
LX, LW = 43, 790
rows = [
    ("Framework", "your code calls all_reduce / all_gather", "you call it", GRN),
    ("Neuron compiler", "builds the NEFF; puts a CC trigger (PTC2) at each collective", "compile", GRN),
    ("Neuron Runtime", "loads the NEFF, CC program and DMA rings", "load + run", LINK),
    ("aws-neuronx-collectives", "implements Ring / Mesh / RDH / KangaRing", "using a collective", LINK),
]
y = 112
for name, role, when, col in rows:
    box(s.s, LX, y, LW, 58, CHIP, stroke=col, radius=0.08, sw=1.5)
    label(s.s, LX + 16, y, 300, 58, name, size=17, bold=True, align="left")
    label(s.s, LX + 320, y, LW - 330, 58, role, size=14, color=MUTED, align="left")
    box(s.s, LX + LW + 20, y + 12, 190, 34, GRN_DK if col == GRN else LINK_DK,
        stroke=GRN_HI if col == GRN else LINK, body=when, size=13, radius=0.3)
    sh.arrow(s.s, t, LX + LW / 2, y + 58, LX + LW / 2, y + 70, color=MUTED, width=1.5)
    y += 70
# two paths below: NeuronLink inside the scale-up domain, EFA beyond it
HW = 385
for k, (name, role, hw, col, hcol) in enumerate((
        ("aws-neuronx-dkms", "driver for the Neuron devices  (kernel)", "CC-Cores + DMA over NeuronLink", GREY, GRN),
        ("EFA software", "libfabric + rdma-core + driver  (multi-instance)", "EFA network interface", GREY, EDGE))):
    x = LX + k * (HW + 20)
    box(s.s, x, y, HW, 62, CHIP, stroke=col, radius=0.08, sw=1.5)
    label(s.s, x + 14, y + 4, HW - 20, 30, name, size=16, bold=True, align="left")
    label(s.s, x + 14, y + 32, HW - 20, 26, role, size=12, color=MUTED, align="left")
    sh.arrow(s.s, t, x + HW / 2, y + 62, x + HW / 2, y + 84, color=MUTED, width=1.5)
    box(s.s, x, y + 84, HW, 54, GRN_DK if k == 0 else GREY, stroke=hcol, body=hw, size=15)
label(s.s, LX, y + 142, HW, 26, "inside the NeuronLink domain", size=13, color=GRN_HI)
label(s.s, LX + HW + 20, y + 142, HW, 26, "beyond it", size=13, color=MUTED)
box(s.s, LX + LW + 20, y + 14, 190, 124, None, stroke=EDGE, dash=MSO_LINE_DASH_STYLE.DASH,
    body="aws-neuronx-tools\n\ncheck and measure", size=13, color=MUTED, radius=0.1)

# ------------------------------------------------------------------ 3. compute-communication overlap
s = d.slide()
title(s.s, "Where compute and communication contend", "GPU with NCCL default kernels vs Trainium")
for x, head in ((43, "GPU  (NCCL default)"), (665, "Trainium")):
    label(s.s, x, 112, 572, 34, head, size=22, bold=True)
    box(s.s, x, 152, 572, 292, CHIP, stroke=EDGE, radius=0.04)

# GPU: the SMs run both the compute kernels and the NCCL kernels
gx, gy = 70, 180
for r in range(4):
    for c in range(8):
        comm = c >= 6
        box(s.s, gx + c * 44, gy + r * 40, 38, 32, LINK_DK if comm else GRN_DK,
            stroke=LINK if comm else GRN, radius=0.12)
label(s.s, gx, gy + 164, 260, 22, "SMs: compute kernels", size=13, color=GRN_HI, align="left")
label(s.s, gx + 264, gy + 164, 90, 22, "NCCL", size=13, color=LINK, align="left")
box(s.s, gx + 370, gy + 10, 150, 46, GREY, stroke=EDGE, body="HBM", size=15)
box(s.s, gx + 370, gy + 92, 150, 46, GREY, stroke=EDGE, body="NVLink / NIC", size=15)
sh.arrow(s.s, t, gx + 352, gy + 115, gx + 370, gy + 115, color=LINK, width=2)
box(s.s, 43, 456, 572, 36, YEL_DK, stroke=YEL, body="Shared: SMs", size=16, radius=0.3)

# Trainium: compute engines trigger CC-Cores; DMA engines are shared
tx, ty = 690, 180
box(s.s, tx, ty, 232, 94, GRN_DK, stroke=GRN, body="Compute engines", size=18)
box(s.s, tx + 300, ty, 210, 94, LINK_DK, stroke=LINK, body="CC-Cores", size=18)
sh.arrow(s.s, t, tx + 232, ty + 30, tx + 300, ty + 30, color=WHITE, width=1.5)
sh.arrow(s.s, t, tx + 300, ty + 66, tx + 232, ty + 66, color=WHITE, width=1.5)
label(s.s, tx + 232, ty + 2, 70, 22, "trigger", size=11, color=MUTED)
label(s.s, tx + 232, ty + 72, 70, 22, "done", size=11, color=MUTED)
box(s.s, tx, ty + 150, 510, 52, YEL_DK, stroke=YEL, body="DMA engines", size=17)
sh.arrow(s.s, t, tx + 115, ty + 94, tx + 115, ty + 150, color=GRN_HI, width=2)
sh.arrow(s.s, t, tx + 405, ty + 94, tx + 405, ty + 150, color=LINK, width=2)
box(s.s, tx, ty + 222, 245, 40, GREY, stroke=EDGE, body="HBM / SBUF", size=14)
box(s.s, tx + 265, ty + 222, 245, 40, GREY, stroke=EDGE, body="NeuronLink / EFA", size=14)
box(s.s, 665, 456, 572, 36, YEL_DK, stroke=YEL, body="Shared: DMA engines", size=16, radius=0.3)

d.save(OUT)
sys.path.insert(0, "/Users/akazawt/works/data-science/claudecode/.claude/skills/review-html/packs/zenn-doc/templates")
import strip_branding  # noqa: E402
strip_branding.main(OUT, OUT)
print(OUT)
