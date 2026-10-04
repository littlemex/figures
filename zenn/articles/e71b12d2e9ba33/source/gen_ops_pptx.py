"""gen_ops_pptx.py — Collective Communication の 5 演算を、送り手 (rank 0 が誰に何を送るか) と
受け手 (rank 0 が誰から何を受け取るか) の 2 つの見方で並べる pptx 1 枚。

    /usr/bin/python3 gen_ops_pptx.py <out.pptx>

4 ランクの例。各ランクの入力は色で区別し、"1:2" は「rank 1 の入力の 2 番目の切れ端」を表す。
AllGather は自分の切れ端を全員へ配り (送り手)、全員の切れ端を集める (受け手)。
ReduceScatter と All-to-All は、自分の入力を切り分けて j 番目を rank j へ送る (scatter の形)。
ReduceScatter は受け手が足し合わせ、All-to-All は並べるだけ。AllReduce は ReduceScatter の後に AllGather。
Permute は輪の隣へ送るだけ。1 ランクが送る量は Ring で数えた値 (C は 1 ランク分の切れ端)。
図の文字は英語。根拠は neuron-runtime/about/collectives.html と explore/intranode-collective-comm.html。
"""
import sys

sys.path.insert(0, "/Users/akazawt/works/data-science/tools/pptxkit")
from pptx.dml.color import RGBColor  # noqa: E402
from pptx.enum.shapes import MSO_SHAPE  # noqa: E402

from pptxkit import Deck, shapes as sh  # noqa: E402

TPL = "/Users/akazawt/works/data-science/tools/pptxkit/templates/aws-dark-2024.pptx"
OUT = sys.argv[1] if len(sys.argv) > 1 else "ops.pptx"

RANK = [RGBColor(40, 150, 240), RGBColor(240, 110, 150), RGBColor(60, 200, 170), RGBColor(240, 190, 60)]
DARK = [RGBColor(16, 50, 82), RGBColor(80, 30, 48), RGBColor(18, 66, 56), RGBColor(80, 62, 16)]
CHIP = RGBColor(22, 26, 30)
EDGE = RGBColor(70, 74, 80)
MUTED = RGBColor(150, 150, 158)
WHITE = RGBColor(245, 245, 248)
GRN_HI = RGBColor(59, 209, 111)

d = Deck(TPL, check_layout=False)
t = d.theme
s = d.slide()


def label(x, y, w, h, body, size=14, color=WHITE, bold=False, align="center"):
    return sh.textbox(s.s, t, x, y, w, h, body, size=size, color=color, bold=bold, align=align,
                      anchor="middle")


def chip(x, y, r, text, w=46, h=24, size=11):
    sp = sh.rect(s.s, t, x, y, w, h, fill=DARK[r], stroke=RANK[r], sw=1.25,
                 shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.25)
    sh._write(sp.text_frame, t, text, size, WHITE, True, "center", "middle")
    tf = sp.text_frame
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = sh.U(1)


def node(x, y, r):
    sp = sh.rect(s.s, t, x, y, 44, 26, fill=CHIP, stroke=RANK[r], sw=2,
                 shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.3)
    sh._write(sp.text_frame, t, f"R{r}", 11, RANK[r], True, "center", "middle")
    tf = sp.text_frame
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = sh.U(0)


def arrow(x1, y1, x2, y2, r):
    sh.arrow(s.s, t, x1, y1, x2, y2, color=RANK[r], width=1.5)


label(43, 22, 1190, 40, "Two views of each collective (4 ranks, logical)", size=28, bold=True, align="left")
label(43, 60, 1190, 24, "Who ends up with whose data. Ring and other algorithms relay and reduce on the way.  \"1:2\" = chunk 2 of rank 1", size=14, color=MUTED,
      align="left")
COLS = (250, 650, 1050)
label(COLS[0] - 20, 92, 380, 26, "What R0 sends", size=15, color=GRN_HI)
label(COLS[1] - 20, 92, 380, 26, "What R0 receives", size=15, color=GRN_HI)
label(COLS[2] - 20, 92, 220, 26, "Data each rank owes", size=15, color=GRN_HI)

ROWS = [
    ("AllGather", "copy", "(N-1) C"),
    ("ReduceScatter", "scatter + reduce", "(N-1) C"),
    ("AllReduce", "RS then AG", "2 (N-1) C, Ring"),
    ("All-to-All", "scatter", "(N-1) C"),
    ("Permute", "ring", "whole tensor"),
]
y0, rh = 124, 104
for k, (name, kind, vol) in enumerate(ROWS):
    y = y0 + k * rh
    sh.rect(s.s, t, 43, y - 4, 1194, rh - 8, fill=CHIP if k % 2 == 0 else None, stroke=None)
    label(55, y, 170, 40, name, size=17, bold=True, align="left")
    label(55, y + 36, 170, 30, kind, size=12, color=MUTED, align="left")
    if name == "AllReduce":
        label(COLS[0], y + 20, 700, 46, "= ReduceScatter, then AllGather of the reduced chunks", size=17,
              color=WHITE, align="left")
        label(COLS[2] + 10, y + 22, 190, 40, vol, size=16, bold=True)
        continue
    # sender view: R0 on the left, targets on the right
    sx = COLS[0]
    node(sx, y + 30, 0)
    targets = [1, 2, 3] if name != "Permute" else [1]
    for i, r in enumerate(targets):
        ty = y + 4 + i * 30 if name != "Permute" else y + 30
        node(sx + 250, ty, r)
        arrow(sx + 40, y + 43, sx + 250, ty + 13, 0)
        if name in ("AllGather",):
            text = "0:0"
        elif name in ("ReduceScatter", "All-to-All"):
            text = f"0:{r}"
        elif name == "AllReduce":
            text = f"0:{r} / sum:0"
        else:
            text = "0:*"
        chip(sx + 110 if name != "AllReduce" else sx + 92, ty + 1 if name != "Permute" else ty + 1, 0,
             text, w=46 if name != "AllReduce" else 82)
    # receiver view: sources on the left, R0 on the right
    rx = COLS[1]
    node(rx + 250, y + 30, 0)
    sources = [1, 2, 3] if name != "Permute" else [3]
    for i, r in enumerate(sources):
        sy = y + 4 + i * 30 if name != "Permute" else y + 30
        node(rx, sy, r)
        arrow(rx + 40, sy + 13, rx + 250, y + 43, r)
        if name == "AllGather":
            text = f"{r}:0"
        elif name in ("ReduceScatter", "All-to-All"):
            text = f"{r}:0"
        elif name == "AllReduce":
            text = f"{r}:0 / sum:{r}"
        else:
            text = "3:*"
        chip(rx + 110 if name != "AllReduce" else rx + 92, sy + 1, r, text,
             w=46 if name != "AllReduce" else 82)
    result = {"AllGather": "keeps all 4", "ReduceScatter": "reduces chunk 0",
              "AllReduce": "full result", "All-to-All": "places chunk 0s", "Permute": "takes 3's data"}[name]
    label(rx + 296, y + 30, 110, 26, result, size=11, color=MUTED, align="left")
    label(COLS[2] + 10, y + 22, 190, 40, vol, size=16, bold=True)

d.save(OUT)
sys.path.insert(0, "/Users/akazawt/works/data-science/claudecode/.claude/skills/review-html/packs/zenn-doc/templates")
import strip_branding  # noqa: E402
strip_branding.main(OUT, OUT)
print(OUT)
