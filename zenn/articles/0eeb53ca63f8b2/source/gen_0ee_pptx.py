import sys

sys.path.insert(0, "/Users/akazawt/works/data-science/tools/pptxkit")
from pptx.dml.color import RGBColor  # noqa: E402
from pptx.enum.dml import MSO_LINE_DASH_STYLE  # noqa: E402
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE  # noqa: E402
from pptx.util import Pt  # noqa: E402

from pptxkit import Deck, shapes as sh  # noqa: E402

TPL = "/Users/akazawt/works/data-science/tools/pptxkit/templates/aws-dark-2024.pptx"
OUT = sys.argv[1] if len(sys.argv) > 1 else "0ee-figures.pptx"

GRN = RGBColor(36, 160, 70)
GRN_HI = RGBColor(59, 209, 111)
GRN_DK = RGBColor(18, 64, 34)
CHIP = RGBColor(22, 26, 30)
EDGE = RGBColor(70, 74, 80)
GREY = RGBColor(52, 56, 64)
LINK = RGBColor(40, 150, 240)
LINK_DK = RGBColor(16, 44, 72)
YEL = RGBColor(251, 211, 50)
YEL_DK = RGBColor(80, 66, 14)
RED = RGBColor(220, 38, 88)
RED_DK = RGBColor(70, 16, 32)
MUTED = RGBColor(150, 150, 158)
WHITE = RGBColor(245, 245, 248)

d = Deck(TPL, check_layout=False)
t = d.theme


def label(s, x, y, w, h, body, size=16, color=WHITE, bold=False, align="center"):
    return sh.textbox(s, t, x, y, w, h, body, size=size, color=color, bold=bold, align=align,
                      anchor="middle")


def box(s, x, y, w, h, fill, stroke=None, body="", size=15, color=WHITE, radius=0.1,
        dash=None, sw=1.25):
    sp = sh.rect(s, t, x, y, w, h, fill=fill, stroke=stroke, sw=sw, dash=dash,
                 shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=radius)
    if body:
        sh._write(sp.text_frame, t, body, size, color, True, "center", "middle")
        tf = sp.text_frame
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = sh.U(2)
    return sp


def arrow(s, x1, y1, x2, y2, color=WHITE, w=2.0):
    sh.arrow(s, t, x1, y1, x2, y2, color=color, width=w)


def seg(s, x1, y1, x2, y2, color=MUTED, w=1.5):
    c = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, sh.U(x1), sh.U(y1), sh.U(x2), sh.U(y2))
    c.shadow.inherit = False
    c.line.color.rgb = color
    c.line.width = Pt(w)


def title(s, head, sub):
    label(s, 43, 26, 1190, 46, head, size=30, bold=True, align="left")
    label(s, 43, 72, 1190, 26, sub, size=16, color=MUTED, align="left")


def legend(s, y=596):
    box(s, 43, y + 4, 24, 18, *CPU)
    label(s, 74, y, 80, 26, "CPU", size=13, color=MUTED, align="left")
    box(s, 150, y + 4, 24, 18, *NEU)
    label(s, 181, y, 160, 26, "NeuronCores", size=13, color=MUTED, align="left")


CPU = (GREY, EDGE)
PUR = RGBColor(150, 110, 230)
PUR_DK = RGBColor(44, 30, 74)
APP = (GREY, EDGE)
OFI = (LINK_DK, LINK)
RDMA = (PUR_DK, PUR)
KERN = (YEL_DK, YEL)
NITRO = (GRN_DK, GRN_HI)


def legend(s, items, y=596):
    x = 43
    for name, col in items:
        box(s, x, y + 4, 24, 18, *col)
        label(s, x + 31, y, 200, 26, name, size=13, color=MUTED, align="left")
        x += 40 + 9 * len(name) + 40


# ------------------------------------------------------------------ 1. the map
s = d.slide()
title(s.s, "From NCCL to the wire", "Each layer has one job.  Data skips the kernel")
box(s.s, 43, 118, 200, 62, *APP, body="NCCL", size=16)
box(s.s, 263, 118, 200, 62, *APP, body="MPI", size=16)
label(s.s, 480, 118, 420, 62, "collectives, point-to-point", size=14, color=MUTED, align="left")
box(s.s, 43, 194, 200, 62, *OFI, body="aws-ofi-nccl", size=15)
label(s.s, 480, 194, 420, 62, "NCCL net plugin -> libfabric", size=14, color=MUTED, align="left")
arrow(s.s, 363, 180, 363, 270, color=LINK, w=1.5)
rows = [
    ("libfabric  EFA provider", "efa fabric: ordering, large messages", OFI),
    ("rdma-core  libibverbs + efa", "queues and doorbells in user space", RDMA),
]
y = 270
for name, job, col in rows:
    box(s.s, 43, y, 420, 62, *col, body=name, size=16)
    label(s.s, 480, y, 420, 62, job, size=14, color=MUTED, align="left")
    y += 76
# kernel side box
box(s.s, 900, 372, 337, 56, *KERN, body="efa.ko  (kernel driver)", size=16)
label(s.s, 900, 430, 337, 24, "setup: queues, memory registration", size=12, color=MUTED)
seg(s.s, 463, 400, 900, 400, color=YEL, w=1.5)
label(s.s, 640, 402, 240, 24, "control path (syscall)", size=12, color=YEL)
# device
box(s.s, 43, 470, 1194, 64, *NITRO, body="Nitro card:  EFA device  +  SRD  (multipath, retransmit, congestion control)", size=16, sw=2)
arrow(s.s, 253, 408, 253, 470, color=GRN_HI, w=2.5)
label(s.s, 263, 426, 360, 30, "data path: no syscall, no kernel", size=13, color=GRN_HI, align="left")
arrow(s.s, 1068, 454, 1068, 470, color=YEL)
label(s.s, 43, 540, 1194, 26, "spine-leaf Ethernet fabric", size=14, color=MUTED)
legend(s.s, [("app", APP), ("libfabric", OFI), ("rdma-core", RDMA), ("kernel", KERN), ("Nitro hardware", NITRO)])

# ------------------------------------------------------------------ 2. Nitro
s = d.slide()
title(s.s, "Nitro: I/O moves off the host", "The host CPU runs the workload; cards run network and storage")
box(s.s, 43, 130, 520, 232, CHIP, stroke=EDGE, radius=0.04)
label(s.s, 63, 136, 480, 30, "EC2 instance (guest)", size=16, bold=True, align="left")
box(s.s, 43, 372, 520, 70, CHIP, stroke=GRN, radius=0.06)
box(s.s, 73, 180, 460, 70, *APP, body="host CPU / GPU / Trainium  (workload)", size=15)
box(s.s, 73, 262, 460, 40, *KERN, body="ENA / EFA driver", size=14)
box(s.s, 73, 310, 460, 40, *KERN, body="NVMe driver", size=14)
box(s.s, 73, 382, 460, 50, *NITRO, body="host: Nitro Hypervisor (CPU and memory only)", size=14)
box(s.s, 683, 130, 554, 300, CHIP, stroke=GRN, radius=0.04)
label(s.s, 703, 136, 500, 30, "Nitro cards (PCIe)", size=16, bold=True, color=GRN_HI, align="left")
box(s.s, 713, 180, 494, 70, *NITRO, body="Card for VPC: ENA, EFA, SRD, encryption", size=15)
box(s.s, 713, 270, 494, 60, *NITRO, body="Card for EBS: NVMe, encryption", size=15)
box(s.s, 713, 350, 494, 60, *NITRO, body="Nitro Controller", size=15)
arrow(s.s, 533, 282, 713, 215, color=WHITE)
arrow(s.s, 533, 330, 713, 300, color=WHITE)
label(s.s, 560, 232, 140, 22, "SR-IOV VF", size=12, color=MUTED)
label(s.s, 560, 322, 140, 22, "SR-IOV VF", size=12, color=MUTED)
legend(s.s, [("workload", APP), ("guest driver", KERN), ("Nitro", NITRO)])

# ------------------------------------------------------------------ 3. loss recovery
s = d.slide()
title(s.s, "When a lost packet needs a timeout", "TCP waits on the same path; SRD resends fast, and reroutes if the path failed")
SC = 9.0
def lane(y, name, col, items):
    label(s.s, 43, y, 120, 60, name, size=18, bold=True, color=col, align="left")
    x = 170
    for w, txt, fill in items:
        box(s.s, x, y, w, 60, fill[0], stroke=fill[1], body=txt, size=14, radius=0.06)
        x += w + 6
lane(170, "TCP", WHITE, [(90, "send", APP), (560, "wait for timeout  (min RTO 50 ms in the paper)", (RED_DK, RED)), (170, "resend, same path", APP)])
lane(300, "SRD", GRN_HI, [(90, "send", NITRO), (260, "fast resend, reroute if failed", NITRO)])

# ------------------------------------------------------------------ 4. how SRD picks paths
s = d.slide()
title(s.s, "How SRD uses paths on ordinary switches", "Switches still hash; the sender changes what they hash")
box(s.s, 43, 150, 300, 90, *NITRO, body="sender's Nitro card\nRTT per path", size=16)
arrow(s.s, 343, 195, 413, 195)
box(s.s, 413, 150, 330, 90, CHIP, stroke=EDGE, body="packet header\n(encapsulation varies)", size=16)
arrow(s.s, 743, 195, 813, 195)
box(s.s, 813, 150, 424, 90, GREY, stroke=EDGE, body="switch ECMP hash -> spine 1..4", size=16)
for i, (lab, col) in enumerate([("S1", LINK), ("S2", LINK), ("S3", GRN_HI), ("S4", YEL)]):
    box(s.s, 853 + i * 96, 290, 80, 50, CHIP, stroke=col, body=lab, size=15)
    arrow(s.s, 1025, 240, 893 + i * 96, 290, color=col, w=1.25)
for i, tx in enumerate(["high RTT -> fewer packets", "path failed -> reroute", "rate + inflight limit per connection"]):
    box(s.s, 43 + i * 405, 420, 385, 70, CHIP, stroke=GRN_HI, body=tx, size=15)
legend(s.s, [("Nitro hardware", NITRO)])

# ------------------------------------------------------------------ 5. results
s = d.slide()
title(s.s, "Measured: SRD stays near ideal", "Flow completion time (FCT) relative to ideal.  Shalev et al., IEEE Micro 2020")
import math
X0, X1 = 330, 1200
def xof(v):
    return X0 + (X1 - X0) * math.log10(v) / 2.0  # 1x .. 100x
for v in (1, 2, 5, 10, 20, 50, 100):
    seg(s.s, xof(v), 140, xof(v), 520, color=GREY, w=0.75)
    label(s.s, xof(v) - 30, 524, 60, 24, f"{v}x", size=12, color=MUTED)
def bar(y, name, lo, hi, col, txt):
    label(s.s, 43, y, 280, 44, name, size=14, align="left")
    x1, x2 = xof(lo), xof(hi)
    box(s.s, x1, y + 6, max(x2 - x1, 10), 32, col[0], stroke=col[1], radius=0.2)
    label(s.s, x2 + 8, y, 260, 44, txt, size=13, color=col[1], align="left")
label(s.s, 43, 140, 600, 26, "48 flows into one 100 Gb/s link: max FCT", size=15, bold=True, color=YEL, align="left")
bar(172, "SRD", 1.0, 1.08, NITRO, "close to ideal")
bar(222, "TCP", 3, 20, (RED_DK, RED), "3-20x")
label(s.s, 43, 300, 700, 26, "8 server pairs across racks, uplinks 50% busy", size=15, bold=True, color=YEL, align="left")
bar(332, "SRD median", 1.0, 1.15, NITRO, "+15%")
bar(382, "TCP mean", 1.0, 1.5, (RED_DK, RED), "+50%")
bar(432, "TCP tail", 10, 100, (RED_DK, RED), "10-100x")

# ------------------------------------------------------------------ 6. control vs data path (details)
s = d.slide()
title(s.s, "Setup through the kernel, data straight to the card", "Setup when resources are created; the data path on every send")
box(s.s, 43, 130, 560, 400, CHIP, stroke=YEL, radius=0.04)
label(s.s, 63, 136, 500, 30, "setup  (when needed)", size=17, bold=True, color=YEL, align="left")
for i, tx in enumerate(["create queue pair, completion queue", "register memory: pin pages, get a key", "map queues and doorbell into user space"]):
    box(s.s, 73, 186 + i * 80, 500, 60, *KERN, body=tx, size=14)
label(s.s, 63, 440, 520, 70, "libibverbs -> /dev/infiniband/uverbs -> efa.ko\n-> admin queue -> device", size=13, color=MUTED, align="left")
box(s.s, 677, 130, 560, 400, CHIP, stroke=GRN_HI, radius=0.04)
label(s.s, 697, 136, 500, 30, "data transfer  (every message)", size=17, bold=True, color=GRN_HI, align="left")
for i, tx in enumerate(["write work request into the mapped queue", "ring the doorbell (MMIO write)", "device DMAs from registered memory", "poll completion queue in user space"]):
    box(s.s, 707, 186 + i * 80, 500, 60, *NITRO, body=tx, size=14)
legend(s.s, [("kernel", KERN), ("user space + device", NITRO)])

# ------------------------------------------------------------------ 7. what the EFA provider adds (details)
s = d.slide()
title(s.s, "What the efa fabric adds on top of SRD", "libfabric EFA provider.  efa-direct skips these protocols")
box(s.s, 43, 130, 360, 380, *NITRO, body="SRD (device)\n\nreliable\nunordered\nsmall datagrams", size=17, radius=0.06)
items = [("ordering", "reorder window per sender  (FI_EFA_RECVWIN_SIZE)"),
         ("large messages", "eager / medium / long (CTS) / read protocols"),
         ("same node", "shm provider instead of the NIC"),
         ("GPU / Trainium memory", "HMEM: DMA to accelerator memory (peer to peer)")]
y = 130
for k, v in items:
    box(s.s, 463, y, 250, 80, *OFI, body=k, size=15)
    label(s.s, 723, y, 514, 80, v, size=14, align="left")
    arrow(s.s, 403, 320, 463, y + 40, color=LINK, w=1.25)
    y += 95
legend(s.s, [("libfabric EFA provider", OFI), ("device", NITRO)])

# ------------------------------------------------------------------ 8. who tells the sender to slow down
s = d.slide()
title(s.s, "Who tells the sender to slow down?", "Credits: the receiver says how much it can take.  SRD: the sender infers it from ACKs")
for k, (head, sub) in enumerate([("Credit-based flow control", "e.g. InfiniBand link level"), ("SRD congestion control", "Nitro card, per connection")]):
    x = 43 + k * 612
    box(s.s, x, 120, 582, 430, CHIP, stroke=EDGE, radius=0.03)
    label(s.s, x + 16, 128, 550, 30, head, size=18, bold=True, align="left")
    label(s.s, x + 16, 158, 550, 24, sub, size=13, color=MUTED, align="left")
    box(s.s, x + 30, 210, 180, 70, *NITRO if k else CPU, body="Sender", size=17)
    box(s.s, x + 372, 210, 180, 70, *CPU, body="Receiver", size=17)
    arrow(s.s, x + 210, 232, x + 372, 232, color=WHITE)
    label(s.s, x + 210, 202, 162, 26, "data", size=13, color=MUTED)
    arrow(s.s, x + 372, 262, x + 210, 262, color=YEL)
    label(s.s, x + 210, 266, 162, 26, "credits: N free" if k == 0 else "ACK", size=13, color=YEL)
    if k == 0:
        box(s.s, x + 30, 330, 522, 60, YEL_DK, stroke=YEL, body="credits = 0  ->  sender waits", size=16)
        label(s.s, x + 30, 400, 522, 30, "receiver decides, explicitly", size=15, color=MUTED)
    else:
        box(s.s, x + 30, 330, 522, 60, GRN_DK, stroke=GRN_HI, body="ACK timing -> delivery rate,  RTT per path", size=15)
        box(s.s, x + 30, 410, 522, 60, YEL_DK, stroke=YEL, body="sender lowers its own rate and inflight", size=16)
        label(s.s, x + 30, 480, 522, 30, "sender decides, by inference", size=15, color=MUTED)

# ------------------------------------------------------------------ 9. SRD congestion decision
s = d.slide()
title(s.s, "How SRD decides it is congested", "Shalev et al., IEEE Micro 2020")
label(s.s, 43, 108, 400, 28, "Inputs", size=16, bold=True, color=GRN_HI, align="left")
for k, txt in enumerate(["RTT of each path", "delivery rate (ACK timing)", "own recent send rate"]):
    box(s.s, 43, 142 + k * 70, 300, 56, *NITRO, body=txt, size=15)
# branch 1: one path
box(s.s, 420, 142, 360, 70, CHIP, stroke=EDGE, body="RTT up on one path", size=16)
arrow(s.s, 343, 170, 420, 177)
arrow(s.s, 780, 177, 840, 177)
box(s.s, 840, 142, 397, 70, LINK_DK, stroke=LINK, body="that path is busy\n-> move packets to other paths", size=15)
# branch 2: whole connection
box(s.s, 420, 270, 360, 110, CHIP, stroke=EDGE, body="RTT up on most paths\nor\ndelivery rate < send rate", size=16)
arrow(s.s, 343, 240, 420, 300)
arrow(s.s, 343, 182, 420, 285)
arrow(s.s, 343, 310, 420, 320)
arrow(s.s, 780, 325, 840, 325)
box(s.s, 840, 270, 397, 110, YEL_DK, stroke=YEL, body="whole connection is congested\n(e.g. incast)\n-> lower rate limit and inflight limit", size=15)
box(s.s, 43, 440, 1194, 60, CHIP, stroke=GRN_HI, body="goal: fair share with minimum bytes in flight  (no queue build-up, no drops as the signal)", size=16)

d.save(OUT)
sys.path.insert(0, "/Users/akazawt/works/data-science/claudecode/.claude/skills/review-html/packs/zenn-doc/templates")
import strip_branding  # noqa: E402
strip_branding.main(OUT, OUT)
