"""gen_nvshmem_pptx.py -- NVSHMEM の上級編の図を、ネイティブ shape の pptx 5 枚で組む。

    /usr/bin/python3 gen_nvshmem_pptx.py <out.pptx>

1. PGAS と対称ヒープ: PE ごとの GPU メモリに、同じ symmetric address で dst が並ぶ
2. CPU が通信を指揮する場合と、カーネルの中から put する場合の時間の流れ
3. ビルド: -rdc=true と、libnvshmem_device.a (デバイスリンク) / libnvshmem_host.so
4. 完了の待ち方: GPU が出した put と、CPU 側の barrier_all が順序づける範囲
5. 通信路: ノード内は P2P、ノード間は NVSHMEM_REMOTE_TRANSPORT (既定 ibrc、EFA は libfabric + efa)
根拠: NVSHMEM API ドキュメント (docs.nvidia.com/nvshmem/api/latest、2026-10-06 参照) と記事の実機ログ (p4d, g4dn)。
図の文字は英語。色の役割は ZN-25 (緑 = 自分が書く、青 = NVSHMEM と CUDA が受け持つ、黄 = 注意)。
"""
import sys

sys.path.insert(0, "/Users/akazawt/works/data-science/tools/pptxkit")
from pptx.dml.color import RGBColor  # noqa: E402
from pptx.enum.dml import MSO_LINE_DASH_STYLE  # noqa: E402
from pptx.enum.shapes import MSO_SHAPE  # noqa: E402

from pptxkit import Deck, shapes as sh  # noqa: E402

TPL = "/Users/akazawt/works/data-science/tools/pptxkit/templates/aws-dark-2024.pptx"
OUT = sys.argv[1] if len(sys.argv) > 1 else "nvshmem-figures.pptx"

GRN = RGBColor(36, 160, 70)
GRN_HI = RGBColor(59, 209, 111)
GRN_DK = RGBColor(18, 64, 34)
CHIP = RGBColor(22, 26, 30)
EDGE = RGBColor(70, 74, 80)
GREY = RGBColor(52, 56, 64)
LINK = RGBColor(40, 150, 240)
LINK_DK = RGBColor(16, 44, 72)
YEL = RGBColor(251, 211, 50)
YEL_DK = RGBColor(80, 66, 14)
MUTED = RGBColor(150, 150, 158)
WHITE = RGBColor(245, 245, 248)

d = Deck(TPL, check_layout=False)
t = d.theme


def label(s, x, y, w, h, body, size=16, color=WHITE, bold=False, align="center"):
    return sh.textbox(s, t, x, y, w, h, body, size=size, color=color, bold=bold, align=align,
                      anchor="middle")


def box(s, x, y, w, h, fill, stroke=None, body="", size=15, color=WHITE, radius=0.1,
        dash=None, sw=1.25):
    sp = sh.rect(s, t, x, y, w, h, fill=fill, stroke=stroke, sw=sw, dash=dash,
                 shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=radius)
    if body:
        sh._write(sp.text_frame, t, body, size, color, True, "center", "middle")
        tf = sp.text_frame
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = sh.U(2)
    return sp


def title(s, head, sub):
    label(s, 43, 26, 1190, 46, head, size=30, bold=True, align="left")
    label(s, 43, 72, 1190, 26, sub, size=16, color=MUTED, align="left")


# ------------------------------------------------------------------ 1. PGAS and the symmetric heap
s = d.slide()
title(s.s, "One address space, split across GPUs",
      "Blue: GPU memory of each PE.  Green: dst from nvshmem_malloc, one symmetric address for every PE")
W, GAP, X0 = 360, 57, 43
for i in range(3):
    x = X0 + i * (W + GAP)
    label(s.s, x, 116, W, 30, f"PE {i}  (GPU {i})", size=18, bold=True)
    box(s.s, x, 150, W, 330, LINK_DK, stroke=LINK, radius=0.05)
    label(s.s, x + 16, 160, W - 32, 26, "private memory", size=14, color=MUTED, align="left")
    box(s.s, x + 20, 196, W - 40, 80, CHIP, stroke=EDGE, body="cudaMalloc, stack, ...", size=14, color=MUTED)
    label(s.s, x + 16, 292, W - 32, 26, "symmetric heap", size=14, color=MUTED, align="left")
    box(s.s, x + 20, 326, W - 40, 130, CHIP, stroke=EDGE, radius=0.06)
    box(s.s, x + 40, 360, W - 80, 56, GRN_DK, stroke=GRN, body="dst", size=18)
for i in range(2):
    xa = X0 + i * (W + GAP) + W - 40
    xb = X0 + (i + 1) * (W + GAP) + 40
    sh.arrow(s.s, t, xa, 380, xb, 380, color=WHITE, width=2.5)
    sh.arrow(s.s, t, xb, 396, xa, 396, color=WHITE, width=2.5)
    label(s.s, xa - 20, 330, xb - xa + 40, 26, "put / get", size=14)
label(s.s, 43, 500, 1194, 30, "local: fast        remote: NVLink, PCIe or network", size=15, color=MUTED)

# ------------------------------------------------------------------ 2. who starts the transfer
s = d.slide()
title(s.s, "Who starts the transfer", "Green: work in a CUDA kernel.  Grey: work on the CPU.  Blue: the transfer")
label(s.s, 43, 120, 600, 30, "CPU drives it (MPI + CUDA)", size=19, bold=True, align="left")
y = 160
segs = [("kernel: compute", GRN_DK, GRN, 250), ("kernel ends", GREY, EDGE, 150), ("CPU: MPI send", GREY, EDGE, 190),
        ("transfer", LINK_DK, LINK, 180), ("CPU: launch", GREY, EDGE, 160), ("kernel: compute", GRN_DK, GRN, 230)]
x = 43
for nm, fill, edge, w in segs:
    box(s.s, x, y, w - 8, 64, fill, stroke=edge, body=nm, size=14)
    x += w
label(s.s, 43, 236, 1194, 28, "every exchange returns to the CPU", size=15, color=MUTED, align="left")
label(s.s, 43, 300, 600, 30, "The kernel drives it (NVSHMEM)", size=19, bold=True, align="left")
y = 340
box(s.s, 43, y, 760, 64, GRN_DK, stroke=GRN, body="kernel: compute  +  nvshmem_int_p  +  compute", size=15)
box(s.s, 330, y + 84, 260, 54, LINK_DK, stroke=LINK, body="transfer", size=14)
sh.arrow(s.s, t, 380, y + 66, 380, y + 82, color=LINK, width=2)
label(s.s, 43, y + 150, 1194, 28, "no return to the CPU per exchange; one kernel launch", size=15, color=MUTED, align="left")

# ------------------------------------------------------------------ 3. build
s = d.slide()
title(s.s, "Why -rdc=true", "Green: your file.  Blue: the two NVSHMEM libraries")
box(s.s, 43, 150, 280, 120, GRN_DK, stroke=GRN, body="ring.cu\nkernel calls\nnvshmem_int_p", size=16)
box(s.s, 43, 330, 280, 90, GRN_DK, stroke=GRN, body="ring.cu\nhost code", size=16)
box(s.s, 420, 196, 300, 74, LINK_DK, stroke=LINK, body="libnvshmem_device.a  (static)", size=15)
box(s.s, 420, 366, 300, 54, LINK_DK, stroke=LINK, body="libnvshmem_host.so", size=15)
box(s.s, 830, 150, 230, 120, YEL_DK, stroke=YEL, body="device link\n(needs -rdc=true)", size=16)
box(s.s, 830, 330, 230, 90, GREY, stroke=EDGE, body="host link", size=16)
sh.arrow(s.s, t, 323, 172, 826, 172, color=WHITE, width=2)
sh.arrow(s.s, t, 720, 233, 826, 233, color=WHITE, width=2)
sh.arrow(s.s, t, 323, 346, 826, 346, color=WHITE, width=2)
sh.arrow(s.s, t, 720, 393, 826, 393, color=WHITE, width=2)
box(s.s, 1110, 230, 127, 110, CHIP, stroke=EDGE, body="./ring", size=18)
sh.arrow(s.s, t, 1060, 210, 1106, 260, color=WHITE, width=2)
sh.arrow(s.s, t, 1060, 375, 1106, 310, color=WHITE, width=2)
label(s.s, 43, 470, 1194, 30, "-gencode=arch=compute_80,code=sm_80  must match the GPU (A100: 80)", size=16,
      color=MUTED)

# ------------------------------------------------------------------ 4. completion
s = d.slide()
title(s.s, "When is the put done?", "Blue: CPU-side calls.  Yellow: the GPU-side quiet the spec asks for")
label(s.s, 43, 118, 200, 30, "GPU (PE 0)", size=17, bold=True, align="left")
label(s.s, 43, 300, 200, 30, "CPU (PE 0)", size=17, bold=True, align="left")
box(s.s, 220, 112, 330, 60, GRN_DK, stroke=GRN, body="nvshmem_int_p: starts the put", size=15)
label(s.s, 220, 176, 330, 26, "may return before delivery", size=14, color=MUTED)
box(s.s, 590, 112, 250, 60, YEL_DK, stroke=YEL, body="nvshmem_quiet()", size=15, dash=MSO_LINE_DASH_STYLE.DASH)
label(s.s, 590, 176, 250, 26, "GPU-side: completes it", size=14, color=MUTED)
box(s.s, 880, 112, 357, 60, GREY, stroke=EDGE, body="kernel ends", size=15)
box(s.s, 220, 294, 300, 60, LINK_DK, stroke=LINK, body="cudaDeviceSynchronize()", size=15)
label(s.s, 220, 358, 300, 26, "waits for the kernel", size=14, color=MUTED)
box(s.s, 560, 294, 300, 60, LINK_DK, stroke=LINK, body="nvshmem_barrier_all()", size=15)
label(s.s, 560, 358, 300, 26, "orders CPU-issued ops", size=14, color=MUTED)
box(s.s, 900, 294, 337, 60, LINK_DK, stroke=LINK, body="cudaMemcpy dst -> host", size=15)
sh.arrow(s.s, t, 520, 324, 556, 324, color=WHITE, width=2)
sh.arrow(s.s, t, 860, 324, 896, 324, color=WHITE, width=2)
sh.arrow(s.s, t, 1058, 172, 370, 290, color=MUTED, width=1.5)
box(s.s, 43, 430, 1194, 70, CHIP, stroke=EDGE, radius=0.2,
    body="official example:  kernel on stream  ->  barrier_all_on_stream  ->  copy  ->  sync",
    size=15)

# ------------------------------------------------------------------ 5. transports
s = d.slide()
title(s.s, "Which path a put takes", "Blue: what happens by default.  Yellow: watch out")
box(s.s, 43, 120, 560, 380, CHIP, stroke=EDGE, radius=0.04)
label(s.s, 43, 128, 560, 32, "inside one node", size=19, bold=True)
box(s.s, 83, 180, 200, 70, GRN_DK, stroke=GRN, body="GPU 0", size=17)
box(s.s, 363, 180, 200, 70, GRN_DK, stroke=GRN, body="GPU 1", size=17)
sh.arrow(s.s, t, 283, 215, 359, 215, color=WHITE, width=3)
label(s.s, 83, 260, 480, 30, "P2P: NVLink or PCIe P2P", size=16, color=LINK)
box(s.s, 83, 310, 480, 60, LINK_DK, stroke=LINK, body="p4d (NVLink + NVSwitch): 56 / 56 pairs", size=15)
box(s.s, 83, 390, 480, 60, YEL_DK, stroke=YEL, body="g4dn (PCIe): 0 / 12 pairs  ->  init exits", size=15)
box(s.s, 677, 120, 560, 380, CHIP, stroke=EDGE, radius=0.04)
label(s.s, 677, 128, 560, 32, "across nodes", size=19, bold=True)
box(s.s, 717, 180, 480, 64, LINK_DK, stroke=LINK, body="NVSHMEM_REMOTE_TRANSPORT  (default: ibrc)", size=15)
label(s.s, 717, 250, 480, 30, "ibrc, ucx, libfabric, ibdevx, gpunetio, none", size=14, color=MUTED)
box(s.s, 717, 300, 480, 80, YEL_DK, stroke=YEL,
    body="EFA:  NVSHMEM_REMOTE_TRANSPORT=libfabric\nNVSHMEM_LIBFABRIC_PROVIDER=efa", size=14)
label(s.s, 717, 400, 480, 60, "seen on p4d:  \"init failed for remote transport: ibrc\"", size=14, color=MUTED)

d.save(OUT)
# ZN-24: テンプレートのロゴ、著作権表示、ヘッダー、フッターを必ず消してから使う
sys.path.insert(0, "/Users/akazawt/works/data-science/claudecode/.claude/skills/review-html/packs/zenn-doc/templates")
import strip_branding  # noqa: E402
strip_branding.main(OUT, OUT)
print(OUT)
