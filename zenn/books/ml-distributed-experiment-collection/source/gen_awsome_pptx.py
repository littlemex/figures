import sys

sys.path.insert(0, "/Users/akazawt/works/data-science/tools/pptxkit")
from pptx.dml.color import RGBColor  # noqa: E402
from pptx.enum.dml import MSO_LINE_DASH_STYLE  # noqa: E402
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE  # noqa: E402
from pptx.util import Pt  # noqa: E402

from pptxkit import Deck, shapes as sh  # noqa: E402

TPL = "/Users/akazawt/works/data-science/tools/pptxkit/templates/aws-dark-2024.pptx"
OUT = sys.argv[1] if len(sys.argv) > 1 else "awsome-figures.pptx"

GRN = RGBColor(36, 160, 70)
GRN_HI = RGBColor(59, 209, 111)
GRN_DK = RGBColor(18, 64, 34)
CHIP = RGBColor(22, 26, 30)
EDGE = RGBColor(70, 74, 80)
GREY = RGBColor(52, 56, 64)
LINK = RGBColor(40, 150, 240)
LINK_DK = RGBColor(16, 44, 72)
YEL = RGBColor(251, 211, 50)
YEL_DK = RGBColor(80, 66, 14)
RED = RGBColor(220, 38, 88)
RED_DK = RGBColor(70, 16, 32)
MUTED = RGBColor(150, 150, 158)
WHITE = RGBColor(245, 245, 248)

d = Deck(TPL, check_layout=False)
t = d.theme


def label(s, x, y, w, h, body, size=16, color=WHITE, bold=False, align="center"):
    return sh.textbox(s, t, x, y, w, h, body, size=size, color=color, bold=bold, align=align,
                      anchor="middle")


def box(s, x, y, w, h, fill, stroke=None, body="", size=15, color=WHITE, radius=0.1,
        dash=None, sw=1.25):
    sp = sh.rect(s, t, x, y, w, h, fill=fill, stroke=stroke, sw=sw, dash=dash,
                 shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=radius)
    if body:
        sh._write(sp.text_frame, t, body, size, color, True, "center", "middle")
        tf = sp.text_frame
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = sh.U(2)
    return sp


def arrow(s, x1, y1, x2, y2, color=WHITE, w=2.0):
    sh.arrow(s, t, x1, y1, x2, y2, color=color, width=w)


def seg(s, x1, y1, x2, y2, color=MUTED, w=1.5):
    c = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, sh.U(x1), sh.U(y1), sh.U(x2), sh.U(y2))
    c.shadow.inherit = False
    c.line.color.rgb = color
    c.line.width = Pt(w)


def title(s, head, sub):
    label(s, 43, 26, 1190, 46, head, size=30, bold=True, align="left")
    label(s, 43, 72, 1190, 26, sub, size=16, color=MUTED, align="left")




HOST = (GREY, EDGE)
CTR = (LINK_DK, LINK)
FIX = (GRN_DK, GRN_HI)
BAD = (RED_DK, RED)
NOTE = (YEL_DK, YEL)


def chain(s, xs, y, items, w=200, h=64, size=14):
    for i, (txt, col) in enumerate(items):
        box(s, xs[i], y, w, h, *col, body=txt, size=size)
        if i:
            arrow(s, xs[i - 1] + w, y + h / 2, xs[i], y + h / 2)


# 1. picotron: how a Slurm job reaches the GPUs
s = d.slide()
title(s.s, "From sbatch to NCCL over EFA", "Who starts what, and where each fix lands")
X = [43, 283, 523, 763, 1003]
chain(s.s, X, 150, [("sbatch script\n(host shell)", HOST), ("srun\n(one task per node)", HOST),
                    ("enroot container\n(pyxis)", CTR), ("torchrun\nsets RANK", CTR),
                    ("train.py\nNCCL", CTR)], w=234, size=14)
box(s.s, 1003, 270, 234, 60, *CTR, body="aws-ofi-nccl -> EFA", size=14)
arrow(s.s, 1120, 214, 1120, 270)
label(s.s, 43, 236, 234, 40, "pitfall (not in PR): PATH lacks\n/opt/slurm/bin on ParallelCluster", size=12, color=YEL)
label(s.s, 763, 236, 234, 40, "pitfall (not in PR): srun alone\ndoes not set RANK / WORLD_SIZE", size=12, color=YEL)
label(s.s, 43, 380, 600, 30, "Fixes in PR #1019", size=17, bold=True, color=GRN_HI, align="left")
FIXES = [("Dockerfile", "pin picotron to 59714b1"), ("create_config.py", "template path from __file__"),
         ("train.sbatch", "EFA provider, huge pages off,\nskip docker/lo/veth sockets"),
         ("READMEs", "config path, duplicate step")]
for k, (f, what) in enumerate(FIXES):
    x = 43 + k * 300
    box(s.s, x, 420, 284, 50, *FIX, body=f, size=15)
    label(s.s, x, 474, 284, 60, what, size=13, color=WHITE)

# 2. ddp: two failures behind one symptom
s = d.slide()
title(s.s, "DDP container training: what broke and what changed", "Issues #1038 and #1047")
chain(s.s, [43, 323, 603], 140, [("2.create-enroot-image.sh\npulls pytorch:latest", BAD),
                                 ("Dockerfile ignored\n-> no mlflow", BAD), ("build from Dockerfile\n(dockerd://)", FIX)], w=260)
chain(s.s, [43, 323, 603], 240, [("sbatch passes\n--use-mlflow", BAD), ("argparse defines\n--use_mlflow", BAD),
                                 ("--use_mlflow", FIX)], w=260)
label(s.s, 43, 330, 900, 30, "Intermittent failures on CPU nodes (#1047)", size=17, bold=True, color=YEL, align="left")
chain(s.s, [43, 323, 603, 883], 370, [("nproc_per_node=4\n8 ranks", BAD), ("some ranks: MNIST\nfile corrupted", BAD),
                                      ("other ranks: GLOO\nconnection closed", BAD), ("pin 2.10.0 +\nNPROC_PER_NODE knob", FIX)], w=260)
box(s.s, 43, 470, 1100, 50, *NOTE, body="open (hypothesis): concurrent MNIST download on shared storage  ->  follow-up", size=15)

# 3. ParallelCluster monitoring
s = d.slide()
title(s.s, "Monitoring ParallelCluster compute nodes", "PR #1043: same collector as HyperPod, a different install hook")
box(s.s, 43, 140, 300, 300, CHIP, stroke=EDGE, radius=0.04)
label(s.s, 63, 148, 260, 30, "Compute node", size=16, bold=True, align="left")
box(s.s, 63, 190, 260, 64, *FIX, body="OnNodeConfigured\ninstall-node-exporter.sh", size=14)
box(s.s, 63, 290, 260, 64, *HOST, body="node_exporter :9100\n(systemd, own user)", size=14)
arrow(s.s, 193, 254, 193, 290)
box(s.s, 450, 260, 300, 90, *CTR, body="Prometheus agent collector\nscrape 1m, timeout 30s", size=15)
arrow(s.s, 323, 322, 450, 305)
box(s.s, 860, 220, 330, 64, *HOST, body="Amazon Managed Prometheus", size=15)
box(s.s, 860, 320, 330, 64, *HOST, body="Grafana dashboards", size=15)
arrow(s.s, 750, 290, 860, 252)
arrow(s.s, 1025, 284, 1025, 320)
label(s.s, 450, 360, 300, 50, "no instance-type filter:\nCPU, GPU, Trainium alike", size=13, color=GRN_HI)
label(s.s, 43, 460, 1150, 40, "checksum verified, 3 retries, arm64 and amd64, skips if already running", size=14, color=MUTED, align="left")

# 4. HyperPod EKS: Lambda INIT failure
s = d.slide()
title(s.s, "helm-chart-installer: why import ssl failed", "PR #1206: one line removed")
box(s.s, 43, 140, 380, 300, CHIP, stroke=EDGE, radius=0.04)
label(s.s, 63, 148, 340, 30, "Lambda python3.12", size=16, bold=True, align="left")
box(s.s, 63, 190, 340, 60, *HOST, body="_ssl (needs OPENSSL_3.3.0)", size=15)
box(s.s, 63, 290, 160, 60, *HOST, body="runtime\nlibcrypto", size=14)
box(s.s, 243, 290, 160, 60, *BAD, body="layer\nlibcrypto (old)", size=14)
arrow(s.s, 233, 250, 323, 290, color=RED)
label(s.s, 63, 370, 340, 50, "LD_LIBRARY_PATH=/opt/python/lib\nputs the layer first", size=13, color=YEL)
box(s.s, 520, 140, 380, 300, CHIP, stroke=EDGE, radius=0.04)
label(s.s, 540, 148, 340, 30, "After the fix", size=16, bold=True, color=GRN_HI, align="left")
box(s.s, 540, 190, 340, 60, *HOST, body="_ssl", size=15)
box(s.s, 540, 290, 160, 60, *FIX, body="runtime\nlibcrypto", size=14)
box(s.s, 720, 290, 160, 60, *HOST, body="layer: git,\nhelm, kubectl", size=14)
arrow(s.s, 710, 250, 620, 290, color=GRN_HI)
box(s.s, 960, 190, 277, 160, *NOTE, body="same installer also in\naws/sagemaker-hyperpod-\ncluster-setup\n-> RFC #1205", size=14)

# 5. slime: MODEL_ARGS lost at a shell boundary
s = d.slide()
title(s.s, "Why MODEL_ARGS arrived empty", "ray job submit -- bash -c \"... ${MODEL_ARGS[@]} ...\"")
chain(s.s, [43, 323, 603, 883], 150, [("recipe builds\none string", HOST), ("Ray: list2cmdline\n+ /bin/sh -c", HOST),
                                      ("outer sh expands\nMODEL_ARGS: empty", BAD), ("train.py: hidden_size None", BAD)], w=260)
label(s.s, 43, 260, 900, 30, "Fix: follow SLIME's own launch pattern", size=17, bold=True, color=GRN_HI, align="left")
chain(s.s, [43, 323, 603], 300, [("ray job submit --\nlauncher.sh", FIX), ("same shell sources\nthe model script", FIX),
                                 ("argv tokens reach\ntrain.py", FIX)], w=260)
label(s.s, 43, 400, 900, 30, "Other fixes in the epic (#1164)", size=17, bold=True, color=YEL, align="left")
for k, txt in enumerate(["log level WARN -> warning\n(uvicorn)", "torch_memory_saver .so\nfor CUDA 13", "numpy<2\n(Megatron init)",
                         "30B MoE: mbridge,\nGPU-less validate_args"]):
    box(s.s, 43 + k * 300, 440, 284, 70, *NOTE, body=txt, size=14)

# 6. miles: the test case layout
s = d.slide()
title(s.s, "miles GRPO test case on EKS", "PR #1225: same shape as the slime test case")
box(s.s, 43, 140, 360, 340, CHIP, stroke=EDGE, radius=0.04)
label(s.s, 63, 148, 320, 30, "RayCluster", size=16, bold=True, align="left")
box(s.s, 63, 190, 320, 60, *HOST, body="head: coordinator (num-cpus 0)", size=14)
box(s.s, 63, 270, 320, 90, *CTR, body="GPU workers\nSGLang rollout + Megatron train", size=14)
box(s.s, 63, 380, 320, 70, *CTR, body="colocated or\ndisaggregated", size=14)
box(s.s, 480, 190, 300, 70, *HOST, body="reward service\n(CPU, FastAPI)", size=14)
arrow(s.s, 383, 315, 480, 225)
box(s.s, 480, 300, 300, 70, *HOST, body="FSx for Lustre\ncheckpoints, data", size=14)
arrow(s.s, 383, 335, 480, 335)
box(s.s, 480, 410, 300, 70, *HOST, body="BuildKit Job\n(in-cluster image)", size=14)
label(s.s, 860, 140, 377, 30, "Review: 18 findings", size=17, bold=True, color=YEL, align="left")
for k, txt in enumerate(["wrong namespace / head pod", "reward errors scored as 0.0", "docs vs pinned image",
                         "validation rows not reproducible"]):
    box(s.s, 860, 180 + k * 75, 377, 60, *NOTE, body=txt, size=14)

# 7. EKS CloudFormation split
s = d.slide()
title(s.s, "EKS GPU cluster as five templates", "PR #1269: a GPU node group can be deployed on its own")
box(s.s, 43, 150, 300, 70, *CTR, body="root\neks-gpu-cluster-deploy-all", size=14)
kids = [("prerequisites\nVPC, subnets, SG", HOST), ("node AMI\nImage Builder", HOST), ("cluster\nEKS control plane", HOST),
        ("GPU node group\naccess entry, plugins", FIX)]
for k, (txt, col) in enumerate(kids):
    x = 420 + k * 210
    box(s.s, x, 150, 196, 70, *col, body=txt, size=13)
seg(s.s, 343, 185, 420, 185, color=WHITE)
label(s.s, 420, 230, 820, 30, "created by URL from an S3 bucket", size=13, color=MUTED, align="left")
box(s.s, 43, 300, 300, 70, *HOST, body="existing EKS cluster", size=15)
arrow(s.s, 343, 335, 1050, 335, color=GRN_HI)
label(s.s, 420, 300, 600, 30, "standalone path: add GPU capacity to a cluster you already run", size=13, color=GRN_HI, align="left")
box(s.s, 1050, 300, 187, 70, *FIX, body="node group\nalone", size=14)
label(s.s, 43, 410, 1194, 30, "Bootstrap fails the stack unless every GPU node is Ready and advertises", size=15, color=WHITE, align="left")
box(s.s, 43, 450, 560, 60, *NOTE, body="nvidia.com/gpu = GPU count", size=15)
box(s.s, 640, 450, 560, 60, *NOTE, body="vpc.amazonaws.com/efa = NIC layout", size=15)

d.save(OUT)
sys.path.insert(0, "/Users/akazawt/works/data-science/claudecode/.claude/skills/review-html/packs/zenn-doc/templates")
import strip_branding  # noqa: E402
strip_branding.main(OUT, OUT)
