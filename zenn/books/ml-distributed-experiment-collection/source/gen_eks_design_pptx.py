"""gen_eks_design_pptx.py: awsome-eks-gpu-cfn design figures (AMI paths, multiple node groups)."""
import os
import sys

sys.path.insert(0, "/Users/akazawt/works/data-science/tools/pptxkit")
from pptx.dml.color import RGBColor  # noqa: E402
from pptx.enum.dml import MSO_LINE_DASH_STYLE  # noqa: E402
from pptx.enum.shapes import MSO_SHAPE  # noqa: E402

from pptxkit import Deck, shapes as sh  # noqa: E402

TPL = "/Users/akazawt/works/data-science/tools/pptxkit/templates/aws-dark-2024.pptx"
OUT = sys.argv[1] if len(sys.argv) > 1 else "looped-figures.pptx"

GRN = RGBColor(36, 160, 70)
GRN_HI = RGBColor(59, 209, 111)
GRN_DK = RGBColor(18, 64, 34)
CHIP = RGBColor(22, 26, 30)
EDGE = RGBColor(70, 74, 80)
GREY = RGBColor(52, 56, 64)
LINK = RGBColor(40, 150, 240)
LINK_DK = RGBColor(16, 44, 72)
YEL = RGBColor(251, 211, 50)
YEL_DK = RGBColor(80, 66, 14)
MUTED = RGBColor(150, 150, 158)
WHITE = RGBColor(245, 245, 248)
DASH = MSO_LINE_DASH_STYLE.DASH

d = Deck(TPL, check_layout=False)
t = d.theme


def label(s, x, y, w, h, body, size=16, color=WHITE, bold=False, align="center"):
    return sh.textbox(s, t, x, y, w, h, body, size=size, color=color, bold=bold, align=align,
                      anchor="middle")


def box(s, x, y, w, h, fill, stroke=None, body="", size=15, color=WHITE, radius=0.1,
        dash=None, sw=1.25):
    sp = sh.rect(s, t, x, y, w, h, fill=fill, stroke=stroke, sw=sw, dash=dash,
                 shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=radius)
    if body:
        sh._write(sp.text_frame, t, body, size, color, True, "center", "middle")
        tf = sp.text_frame
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = sh.U(2)
    return sp


def title(s, head, sub):
    label(s, 43, 26, 1190, 46, head, size=30, bold=True, align="left")
    label(s, 43, 72, 1190, 26, sub, size=16, color=MUTED, align="left")


def arrow(s, x1, y1, x2, y2, color=WHITE, width=2, dashed=False):
    sh.arrow(s, t, x1, y1, x2, y2, color=color, width=width, dashed=dashed)




ORG = RGBColor(236, 114, 17)

# ------------------------------------------------------------------ 1. where the node AMI comes from
s = d.slide().s
title(s, "Where the GPU node AMI comes from",
      "Green: built by the AMI stack.  Blue: you bring it.  Grey: EKS picks it.  Yellow: the node group")
rows = [(130, "NodeAmiId set", "existing AMI", LINK_DK, LINK, None),
        (240, "NodeImagePackages set", "Image Builder: EKS standard AMI + packages", GRN_DK, GRN, "requires NodeImageAssertPaths"),
        (350, "NodeImageRecipeArn set", "Image Builder: your recipe", GRN_DK, GRN, "your recipe carries its tests"),
        (480, "all empty", "EKS picks the AMI from AmiType", GREY, EDGE, None)]
for y, cond, what, fill, st, note in rows:
    box(s, 43, y, 250, 80, CHIP, stroke=EDGE, body=cond, size=15)
    arrow(s, 293, y + 40, 343, y + 40)
    box(s, 343, y, 400, 80, fill, stroke=st, body=what, size=15)
    if note:
        label(s, 343, y + 82, 400, 22, note, size=12, color=MUTED)
    if y < 480:
        arrow(s, 743, y + 40, 803, 270, color=st)
box(s, 803, 200, 220, 140, CHIP, stroke=WHITE, body="launch template\nwith the AMI ID", size=15)
label(s, 803, 344, 220, 44, "AmiType CUSTOM\nfull NodeConfig in user data", size=12, color=MUTED)
arrow(s, 743, 520, 1063, 420, color=MUTED)
label(s, 803, 520, 260, 30, "no AMI ID in launch template", size=12, color=MUTED)
arrow(s, 1023, 270, 1063, 330)
box(s, 1063, 290, 174, 140, YEL_DK, stroke=YEL, body="GPU managed\nnode group", size=16)
label(s, 43, 590, 1190, 30, "Two or more set: rejected before any resource is created", size=14, color=MUTED, align="left")

# ------------------------------------------------------------------ 2. one cluster, several GPU node groups
s = d.slide().s
title(s, "One cluster, one node group stack per GPU group",
      "Blue: shared, installed by the first stack.  Green: one per stack.  Yellow: how capacity is obtained")
box(s, 43, 120, 1194, 110, CHIP, stroke=LINK, radius=0.04)
label(s, 63, 126, 1100, 28, "EKS cluster  (each GPU group: an eks-add-gpu-nodegroup.yaml stack)", size=16, bold=True, color=LINK, align="left")
box(s, 63, 162, 360, 54, LINK_DK, stroke=LINK, body="NVIDIA device plugin release", size=14)
box(s, 443, 162, 360, 54, LINK_DK, stroke=LINK, body="EFA device plugin release", size=14)
label(s, 823, 162, 400, 54, "one release per plugin;\nlater stacks reuse it if versions match", size=13, color=MUTED, align="left")
cols = [(43, "NodeGroupName = gpu", "nested stack created by root", "g7e.12xlarge", "On-Demand\n+ placement group"),
        (448, "NodeGroupName = ng-g7", "standalone stack", "g7.48xlarge", "targeted ODCR\n(no placement group)"),
        (853, "NodeGroupName = ng-p5en", "standalone stack", "p5en.48xlarge", "Capacity Block, CAPACITY_BLOCK\n(no placement group)")]
for x, name, kind, itype, cap in cols:
    arrow(s, x + 192, 230, x + 192, 270, color=LINK)
    box(s, x, 270, 384, 330, CHIP, stroke=GRN, radius=0.04)
    label(s, x, 278, 384, 26, name, size=15, bold=True, color=GRN_HI)
    label(s, x, 304, 384, 22, kind, size=12, color=MUTED)
    box(s, x + 20, 334, 344, 46, GRN_DK, stroke=GRN, body=f"managed node group ({itype})", size=13)
    box(s, x + 20, 388, 344, 46, GRN_DK, stroke=GRN, body="CodeBuild bootstrap + access entry", size=13)
    box(s, x + 20, 442, 344, 46, GRN_DK, stroke=GRN, body="verify this group: Ready + GPU/EFA", size=13)
    box(s, x + 20, 496, 344, 40, GRN_DK, stroke=GRN, body="pre-pull DaemonSet (if PrePullImage)", size=13)
    box(s, x + 20, 544, 344, 48, YEL_DK, stroke=YEL, body=cap, size=12)

d.save(OUT)
# ZN-24: テンプレートのロゴ、著作権表示、ヘッダー、フッターを必ず消してから使う
sys.path.insert(0, "/Users/akazawt/works/data-science/claudecode/.claude/skills/review-html/packs/zenn-doc/templates")
import strip_branding  # noqa: E402
strip_branding.main(OUT, OUT)
print(OUT)
