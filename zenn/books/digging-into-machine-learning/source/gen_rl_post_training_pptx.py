"""gen_rl_post_training_pptx.py: figures for the RL post-training chapters (6 slides).

Slides 1-3: rl-objectives-after-grpo (grpo-anatomy, clip-shapes, token-vs-sequence-ratio)
Slides 4-6: rl-mismatch-and-on-policy-distillation (mismatch-sources, sft-rl-opd, multi-teacher-opd)
"""
import sys

sys.path.insert(0, "/Users/akazawt/works/data-science/tools/pptxkit")
from pptx.dml.color import RGBColor  # noqa: E402
from pptx.enum.dml import MSO_LINE_DASH_STYLE  # noqa: E402
from pptx.enum.shapes import MSO_SHAPE  # noqa: E402

from pptxkit import Deck, shapes as sh  # noqa: E402

TPL = "/Users/akazawt/works/data-science/tools/pptxkit/templates/aws-dark-2024.pptx"
OUT = sys.argv[1] if len(sys.argv) > 1 else "rl-post-training-figures.pptx"

GRN = RGBColor(36, 160, 70)
GRN_HI = RGBColor(59, 209, 111)
GRN_MID = RGBColor(26, 110, 52)
GRN_DK = RGBColor(18, 64, 34)
CHIP = RGBColor(22, 26, 30)
EDGE = RGBColor(70, 74, 80)
GREY = RGBColor(52, 56, 64)
LINK = RGBColor(40, 150, 240)
LINK_DK = RGBColor(16, 44, 72)
YEL = RGBColor(251, 211, 50)
YEL_DK = RGBColor(80, 66, 14)
MUTED = RGBColor(150, 150, 158)
WHITE = RGBColor(245, 245, 248)
DASH = MSO_LINE_DASH_STYLE.DASH

d = Deck(TPL, check_layout=False)
t = d.theme


def label(s, x, y, w, h, body, size=16, color=WHITE, bold=False, align="center"):
    return sh.textbox(s, t, x, y, w, h, body, size=size, color=color, bold=bold, align=align,
                      anchor="middle")


def box(s, x, y, w, h, fill, stroke=None, body="", size=15, color=WHITE, radius=0.1,
        dash=None, sw=1.25, bold=True):
    sp = sh.rect(s, t, x, y, w, h, fill=fill, stroke=stroke, sw=sw, dash=dash,
                 shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=radius)
    if body:
        sh._write(sp.text_frame, t, body, size, color, bold, "center", "middle")
        tf = sp.text_frame
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = sh.U(2)
    return sp


def title(s, head, sub):
    label(s, 43, 26, 1190, 46, head, size=30, bold=True, align="left")
    label(s, 43, 72, 1190, 26, sub, size=16, color=MUTED, align="left")


def arrow(s, x1, y1, x2, y2, color=WHITE, width=2, dashed=False):
    sh.arrow(s, t, x1, y1, x2, y2, color=color, width=width, dashed=dashed)


# ------------------------------------------------------------------ 1. GRPO anatomy
s = d.slide().s
title(s, "Five parts of GRPO, and the work that changed each",
      "Green: a part of the GRPO loss.  Yellow: what that part decides.  Blue: methods that change it")
box(s, 340, 112, 600, 56, CHIP, stroke=EDGE, body="GRPO loss for one group of G responses", size=17)
cols = [
    ("averaging\nover tokens", "how much each\ntoken weighs",
     ["DAPO: per token\nin the group", "Dr. GRPO: no 1/|o|", "proof (binary reward):\nno length weight is\nunbiased and length-free"]),
    ("baseline and\nnormalization", "how good a response\nis, against what",
     ["Dr. GRPO: no std", "RLOO: leave-one-out", "REINFORCE++:\nbatch-wide norm", "DAPO, Magistral, Olmo 3:\nskip all-same groups"]),
    ("clip", "how much an update\nis limited",
     ["DAPO: Clip-Higher", "CISPO: clip the weight,\nkeep the gradient", "SAPO: smooth gate"]),
    ("unit of the\nprobability ratio", "token or\nwhole response",
     ["GSPO: geometric mean\nper response", "GMPO: geometric mean,\nclip per token"]),
    ("KL term", "distance to the\nreference model",
     ["removed: DAPO,\nMagistral, GLM-4.5,\nOlmo 3", "DeepSeek-V3.2:\ncorrected estimator"]),
]
cw, gap, x0 = 222, 20, 43
for i, (name, decides, methods) in enumerate(cols):
    x = x0 + i * (cw + gap)
    arrow(s, 640, 168, x + cw / 2, 206, color=MUTED, width=1.25)
    box(s, x, 206, cw, 74, GRN_DK, stroke=GRN_HI, body=name, size=17, sw=2)
    box(s, x, 292, cw, 60, YEL_DK, stroke=YEL, body=decides, size=14, bold=False)
    y = 372
    for m in methods:
        n = m.count("\n") + 1
        h = 30 + 22 * n
        box(s, x, y, cw, h, LINK_DK, stroke=LINK, body=m, size=13, bold=False)
        y += h + 10

# ------------------------------------------------------------------ 2. clip shapes
s = d.slide().s
title(s, "Where the clipped term passes gradient, advantage > 0",
      "Green: gradient flows.  Grey: no gradient from this term.  Yellow: weight held at the cap.  Bars: SAPO weight")
ax0, ax1 = 330, 1210  # ratio 0 .. 2
def rx(r):
    return ax0 + (ax1 - ax0) * r / 2.0
rows = [
    ("GRPO", "common default ε = 0.2", [(0, 1.2, "g"), (1.2, 2.0, "z")]),
    ("DAPO", "Clip-Higher, ε_high = 0.28", [(0, 1.28, "g"), (1.28, 2.0, "z")]),
    ("CISPO", "cap 1 + ε_high (illustrative)", [(0, 1.28, "g"), (1.28, 2.0, "y")]),
    ("SAPO", "τ_pos = 1.0, weight w(r)", None),
]
import math
for k, (name, sub, segs) in enumerate(rows):
    y = 140 + k * 108
    label(s, 43, y, 280, 34, name, size=20, bold=True, align="left")
    label(s, 43, y + 34, 280, 26, sub, size=14, color=MUTED, align="left")
    if segs:
        for a, b, kind in segs:
            fill, stroke, body = {
                "g": (GRN_DK, GRN_HI, "gradient flows"),
                "z": (CHIP, EDGE, "no gradient"),
                "y": (YEL_DK, YEL, "weight = cap, gradient flows"),
            }[kind]
            box(s, rx(a), y + 6, rx(b) - rx(a), 52, fill, stroke=stroke, body=body, size=15,
                color=WHITE if kind != "z" else MUTED, bold=False, radius=0.08)
    else:
        n = 20
        base = y + 74
        for j in range(n):
            rc = 2.0 * (j + 0.5) / n
            p = 1 / (1 + math.exp(-1.0 * (rc - 1.0)))
            w = 4 * p * (1 - p)
            h = 70 * w
            sh.rect(s, t, rx(2.0 * j / n) + 2, base - h, (ax1 - ax0) / n - 4, h, fill=GRN_DK, stroke=GRN_HI)
        label(s, rx(1.0) - 60, y - 22, 120, 22, "w = 1", size=13, color=WHITE)
        label(s, rx(0.0) - 10, y - 22, 160, 22, "w ≈ 0.79", size=13, color=MUTED, align="left")
        label(s, rx(2.0) - 150, y - 22, 160, 22, "w ≈ 0.79", size=13, color=MUTED, align="right")
ay = 140 + 4 * 108 + 6
sh.rect(s, t, ax0, ay, ax1 - ax0, 2, fill=WHITE)
for r in (0, 0.5, 1.0, 1.5, 2.0):
    sh.rect(s, t, rx(r) - 1, ay - 6, 2, 14, fill=WHITE)
    label(s, rx(r) - 40, ay + 12, 80, 26, f"{r:g}", size=15, color=WHITE)
label(s, ax0, ay + 40, ax1 - ax0, 26, "probability ratio  r = new / old", size=16, color=MUTED)

# ------------------------------------------------------------------ 3. token vs sequence ratio
s = d.slide().s
title(s, "A ratio per token, or one ratio per response",
      "Blue: one response's tokens.  Green: kept.  Grey: removed from the update.  Advantage > 0, example numbers")
ratios = ["1.01", "0.99", "1.02", "1.60", "1.00", "0.98"]
label(s, 43, 120, 1190, 30, "one response, six tokens, each with ratio r = new / old", size=17,
      color=MUTED, align="left")
for j, r in enumerate(ratios):
    box(s, 43 + j * 120, 160, 108, 60, LINK_DK, stroke=LINK, body=f"r = {r}", size=16)
box(s, 43, 260, 580, 330, CHIP, stroke=EDGE, radius=0.04)
label(s, 63, 272, 540, 32, "Token ratios (GRPO, DAPO)", size=19, bold=True)
for j, r in enumerate(ratios):
    out = r == "1.60"
    box(s, 73 + j * 88, 330, 78, 56, CHIP if out else GRN_DK, stroke=EDGE if out else GRN_HI,
        body=r, size=15, color=MUTED if out else WHITE)
label(s, 63, 400, 540, 30, "each token is clipped on its own", size=16, color=MUTED)
label(s, 63, 440, 540, 120, "the 1.60 token leaves the update\nthe other five stay", size=17)
box(s, 657, 260, 580, 330, CHIP, stroke=EDGE, radius=0.04)
label(s, 677, 272, 540, 32, "Sequence ratio (GSPO)", size=19, bold=True)
box(s, 717, 330, 460, 56, CHIP, stroke=EDGE, body="s = geometric mean = 1.08", size=17, color=MUTED)
for j in range(6):
    box(s, 717 + j * 78, 476, 68, 40, CHIP, stroke=EDGE, body="", radius=0.1)
label(s, 677, 396, 540, 60, "the outlier is diluted to 1.08, still above\nGSPO's cap 1 + 4e-4", size=16)
label(s, 677, 520, 540, 40, "all six tokens leave together", size=17)

# ------------------------------------------------------------------ 4. mismatch sources
s = d.slide().s
title(s, "Two engines, two probabilities for the same tokens",
      "Blue: rollout and its probability μ.  Green: training engine.  Yellow: what reduces each source")
box(s, 43, 140, 330, 110, LINK_DK, stroke=LINK, body="inference engine\nvLLM, SGLang\nsamples tokens, returns μ", size=16)
box(s, 907, 140, 330, 110, GRN_DK, stroke=GRN_HI, body="training engine\nMegatron, FSDP\nrecomputes π_old, gradient", size=16, sw=2)
arrow(s, 373, 180, 907, 180)
label(s, 380, 142, 520, 30, "tokens + μ (+ routing, sampling mask if replayed)", size=15)
arrow(s, 907, 220, 373, 220, color=MUTED, dashed=True)
label(s, 400, 226, 480, 30, "updated weights", size=15, color=MUTED)
box(s, 440, 276, 400, 60, CHIP, stroke=EDGE, body="ρ = π_old / μ ≠ 1   (same weights)", size=18)
srcs = [("summation order\nvaries with batch", "batch-invariant kernels\n(Thinking Machines,\nDeepSeek-V4)"),
        ("BF16 rounding", "FP16 for training\nand rollout"),
        ("LM head\nprecision", "FP32 LM head\n(MiniMax-M1)"),
        ("MoE picks other\nexperts", "replay rollout\nrouting (R3,\nKeep Routing)"),
        ("top-p / top-k cut\nonly in rollout", "Keep Sampling\nMask\n(DeepSeek-V3.2)")]
cw4 = 190
for i, (src, fix) in enumerate(srcs):
    x = 43 + i * (cw4 + 12)
    box(s, x, 380, cw4, 70, CHIP, stroke=EDGE, body=src, size=14, bold=False)
    box(s, x, 480, cw4, 92, YEL_DK, stroke=YEL, body=fix, size=13, bold=False)
    arrow(s, x + cw4 / 2, 450, x + cw4 / 2, 480, color=YEL, width=1.5)
    arrow(s, x + cw4 / 2, 380, 470 + i * 85, 336, color=MUTED, width=1.0)
xs = 43 + 5 * (cw4 + 12)
box(s, xs, 380, 184, 70, CHIP, stroke=MUTED, dash=DASH, body="stale policy:\nweights differ", size=14, bold=False)
box(s, xs, 480, 184, 92, YEL_DK, stroke=YEL, dash=DASH, body="π_θ / μ differs too\nAReaL, M2PO", size=13, bold=False)
arrow(s, xs + 92, 450, xs + 92, 480, color=YEL, width=1.5)
box(s, 43, 600, 1194, 50, YEL_DK, stroke=YEL,
    body="what is left: weight it  (TIS: min(ρ, C))  or drop it  (IcePop, masked IS, sequence masking)", size=16,
    bold=False)

# ------------------------------------------------------------------ 5. SFT vs RL vs OPD
s = d.slide().s
title(s, "Who writes the response, and what signal each token gets",
      "Blue: teacher.  Green: student.  Yellow: the learning signal")
cols3 = [("Off-policy distillation", "teacher writes", True, "match teacher's tokens\non teacher's prefix"),
         ("Outcome-reward RL (GRPO)", "student writes", False, "advantage from one\nreward, same for every token"),
         ("On-policy distillation", "student writes", False, "sampled-token form:\nlog π_T − log π_θ per token")]
for i, (head, who, teacher, sig) in enumerate(cols3):
    x = 43 + i * 405
    box(s, x, 120, 385, 480, CHIP, stroke=EDGE, radius=0.04)
    label(s, x + 10, 132, 365, 32, head, size=19, bold=True)
    label(s, x + 10, 170, 365, 26, who, size=15, color=MUTED)
    for j in range(6):
        box(s, x + 22 + j * 58, 214, 50, 50, LINK_DK if teacher else GRN_DK,
            stroke=LINK if teacher else GRN_HI, body="", radius=0.1)
    if i == 0:
        label(s, x + 10, 280, 365, 60, "no student rollout needed\nlearns on teacher's prefixes", size=14, color=MUTED)
    if i == 1:
        box(s, x + 22, 300, 341, 50, YEL_DK, stroke=YEL, body="one reward R → advantage", size=15, bold=False)
        for j in range(6):
            arrow(s, x + 192, 300, x + 47 + j * 58, 264, color=YEL, width=1.25)
    if i == 2:
        for j in range(6):
            box(s, x + 22 + j * 58, 300, 50, 50, YEL_DK, stroke=YEL, body=f"A{j + 1}", size=14, bold=False)
            arrow(s, x + 47 + j * 58, 300, x + 47 + j * 58, 264, color=YEL, width=1.25)
        box(s, x + 22, 372, 341, 46, LINK_DK, stroke=LINK, body="teacher scores each token", size=14, bold=False)
    box(s, x + 22, 470, 341, 90, YEL_DK, stroke=YEL, body=sig, size=15, bold=False)

# ------------------------------------------------------------------ 6. multi-teacher OPD
s = d.slide().s
title(s, "Distill RL-trained domain experts into one student",
      "Green: student, starting from the SFT model.  Blue: domain teachers.  Yellow: per-token signal in each report")
box(s, 43, 280, 170, 80, GRN_DK, stroke=GRN_HI, body="SFT model", size=17, sw=2)
teachers = ["math expert", "code expert", "agent expert", "... more domains"]
for i, name in enumerate(teachers):
    y = 130 + i * 100
    box(s, 300, y, 210, 70, LINK_DK, stroke=LINK, body=f"{name}\ndomain RL" if i < 3 else name, size=15)
    arrow(s, 213, 320, 300, y + 35, color=MUTED, width=1.25)
box(s, 640, 260, 230, 120, GRN_DK, stroke=GRN_HI, body="student\nwrites a response\nto prompt x", size=16, sw=2)
arrow(s, 213, 330, 640, 330, color=GRN_HI, width=1.5)
for i in range(4):
    arrow(s, 510, 165 + i * 100, 640, 280 + i * 20, color=LINK, width=1.0, dashed=True)
label(s, 520, 120, 150, 50, "teacher for\nx's domain", size=13, color=LINK)
label(s, 960, 96, 277, 30, "what the student learns from", size=15, color=MUTED)
box(s, 960, 130, 277, 110, YEL_DK, stroke=YEL,
    body="sampled token only\nlog π_T(y_t) − log π_θ(y_t)\nMiMo-V2-Flash, Kimi K3", size=14, bold=False)
box(s, 960, 270, 277, 110, YEL_DK, stroke=YEL,
    body="full vocabulary\nreverse KL per position\nDeepSeek-V4", size=14, bold=False)
box(s, 960, 410, 277, 110, YEL_DK, stroke=YEL,
    body="plus outcome-reward\nadvantage (GRPO)\nMiMo-V2-Flash", size=14, bold=False)
arrow(s, 960, 185, 870, 300, color=YEL, width=1.5)
arrow(s, 960, 325, 870, 320, color=YEL, width=1.5)
arrow(s, 960, 465, 870, 345, color=YEL, width=1.5)
label(s, 43, 560, 900, 30, "schematic.  GLM-5 and Nemotron-Cascade 2 take teachers from earlier training stages",
      size=14, color=MUTED, align="left")

d.save(OUT)
# ZN-24: テンプレートのロゴ、著作権表示、ヘッダー、フッターを必ず消してから使う
sys.path.insert(0, "/Users/akazawt/works/data-science/claudecode/.claude/skills/review-html/packs/zenn-doc/templates")
import strip_branding  # noqa: E402
strip_branding.main(OUT, OUT)
print(OUT)
