"""gen_rl_stats_pptx.py: figures for the mathematical-statistics chapter (4 slides).

Slide 1 stats-map, 2 objective-weights, 3 neyman-by-objective, 4 is-ess
"""
import math
import sys

sys.path.insert(0, "/Users/akazawt/works/data-science/tools/pptxkit")
from pptx.dml.color import RGBColor  # noqa: E402
from pptx.enum.dml import MSO_LINE_DASH_STYLE  # noqa: E402
from pptx.enum.shapes import MSO_SHAPE  # noqa: E402

from pptxkit import Deck, shapes as sh  # noqa: E402

TPL = "/Users/akazawt/works/data-science/tools/pptxkit/templates/aws-dark-2024.pptx"
OUT = sys.argv[1] if len(sys.argv) > 1 else "rl-stats-figures.pptx"

GRN = RGBColor(36, 160, 70)
GRN_HI = RGBColor(59, 209, 111)
GRN_MID = RGBColor(26, 110, 52)
GRN_DK = RGBColor(18, 64, 34)
CHIP = RGBColor(22, 26, 30)
EDGE = RGBColor(70, 74, 80)
GREY = RGBColor(52, 56, 64)
LINK = RGBColor(40, 150, 240)
LINK_DK = RGBColor(16, 44, 72)
YEL = RGBColor(251, 211, 50)
YEL_DK = RGBColor(80, 66, 14)
MUTED = RGBColor(150, 150, 158)
WHITE = RGBColor(245, 245, 248)
DASH = MSO_LINE_DASH_STYLE.DASH

d = Deck(TPL, check_layout=False)
t = d.theme


def label(s, x, y, w, h, body, size=16, color=WHITE, bold=False, align="center"):
    return sh.textbox(s, t, x, y, w, h, body, size=size, color=color, bold=bold, align=align,
                      anchor="middle")


def box(s, x, y, w, h, fill, stroke=None, body="", size=15, color=WHITE, radius=0.1,
        dash=None, sw=1.25, bold=True):
    sp = sh.rect(s, t, x, y, w, h, fill=fill, stroke=stroke, sw=sw, dash=dash,
                 shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=radius)
    if body:
        sh._write(sp.text_frame, t, body, size, color, bold, "center", "middle")
        tf = sp.text_frame
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = sh.U(2)
    return sp


def title(s, head, sub):
    label(s, 43, 26, 1190, 46, head, size=30, bold=True, align="left")
    label(s, 43, 72, 1190, 26, sub, size=16, color=MUTED, align="left")


def arrow(s, x1, y1, x2, y2, color=WHITE, width=2, dashed=False):
    sh.arrow(s, t, x1, y1, x2, y2, color=color, width=width, dashed=dashed)



# ------------------------------------------------------------------ 1. stats map
s = d.slide().s
title(s, "Statistics concepts behind RL post-training choices",
      "Green: the statistical idea.  Blue: where the RL methods use it.  Schematic")
rows = [("control variate", "group mean as baseline; RLOO = mean-centred × G/(G−1)"),
        ("variance-stabilizing transform", "large G: dividing by group std ascends 2·arcsin√p"),
        ("ratio estimator vs mean of ratios", "token-level mean (DAPO) vs per-answer mean (GRPO)"),
        ("χ² divergence and effective sample size", "sequence weights explode; masks threshold a KL estimate"),
        ("winsorizing vs trimming", "TIS, CISPO cap the weight; IcePop, masking drop the token"),
        ("estimators of KL", "k3 value is unbiased; its loss gradient is ∇KL(π_ref‖π)"),
        ("Rao–Blackwell", "full-vocabulary OPD = per-position conditional mean"),
        ("Neyman allocation", "rollouts per prompt ≈ ∝ |f′(p)|·σ under the objective f"),
        ("UMVUE, paired tests, FDR", "pass@k estimator, simultaneous bands, per-problem tests"),
        ("Condorcet jury theorem", "plurality labels converge to a unique mode, right or wrong")]
y0, hh, gap = 118, 50, 9
for i, (a, b) in enumerate(rows):
    y = y0 + i * (hh + gap)
    box(s, 43, y, 470, hh, GRN_DK, stroke=GRN, body=a, size=17)
    arrow(s, 518, y + hh / 2, 568, y + hh / 2, color=MUTED)
    box(s, 573, y, 664, hh, LINK_DK, stroke=LINK, body=b, size=16, bold=False)


# ------------------------------------------------------------------ 2. objective weights
s = d.slide().s
title(s, "Weight on ∇p for one prompt, by pass rate p",
      "Bars: d f/dp, computed for binary reward in the large-group limit, scaled so that p = 0.5 gives 1")
ps = [0.05 + 0.1 * i for i in range(10)]
rows = [("Dr. GRPO, f = p", [1.0] * 10, LINK_DK, LINK),
        ("GRPO, f = 2·arcsin√p", [0.5 / math.sqrt(p * (1 - p)) for p in ps], GRN_DK, GRN_HI),
        ("log objective, f = log p", [0.5 / p for p in ps], YEL_DK, YEL)]
bx0, bw, gapb = 330, 78, 14
hmax = 4.0
for r, (name, ws, fill, stroke) in enumerate(rows):
    base = 250 + r * 150
    label(s, 43, base - 80, 280, 34, name, size=18, bold=True, align="left")
    for i, w in enumerate(ws):
        h = 110 * min(w, hmax) / hmax
        sh.rect(s, t, bx0 + i * (bw + gapb), base - h, bw, h, fill=fill, stroke=stroke)
        txt = f"{w:.2f}" if w < hmax else f"{w:.1f}▲"
        label(s, bx0 + i * (bw + gapb), base - h - 24, bw, 22, txt, size=14, color=WHITE)
for i, p in enumerate(ps):
    label(s, bx0 + i * (bw + gapb), 560, bw, 24, f"p={p:.2f}", size=14, color=MUTED)
label(s, 330, 600, 900, 50, "▲ bar cut at 4.  GRPO weights the extremes; the log objective weights hard prompts.\nThe per-prompt optimum is the same for all three; the mix across prompts differs.",
      size=14, color=MUTED, align="left")


# ------------------------------------------------------------------ 3. neyman by objective
s = d.slide().s
title(s, "Neyman allocation of rollouts depends on the objective",
      "Bars: ∝ |f′(p)|·√(p(1−p)), mean 16 over these 10 p.  Assumes equal score norms, drops the ‖∇p‖² term")
def scale(ws):
    m = sum(ws) / len(ws)
    return [16 * w / m for w in ws]
rows = [("Dr. GRPO, f = p", scale([math.sqrt(p * (1 - p)) for p in ps]), LINK_DK, LINK),
        ("GRPO, f = 2·arcsin√p", scale([1.0] * 10), GRN_DK, GRN_HI),
        ("log objective, f = log p", scale([math.sqrt((1 - p) / p) for p in ps]), YEL_DK, YEL)]
hmax = 48.0
for r, (name, gs, fill, stroke) in enumerate(rows):
    base = 250 + r * 150
    label(s, 43, base - 80, 280, 34, name, size=18, bold=True, align="left")
    for i, g in enumerate(gs):
        h = 110 * min(g, hmax) / hmax
        sh.rect(s, t, bx0 + i * (bw + gapb), base - h, bw, h, fill=fill, stroke=stroke)
        label(s, bx0 + i * (bw + gapb), base - h - 24, bw, 22, f"{g:.0f}", size=14, color=WHITE)
for i, p in enumerate(ps):
    label(s, bx0 + i * (bw + gapb), 560, bw, 24, f"p={p:.2f}", size=14, color=MUTED)
label(s, 330, 600, 900, 50, "Under the arcsine objective (large G) the guide is a fixed count per prompt.\nThe log row is close to Reinforce-Ada's 1/√p for small p.",
      size=14, color=MUTED, align="left")


# ------------------------------------------------------------------ 4. IS effective sample size
s = d.slide().s
title(s, "Sequence-level weights: ESS guide for 16 answers",
      "Cells: 16 / (1 + κ)^L, upper bound on the population guide 16 / E[w²]; fixed L, χ² ≥ κ at every prefix")
cs = [1e-4, 1e-3, 1e-2]
Ts = [100, 1000, 10000]
cx0, cy0, cw, ch = 380, 190, 270, 120
for j, T in enumerate(Ts):
    label(s, cx0 + j * (cw + 12), cy0 - 44, cw, 32, f"L = {T:,} tokens", size=18, color=MUTED, bold=True)
for i, c in enumerate(cs):
    label(s, 43, cy0 + i * (ch + 12), 320, ch, f"per-token χ² ≥ {c:g}", size=19, bold=True, align="left")
    for j, T in enumerate(Ts):
        m = (1 + c) ** T
        ess = 16 / m
        if ess >= 8:
            fill, stroke = GRN_DK, GRN_HI
        elif ess >= 1:
            fill, stroke = YEL_DK, YEL
        else:
            fill, stroke = CHIP, MUTED
        txt = f"{ess:.1f}" if ess >= 1 else ("< 0.001" if ess < 1e-3 else f"{ess:.3f}")
        box(s, cx0 + j * (cw + 12), cy0 + i * (ch + 12), cw, ch, fill, stroke=stroke, body=txt, size=26)
label(s, 380, 600, 860, 50, "Green: guide ≥ 8.  Yellow: 1 to 8.\nGrey: below 1, the estimate hinges on rare answers with huge weights.",
      size=14, color=MUTED, align="left")

d.save(OUT)
# ZN-24: テンプレートのロゴ、著作権表示、ヘッダー、フッターを必ず消してから使う
sys.path.insert(0, "/Users/akazawt/works/data-science/claudecode/.claude/skills/review-html/packs/zenn-doc/templates")
import strip_branding  # noqa: E402
strip_branding.main(OUT, OUT)
print(OUT)
