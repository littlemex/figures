"""gen_moe_routing_pptx.py: MoE routing chapter figures (4 slides)."""
import os
import sys

sys.path.insert(0, "/Users/akazawt/works/data-science/tools/pptxkit")
from pptx.dml.color import RGBColor  # noqa: E402
from pptx.enum.dml import MSO_LINE_DASH_STYLE  # noqa: E402
from pptx.enum.shapes import MSO_SHAPE  # noqa: E402

from pptxkit import Deck, shapes as sh  # noqa: E402

TPL = "/Users/akazawt/works/data-science/tools/pptxkit/templates/aws-dark-2024.pptx"
OUT = sys.argv[1] if len(sys.argv) > 1 else "looped-figures.pptx"

GRN = RGBColor(36, 160, 70)
GRN_HI = RGBColor(59, 209, 111)
GRN_DK = RGBColor(18, 64, 34)
CHIP = RGBColor(22, 26, 30)
EDGE = RGBColor(70, 74, 80)
GREY = RGBColor(52, 56, 64)
LINK = RGBColor(40, 150, 240)
LINK_DK = RGBColor(16, 44, 72)
YEL = RGBColor(251, 211, 50)
YEL_DK = RGBColor(80, 66, 14)
MUTED = RGBColor(150, 150, 158)
WHITE = RGBColor(245, 245, 248)
DASH = MSO_LINE_DASH_STYLE.DASH

d = Deck(TPL, check_layout=False)
t = d.theme


def label(s, x, y, w, h, body, size=16, color=WHITE, bold=False, align="center"):
    return sh.textbox(s, t, x, y, w, h, body, size=size, color=color, bold=bold, align=align,
                      anchor="middle")


def box(s, x, y, w, h, fill, stroke=None, body="", size=15, color=WHITE, radius=0.1,
        dash=None, sw=1.25):
    sp = sh.rect(s, t, x, y, w, h, fill=fill, stroke=stroke, sw=sw, dash=dash,
                 shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=radius)
    if body:
        sh._write(sp.text_frame, t, body, size, color, True, "center", "middle")
        tf = sp.text_frame
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = sh.U(2)
    return sp


def title(s, head, sub):
    label(s, 43, 26, 1190, 46, head, size=30, bold=True, align="left")
    label(s, 43, 72, 1190, 26, sub, size=16, color=MUTED, align="left")


def arrow(s, x1, y1, x2, y2, color=WHITE, width=2, dashed=False):
    sh.arrow(s, t, x1, y1, x2, y2, color=color, width=width, dashed=dashed)



L1 = "Green: router, selected experts.  Grey: idle experts, top-2 step.  Blue: token, output.  Yellow: backward"
L2 = "Green: the loop between router and expert.  Yellow: where each balancing method enters"
L3 = "Blue: a pattern the paper reports.  Grey: the paper reports little pattern.  Dashed: not reported"
L4 = "Green: router and experts.  Blue: data points and active units"

# ------------------------------------------------------------------ 1. the routing layer
s = d.slide().s
title(s, "One token through a top-2 MoE layer", L1)
box(s, 43, 300, 130, 70, LINK_DK, stroke=LINK, body="token x", size=17)
box(s, 223, 290, 170, 90, GRN_DK, stroke=GRN_HI, body="router\nh = W_r · x", size=17, sw=2)
arrow(s, 173, 335, 223, 335)
box(s, 443, 290, 170, 90, CHIP, stroke=EDGE, body="softmax\nthen top-2", size=17)
arrow(s, 393, 335, 443, 335)
ys = [140, 196, 252, 308, 364, 420, 476, 532]
for i, y in enumerate(ys):
    on = i in (2, 4)
    box(s, 693, y, 170, 50, GRN_DK if on else CHIP, stroke=GRN_HI if on else EDGE,
        body=f"expert {i + 1}", size=15, color=WHITE if on else MUTED, sw=2 if on else 1.25)
arrow(s, 613, 320, 693, 277, color=WHITE)
arrow(s, 613, 345, 693, 389, color=WHITE)
label(s, 615, 252, 80, 26, "p3", size=15, color=GRN_HI)
label(s, 615, 380, 80, 26, "p5", size=15, color=GRN_HI)
box(s, 943, 300, 290, 70, LINK_DK, stroke=LINK, body="y = p3 · E3(x) + p5 · E5(x)", size=17)
arrow(s, 863, 277, 943, 320)
arrow(s, 863, 389, 943, 350)
# backward
arrow(s, 1088, 370, 1088, 660, color=YEL, dashed=True)
arrow(s, 1088, 660, 308, 660, color=YEL, dashed=True)
arrow(s, 308, 660, 308, 380, color=YEL, dashed=True)
label(s, 360, 618, 700, 30, "backward pass: via p3 and p5 only", size=15, color=YEL)
label(s, 648, 108, 260, 30, "8 experts, 2 selected", size=14, color=MUTED)

# ------------------------------------------------------------------ 2. the imbalance loop
s = d.slide().s
title(s, "Imbalance reinforces itself, and where each fix enters", L2)
cyc = [(470, 140, "router picks\nexpert A more"), (820, 300, "A receives\nmore gradient"),
       (470, 460, "A gets better\nat its tokens"), (120, 300, "router scores\nA even higher")]
W, H = 300, 90
for x, y, txt in cyc:
    box(s, x, y, W, H, GRN_DK, stroke=GRN, body=txt, size=17)
arrow(s, 770, 185, 900, 300, color=WHITE, width=2)
arrow(s, 900, 390, 770, 505, color=WHITE, width=2)
arrow(s, 470, 505, 340, 390, color=WHITE, width=2)
arrow(s, 230, 300, 470, 190, color=WHITE, width=2)
box(s, 43, 128, 260, 64, YEL_DK, stroke=YEL, body="noise on logits", size=15)
box(s, 43, 204, 260, 64, YEL_DK, stroke=YEL, body="loss-free bias", size=15)
arrow(s, 303, 160, 352, 240, color=YEL)
arrow(s, 303, 236, 352, 248, color=YEL)
label(s, 360, 252, 160, 26, "before top-k", size=13, color=YEL, align="left")
box(s, 920, 128, 317, 76, YEL_DK, stroke=YEL, body="balanced assignment\n(BASE, Expert Choice)", size=15)
arrow(s, 920, 166, 770, 175, color=YEL)
label(s, 790, 128, 130, 26, "replaces top-k", size=13, color=YEL, align="left")
box(s, 43, 500, 260, 64, YEL_DK, stroke=YEL, body="aux balance loss", size=15)
arrow(s, 200, 500, 230, 390, color=YEL)
label(s, 215, 440, 200, 26, "extra gradient", size=13, color=YEL, align="left")
label(s, 380, 320, 520, 50, "never picked: no gradient", size=16, color=MUTED)

# ------------------------------------------------------------------ 3. what routers follow
s = d.slide().s
title(s, "What routing follows in released models", L3)
colx = [(323, "Mixtral 8x7B", "8 experts, top-2"), (623, "OpenMoE", "650M-34B models"),
        (923, "OLMoE-1B-7B", "64 experts, top-8")]
for x, head, sub in colx:
    label(s, x, 112, 290, 30, head, size=19, bold=True)
    label(s, x, 142, 290, 24, sub, size=14, color=MUTED)
rows = [(180, "domain or topic"), (290, "token and syntax"), (400, "when it settles")]
cells = {
    (0, 0): ("no obvious topic pattern\n(DM Math differs slightly)", False), (1, 0): ("close to uniform", False), (2, 0): ("clear experts\nfor arXiv, GitHub", True),
    (0, 1): ("same expert for\n'self', indentation", True), (1, 1): ("mostly by token ID", True),
    (2, 1): ("later layers follow\nthe predicted token", True),
    (0, 2): ("not reported", False), (1, 2): ("early in pretraining", True), (2, 2): ("up to ~60% overlap\nat 1% of training", True),
}
for y, txt in rows:
    label(s, 43, y, 260, 90, txt, size=17, bold=True, align="left")
    for c, (x, _, _) in enumerate(colx):
        body, strong = cells[(c, rows.index((y, txt)))]
        box(s, x, y, 290, 90, LINK_DK if strong else CHIP, stroke=LINK if strong else EDGE,
            body=body, size=15, color=WHITE if strong else MUTED,
            dash=DASH if body == "not reported" else None)

# ------------------------------------------------------------------ 4. why splitting helps
s = d.slide().s
title(s, "Two explanations for why splitting pays", L4)
box(s, 43, 120, 580, 470, CHIP, stroke=EDGE, radius=0.04)
box(s, 657, 120, 580, 470, CHIP, stroke=EDGE, radius=0.04)
label(s, 63, 132, 540, 32, "Data side: clusters (schematic)", size=19, bold=True)
label(s, 677, 132, 540, 32, "FFN side: activations are already sparse", size=19, bold=True)
for i, (cx, cy) in enumerate([(130, 230), (230, 300), (130, 370)]):
    for dx, dy in [(0, 0), (22, 10), (-14, 18), (10, -16)]:
        box(s, cx + dx, cy + dy, 14, 14, LINK_DK, stroke=LINK, radius=0.5)
box(s, 300, 255, 150, 90, GRN_DK, stroke=GRN_HI, body="router learns\ncluster centers", size=14, sw=2)
arrow(s, 260, 300, 300, 300)
for j, y in enumerate([200, 290, 380]):
    box(s, 483, y, 120, 60, GRN_DK, stroke=GRN, body=f"expert {j + 1}", size=14)
    arrow(s, 450, 300, 483, y + 30)
label(s, 63, 470, 540, 90, "simpler sub-problem per expert\nnonlinear experts required", size=15, color=MUTED)
# FFN side
label(s, 677, 180, 540, 26, "dense FFN (ReLU) hidden units for one token", size=14, color=MUTED)
for k in range(24):
    on = k in (13, 16)
    box(s, 697 + k * 22, 214, 18, 40, LINK if on else GREY, stroke=LINK if on else EDGE, radius=0.1)
label(s, 677, 262, 540, 26, "3.0% non-zero in T5-Base (Lazy Neuron)", size=14, color=LINK)
arrow(s, 947, 300, 947, 350)
for j in range(4):
    on = j == 2
    box(s, 697 + j * 130, 360, 120, 60, GRN_DK if on else CHIP, stroke=GRN_HI if on else EDGE,
        body=f"block {j + 1}", size=14, color=WHITE if on else MUTED)
label(s, 677, 470, 540, 90, "MoEfication: 10-30% of\nFFN params per token", size=15, color=MUTED)

d.save(OUT)
# ZN-24: テンプレートのロゴ、著作権表示、ヘッダー、フッターを必ず消してから使う
sys.path.insert(0, "/Users/akazawt/works/data-science/claudecode/.claude/skills/review-html/packs/zenn-doc/templates")
import strip_branding  # noqa: E402
strip_branding.main(OUT, OUT)
print(OUT)
