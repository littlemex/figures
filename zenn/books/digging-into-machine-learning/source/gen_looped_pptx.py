"""gen_looped_pptx.py — 「思考の状態はどこに置かれるか」の章の図を、ネイティブ shape の pptx 4 枚で組む。

    /usr/bin/python3 gen_looped_pptx.py <out.pptx>

1. 途中の状態の置き場所: thinking は token、loop は hidden state、persistent memory は request をまたぐ状態
2. Looped Transformer の 1 token: prelude / 共有ブロックの反復 / coda と、反復ごとに増える KV と共有 KV
3. 記憶の地図: いつ書くか (事前学習 / request の最中 / request の間) と、何を持つか
4. 考えた結果を次の request へ渡す循環と、既存研究が受け持つ段
図の文字は英語。根拠は章本文と各論文の abstract (arXiv 2502.05171, 2510.25741, 2605.07721, 2610.02383,
2601.07372, 2501.00663, 2601.05505, 2506.06266, 2504.13171, 2605.30757)。
色: 緑 = 重みを使う計算、青 = モデルが運ぶ状態、黄 = request をまたいで残るもの、灰 = 外の入出力。
"""
import os
import sys

sys.path.insert(0, "/Users/akazawt/works/data-science/tools/pptxkit")
from pptx.dml.color import RGBColor  # noqa: E402
from pptx.enum.dml import MSO_LINE_DASH_STYLE  # noqa: E402
from pptx.enum.shapes import MSO_SHAPE  # noqa: E402

from pptxkit import Deck, shapes as sh  # noqa: E402

TPL = "/Users/akazawt/works/data-science/tools/pptxkit/templates/aws-dark-2024.pptx"
OUT = sys.argv[1] if len(sys.argv) > 1 else "looped-figures.pptx"

GRN = RGBColor(36, 160, 70)
GRN_HI = RGBColor(59, 209, 111)
GRN_DK = RGBColor(18, 64, 34)
CHIP = RGBColor(22, 26, 30)
EDGE = RGBColor(70, 74, 80)
GREY = RGBColor(52, 56, 64)
LINK = RGBColor(40, 150, 240)
LINK_DK = RGBColor(16, 44, 72)
YEL = RGBColor(251, 211, 50)
YEL_DK = RGBColor(80, 66, 14)
MUTED = RGBColor(150, 150, 158)
WHITE = RGBColor(245, 245, 248)
DASH = MSO_LINE_DASH_STYLE.DASH

d = Deck(TPL, check_layout=False)
t = d.theme


def label(s, x, y, w, h, body, size=16, color=WHITE, bold=False, align="center"):
    return sh.textbox(s, t, x, y, w, h, body, size=size, color=color, bold=bold, align=align,
                      anchor="middle")


def box(s, x, y, w, h, fill, stroke=None, body="", size=15, color=WHITE, radius=0.1,
        dash=None, sw=1.25):
    sp = sh.rect(s, t, x, y, w, h, fill=fill, stroke=stroke, sw=sw, dash=dash,
                 shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=radius)
    if body:
        sh._write(sp.text_frame, t, body, size, color, True, "center", "middle")
        tf = sp.text_frame
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = sh.U(2)
    return sp


def title(s, head, sub):
    label(s, 43, 26, 1190, 46, head, size=30, bold=True, align="left")
    label(s, 43, 72, 1190, 26, sub, size=16, color=MUTED, align="left")


def arrow(s, x1, y1, x2, y2, color=WHITE, width=2, dashed=False):
    sh.arrow(s, t, x1, y1, x2, y2, color=color, width=width, dashed=dashed)


LEGEND = "Green: weights.  Blue: state the model carries.  Yellow: kept across requests"

# ------------------------------------------------------------------ 1. where the state lives
s = d.slide().s
title(s, "Where the intermediate state lives", LEGEND)
cols = [(43, "Thinking (CoT)", "token axis"),
        (448, "Looped Transformer", "depth axis"),
        (853, "Carried latent state", "request axis (open)")]
for x, head, sub in cols:
    label(s, x, 112, 384, 34, head, size=22, bold=True)
    label(s, x, 144, 384, 24, sub, size=14, color=MUTED)
    box(s, x, 176, 384, 300, CHIP, stroke=EDGE, radius=0.04)
# thinking: tokens grow in the context
box(s, 63, 200, 344, 50, GRN_DK, stroke=GRN, body="model (L layers)", size=15)
toks = ["Q", "t1", "t2", "t3", "...", "A"]
for i, tk in enumerate(toks):
    fill, st = (GREY, EDGE) if tk in ("Q", "A") else (LINK_DK, LINK)
    box(s, 63 + i * 58, 290, 52, 50, fill, stroke=st, body=tk, size=15)
arrow(s, 235, 250, 235, 290)
sh.bracket(s, t, 121, 352, 228, "grows", size=14, down=True, color=LINK)
label(s, 63, 404, 344, 56, "state = tokens in context\nreadable, one token per step", size=14, color=MUTED)
# loop: one shared block, h0 -> h3
box(s, 468, 200, 344, 50, GRN_DK, stroke=GRN, body="shared block  x r", size=15)
for i in range(4):
    box(s, 478 + i * 84, 290, 70, 50, LINK_DK, stroke=LINK, body=f"h{i}", size=15)
    if i < 3:
        arrow(s, 548 + i * 84, 315, 562 + i * 84, 315, color=LINK)
arrow(s, 640, 250, 640, 290)
label(s, 468, 352, 344, 40, "same size every pass", size=14, color=LINK)
label(s, 468, 404, 344, 56, "state = hidden vectors\nnot read out as words", size=14, color=MUTED)
# persistent: request -> state -> request
box(s, 873, 200, 160, 50, GREY, stroke=EDGE, body="request 1", size=14)
box(s, 1053, 200, 164, 50, GREY, stroke=EDGE, body="request 2", size=14)
box(s, 873, 290, 160, 50, YEL_DK, stroke=YEL, body="state A", size=15)
box(s, 1053, 290, 164, 50, YEL_DK, stroke=YEL, body="state B", size=15)
arrow(s, 953, 250, 953, 290)
arrow(s, 1135, 250, 1135, 290)
arrow(s, 1033, 315, 1053, 315, color=YEL)
label(s, 873, 352, 344, 40, "saved, then loaded", size=14, color=YEL)
label(s, 873, 404, 344, 56, "state = what thinking left\nlives beyond one request", size=14, color=MUTED)
# lifetime axis
for x, txt, col in ((43, "lives for one request", LINK), (448, "lives for one token's pass", LINK),
                    (853, "lives across requests", YEL)):
    box(s, x, 496, 384, 50, CHIP, stroke=col, body=txt, size=16, color=col, radius=0.3)

# ------------------------------------------------------------------ 2. looped block and its KV
s = d.slide().s
title(s, "One token through a looped model", LEGEND)
box(s, 43, 150, 150, 70, GREY, stroke=EDGE, body="input e", size=16)
box(s, 233, 150, 170, 70, GRN_DK, stroke=GRN, body="prelude", size=17)
box(s, 463, 135, 330, 100, GRN_DK, stroke=GRN_HI, body="core block R\ns_i = R(e, s_(i-1))", size=17, sw=2)
box(s, 853, 150, 170, 70, GRN_DK, stroke=GRN, body="coda", size=17)
box(s, 1063, 150, 174, 70, GREY, stroke=EDGE, body="next token", size=16)
arrow(s, 193, 185, 233, 185)
arrow(s, 403, 185, 463, 185)
arrow(s, 793, 185, 853, 185)
arrow(s, 1023, 185, 1063, 185)
# loop back
arrow(s, 740, 235, 740, 262, color=LINK)
arrow(s, 740, 262, 516, 262, color=LINK)
arrow(s, 516, 262, 516, 235, color=LINK)
label(s, 463, 266, 330, 30, "repeat r times, same weights", size=14, color=LINK)
# input injection
arrow(s, 118, 220, 118, 300, color=MUTED, dashed=True)
arrow(s, 118, 300, 480, 300, color=MUTED, dashed=True)
arrow(s, 480, 300, 480, 235, color=MUTED, dashed=True)
label(s, 140, 304, 320, 28, "e is fed again every pass", size=13, color=MUTED, align="left")
# exit gate
box(s, 853, 250, 384, 56, CHIP, stroke=MUTED, body="exit gate: learned depth per token", size=15)
# KV
label(s, 43, 360, 1194, 30, "KV cache for that token", size=17, bold=True, align="left")
box(s, 43, 400, 572, 210, CHIP, stroke=EDGE, radius=0.04)
box(s, 665, 400, 572, 210, CHIP, stroke=EDGE, radius=0.04)
label(s, 63, 410, 532, 30, "one cache per pass", size=16, bold=True)
for i in range(4):
    box(s, 83 + i * 126, 456, 110, 50, LINK_DK, stroke=LINK, body=f"KV pass {i + 1}", size=13)
label(s, 63, 530, 532, 56, "memory grows with r", size=15, color=WHITE)
label(s, 685, 410, 532, 30, "shared cache (MELT; LPT adds a window)", size=16, bold=True)
box(s, 705, 456, 236, 50, LINK_DK, stroke=LINK, body="shared KV", size=14)
box(s, 961, 456, 236, 50, CHIP, stroke=EDGE, body="local window (LPT)", size=13, color=MUTED)
label(s, 685, 530, 532, 56, "MELT: flat as r grows.  LPT: + short windows", size=15, color=GRN_HI)

# ------------------------------------------------------------------ 3. memory map
s = d.slide().s
title(s, "Who writes the memory, and when", LEGEND)
xs = [(263, "pretraining"), (583, "while answering"), (903, "between requests")]
ys = [(150, "facts and patterns"), (300, "the context"), (450, "results of thinking")]
for x, txt in xs:
    label(s, x, 112, 300, 30, txt, size=17, bold=True)
for y, txt in ys:
    label(s, 43, y, 200, 120, txt, size=16, bold=True, align="left")
    for x, _ in xs:
        box(s, x, y, 300, 130, CHIP, stroke=EDGE, radius=0.04)
box(s, 283, 175, 260, 80, GRN_DK, stroke=GRN, body="Engram\nn-gram lookup, read-only", size=14)
box(s, 603, 315, 260, 46, LINK_DK, stroke=LINK, body="KV cache", size=14)
box(s, 603, 371, 260, 46, LINK_DK, stroke=LINK, body="Titans / ATLAS", size=14)
box(s, 923, 325, 260, 80, YEL_DK, stroke=YEL, body="Cartridges\nKV trained per corpus", size=14)
box(s, 603, 475, 260, 80, LINK_DK, stroke=LINK, body="FlashMem\nfrom last hidden state", size=14)
box(s, 923, 465, 260, 46, YEL_DK, stroke=YEL, body="sleep-time compute", size=13)
box(s, 923, 521, 260, 46, None, stroke=YEL, body="saved latent thought: open", size=13, color=YEL, dash=DASH)
label(s, 263, 600, 940, 30, "Engram stores what a token sequence means; the open cell stores what the model concluded",
      size=13, color=MUTED, align="left")

# ------------------------------------------------------------------ 4. persistent thought loop
s = d.slide().s
title(s, "Carrying a thought state to the next request", LEGEND)
steps = [("request n", GREY, EDGE), ("load state", YEL_DK, YEL), ("think in loops", GRN_DK, GRN),
         ("consolidate", LINK_DK, LINK), ("save state", YEL_DK, YEL)]
x = 43
W = 210
for i, (name, fill, st) in enumerate(steps):
    box(s, x, 160, W, 70, fill, stroke=st, body=name, size=17)
    if i < len(steps) - 1:
        arrow(s, x + W, 195, x + W + 36, 195)
    x += W + 36
# back to the next request
arrow(s, 1147, 230, 1147, 300, color=YEL)
arrow(s, 1147, 300, 148, 300, color=YEL)
arrow(s, 148, 300, 148, 230, color=YEL)
label(s, 400, 304, 480, 30, "request n + 1 starts from the saved state", size=14, color=YEL)
# existing work per step
label(s, 43, 370, 1194, 30, "Closest existing work", size=17, bold=True, align="left")
rows = [(289, "Cartridges\nload a trained KV"), (535, "MELT (from Ouro)\nloop at fixed memory"),
        (781, "FlashMem\nmemory from last h"), (1027, "open: re-loop a saved\nlatent state")]
for xx, txt in rows:
    open_cell = txt.startswith("open")
    box(s, xx, 410, 210, 90, None if open_cell else CHIP, stroke=YEL if open_cell else EDGE,
        body=txt, size=13, color=YEL if open_cell else WHITE, dash=DASH if open_cell else None)
label(s, 43, 530, 1194, 60, "Each step has a nearby method; re-looping a saved latent state across requests is open",
      size=15, color=MUTED, align="left")

d.save(OUT)
# ZN-24: テンプレートのロゴ、著作権表示、ヘッダー、フッターを必ず消してから使う
sys.path.insert(0, "/Users/akazawt/works/data-science/claudecode/.claude/skills/review-html/packs/zenn-doc/templates")
import strip_branding  # noqa: E402
strip_branding.main(OUT, OUT)
print(OUT)
