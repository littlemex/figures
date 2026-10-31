"""gen_rl_async_pptx.py: before/after figure for synchronous vs asynchronous RL.

Slide 1 sync-vs-async
"""
import math
import sys

sys.path.insert(0, "/Users/akazawt/works/data-science/tools/pptxkit")
from pptx.dml.color import RGBColor  # noqa: E402
from pptx.enum.dml import MSO_LINE_DASH_STYLE  # noqa: E402
from pptx.enum.shapes import MSO_SHAPE  # noqa: E402

from pptxkit import Deck, shapes as sh  # noqa: E402

TPL = "/Users/akazawt/works/data-science/tools/pptxkit/templates/aws-dark-2024.pptx"
OUT = sys.argv[1] if len(sys.argv) > 1 else "rl-async-figures.pptx"

GRN = RGBColor(36, 160, 70)
GRN_HI = RGBColor(59, 209, 111)
GRN_MID = RGBColor(26, 110, 52)
GRN_DK = RGBColor(18, 64, 34)
CHIP = RGBColor(22, 26, 30)
EDGE = RGBColor(70, 74, 80)
GREY = RGBColor(52, 56, 64)
LINK = RGBColor(40, 150, 240)
LINK_DK = RGBColor(16, 44, 72)
YEL = RGBColor(251, 211, 50)
YEL_DK = RGBColor(80, 66, 14)
MUTED = RGBColor(150, 150, 158)
WHITE = RGBColor(245, 245, 248)
DASH = MSO_LINE_DASH_STYLE.DASH

d = Deck(TPL, check_layout=False)
t = d.theme


def label(s, x, y, w, h, body, size=16, color=WHITE, bold=False, align="center"):
    return sh.textbox(s, t, x, y, w, h, body, size=size, color=color, bold=bold, align=align,
                      anchor="middle")


def box(s, x, y, w, h, fill, stroke=None, body="", size=15, color=WHITE, radius=0.1,
        dash=None, sw=1.25, bold=True):
    sp = sh.rect(s, t, x, y, w, h, fill=fill, stroke=stroke, sw=sw, dash=dash,
                 shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=radius)
    if body:
        sh._write(sp.text_frame, t, body, size, color, bold, "center", "middle")
        tf = sp.text_frame
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = sh.U(2)
    return sp


def title(s, head, sub):
    label(s, 43, 26, 1190, 46, head, size=30, bold=True, align="left")
    label(s, 43, 72, 1190, 26, sub, size=16, color=MUTED, align="left")


def arrow(s, x1, y1, x2, y2, color=WHITE, width=2, dashed=False):
    sh.arrow(s, t, x1, y1, x2, y2, color=color, width=width, dashed=dashed)



s = d.slide().s
title(s, "Synchronous vs asynchronous RL",
      "Left: before (synchronous).  Right: after (asynchronous).  Schematic, time runs left to right")
# panel headers
label(s, 43, 112, 580, 30, "before: synchronous", size=20, bold=True, align="left")
label(s, 663, 112, 580, 30, "after: asynchronous", size=20, bold=True, align="left")
# --- left panel: sequential phases with idle GPUs
lx, ly = 43, 160
rows = ["GPU 1", "GPU 2", "GPU 3", "GPU 4"]
lens = [380, 170, 240, 120]
for i, (r, L) in enumerate(zip(rows, lens)):
    y = ly + i * 58
    label(s, lx, y, 80, 44, r, size=14, color=MUTED, align="left")
    box(s, lx + 85, y, L, 44, LINK_DK, stroke=LINK, body="generate", size=14, bold=False)
    if L < 380:
        box(s, lx + 85 + L, y, 380 - L, 44, CHIP, stroke=MUTED, dash=DASH, body="idle", size=13, color=MUTED, bold=False)
box(s, lx + 85, ly + 4 * 58 + 6, 380, 44, GRN_DK, stroke=GRN, body="score + train", size=15)
box(s, lx + 85, ly + 5 * 58 + 6, 380, 44, YEL_DK, stroke=YEL, body="sync weights", size=15)
label(s, lx, ly + 6 * 58 + 18, 560, 60, "every step waits for the longest answer;\nall answers come from the current weights",
      size=15, color=MUTED, align="left")
# --- right panel: three concurrent lanes
rx, ry = 663, 160
box(s, rx, ry, 580, 70, LINK_DK, stroke=LINK, body="1  inference engine keeps generating with the weights it has", size=15, bold=False)
box(s, rx, ry + 110, 580, 70, GRN_DK, stroke=GRN, body="2  trainer learns from finished answers\n(some from older weights, many short ones)", size=15, bold=False)
box(s, rx, ry + 220, 580, 70, YEL_DK, stroke=YEL, body="3  updated weights are sent to the inference engine", size=15, bold=False)
arrow(s, rx + 290, ry + 72, rx + 290, ry + 108, color=MUTED)
arrow(s, rx + 290, ry + 182, rx + 290, ry + 218, color=MUTED)
arrow(s, rx + 590, ry + 255, rx + 590, ry + 35, color=MUTED, dashed=True)
label(s, rx, ry + 320, 580, 60, "less waiting; staleness and short-answer bias\nmust be handled (AReaL, M2PO, resume from saved tokens)",
      size=15, color=MUTED, align="left")

d.save(OUT)
# ZN-24: テンプレートのロゴ、著作権表示、ヘッダー、フッターを必ず消してから使う
sys.path.insert(0, "/Users/akazawt/works/data-science/claudecode/.claude/skills/review-html/packs/zenn-doc/templates")
import strip_branding  # noqa: E402
strip_branding.main(OUT, OUT)
print(OUT)
